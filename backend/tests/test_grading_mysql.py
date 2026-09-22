import os
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
from hashlib import sha256
import json
from io import BytesIO
from pathlib import Path
import re
import tempfile
from threading import Event
from time import sleep
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine, delete, select
from sqlalchemy.orm import Session

from app.common.models import DomainEventOutbox, FileObject
from app.main import app
from app.grading.models import (AnalyticsCourseSummary, AnalyticsLabSummary, AnalyticsSectionSummary, AnalyticsStudentLabSummary,
    AuditEvent, CourseArchive, CourseArchiveArtifact, GradeEvent, Gradebook, GradebookItem, GradingPolicy,
    GradingPolicyItem, StudentCourseScore, StudentRiskFlag)
from app.grading.service import GradingService

pytestmark = pytest.mark.skipif(not os.getenv("YUEKE_DATABASE_URL"), reason="需要专属 MySQL 集成库")
COURSE, CLASS = "course_data_security", "class_2301"
MANAGER = {"X-User-Id":"teacher_f","X-Role":"teacher","X-Teacher-Id":"teacher_f","X-Course-Ids":COURSE,"X-Class-Ids":CLASS,"X-Permissions":"grading:read,grading:policy,grading:recalculate,grading:post,analytics:class,analytics:read,archives:read,archives:write,archives:freeze,audit:read"}
STUDENT = {"X-User-Id":"user_s1","X-Role":"student","X-Student-Id":"student_1","X-Course-Ids":COURSE,"X-Class-Ids":CLASS,"X-Permissions":"grading:read,analytics:read"}
SERVICE = {"X-User-Id":"service_event_consumer","X-Role":"admin","X-Permissions":"grading:consume,audit:ingest"}
BROWSER_SPOOF = {"X-User-Id":"teacher_f","X-Role":"teacher","X-Teacher-Id":"teacher_f","X-Permissions":"grading:consume,audit:ingest"}


def remove_test_archive_files(items):
    allowed_roots = [Path(tempfile.gettempdir()).resolve(), (Path(__file__).parents[1] / "var" / "resource_uploads").resolve()]
    for item in items:
        path = Path(item.object_key).resolve()
        structurally_valid = bool(re.fullmatch(r"[0-9a-f]{64}", path.name)) and path.parent.parent.name == "objects" and path.parent.parent.parent.name == "course_archives"
        if not structurally_valid or not any(root == path or root in path.parents for root in allowed_roots):
            continue
        try:path.chmod(0o600);path.unlink(missing_ok=True)
        except OSError:pass


@pytest.fixture(autouse=True)
def clean_database():
    if not os.getenv("YUEKE_DATABASE_URL"): yield; return
    engine=create_engine(os.environ["YUEKE_DATABASE_URL"])
    order=[CourseArchiveArtifact,CourseArchive,StudentRiskFlag,AnalyticsLabSummary,AnalyticsStudentLabSummary,AnalyticsSectionSummary,AnalyticsCourseSummary,StudentCourseScore,GradebookItem,Gradebook,GradeEvent,GradingPolicyItem,GradingPolicy,AuditEvent]
    with Session(engine) as session:
        for model in order: session.execute(delete(model))
        archive_files=list(session.scalars(select(FileObject).where(FileObject.bucket=="course-archives")))
        session.execute(delete(DomainEventOutbox));session.execute(delete(FileObject).where(FileObject.bucket=="course-archives"));session.execute(delete(FileObject).where(FileObject.storage_provider == "generated-api"));session.commit()
        remove_test_archive_files(archive_files)
    yield
    with Session(engine) as session:
        for model in order: session.execute(delete(model))
        archive_files=list(session.scalars(select(FileObject).where(FileObject.bucket=="course-archives")))
        session.execute(delete(DomainEventOutbox));session.execute(delete(FileObject).where(FileObject.bucket=="course-archives"));session.execute(delete(FileObject).where(FileObject.storage_provider == "generated-api"));session.commit()
        remove_test_archive_files(archive_files)


@pytest.fixture()
def client(): return TestClient(app)


def envelope(event_type, student, source, raw, maximum=100, lesson="lesson_3_2", lab_release_id=None):
    payload={"course_id":COURSE,"class_id":CLASS,"student_id":student,"lesson_id":lesson,"source_id":source,"raw_score":raw,"max_score":maximum}
    if event_type=="lab.submitted":payload["lab_release_id"]=lab_release_id or source
    return {"event_id":str(uuid4()),"event_type":event_type,"aggregate_type":"source_fact","aggregate_id":source,"actor_user_id":"upstream","occurred_at":datetime.now(timezone.utc).isoformat(),"idempotency_key":str(uuid4()),"payload":payload}


def seed_complete_facts(client):
    rows=[]
    for student,scores in {"student_1":[100,80,90,100,70],"student_2":[50,60,70,80,50]}.items():
        for event_type,source,score in zip(["attendance.completed","assignment.submitted","quiz.completed","lab.submitted","poll.completed"],[f"att_{student}",f"ass_{student}",f"quiz_{student}","lab_rsa",f"poll_{student}"],scores):
            data=envelope(event_type,student,source,score);rows.append(data)
            response=client.post("/api/v1/grading/events/consume",headers=SERVICE,json=data);assert response.status_code==200
    return rows


def freeze_envelope(event_type, *, class_id=CLASS, member_count=2):
    if event_type=="course.roster.frozen":
        aggregate_id=class_id
        payload={"course_id":COURSE,"class_id":class_id,"member_count":member_count,"snapshot_hash":sha256(f"{class_id}:{member_count}".encode()).hexdigest(),"frozen_at":datetime.now(timezone.utc).isoformat()}
    else:
        aggregate_id=COURSE
        payload={"course_id":COURSE,"manifest_id":"resource_manifest_1","version_no":1}
    return {"event_id":str(uuid4()),"event_type":event_type,"aggregate_type":"upstream_manifest","aggregate_id":aggregate_id,"actor_user_id":"upstream","occurred_at":datetime.now(timezone.utc).isoformat(),"idempotency_key":str(uuid4()),"payload":payload}


def consume_freeze(client, event_type, **kwargs):
    response=client.post("/api/v1/grading/events/consume",headers=SERVICE,json=freeze_envelope(event_type,**kwargs))
    assert response.status_code==200
    return response


def test_event_replay_three_times_creates_one_grade_event(client):
    data=envelope("quiz.completed","student_1","quiz_once",88)
    assert client.post("/api/v1/grading/events/consume",headers=MANAGER,json=data).status_code==403
    assert client.post("/api/v1/grading/events/consume",headers=BROWSER_SPOOF,json=data).status_code==403
    responses=[client.post("/api/v1/grading/events/consume",headers=SERVICE,json=data).json() for _ in range(3)]
    assert [x["status"] for x in responses]==["CONSUMED","DUPLICATE","DUPLICATE"]
    engine=create_engine(os.environ["YUEKE_DATABASE_URL"])
    with Session(engine) as session: assert len(list(session.query(GradeEvent)))==1


def test_concurrent_event_replay_returns_duplicate_instead_of_database_error():
    data=envelope("quiz.completed","student_1","quiz_concurrent",88)
    def consume_once():
        with TestClient(app) as local_client:
            response=local_client.post("/api/v1/grading/events/consume",headers=SERVICE,json=data)
            return response.status_code,response.json()
    with ThreadPoolExecutor(max_workers=6) as pool:
        results=list(pool.map(lambda _:consume_once(),range(6)))
    assert all(status==200 for status,_ in results)
    assert sorted(body["status"] for _,body in results)==["CONSUMED",*["DUPLICATE"]*5]
    engine=create_engine(os.environ["YUEKE_DATABASE_URL"])
    with Session(engine) as session:assert len(list(session.scalars(select(GradeEvent).where(GradeEvent.event_id==data["event_id"]))))==1


def test_invalid_event_records_one_replayable_error_state(client):
    data=envelope("quiz.completed","student_1","quiz_broken",88);data["payload"].pop("max_score")
    responses=[client.post("/api/v1/grading/events/consume",headers=SERVICE,json=data) for _ in range(3)]
    assert all(x.status_code==422 and x.json()["code"]=="GRADING.EVENT_PAYLOAD_INVALID" for x in responses)
    rows=client.get("/api/v1/audit/events",headers=MANAGER,params={"action":"GRADE_EVENT_REJECTED"}).json()
    assert rows["total"]==1 and rows["items"][0]["result"]=="ERROR"


@pytest.mark.parametrize(("field","value","code"),[("source_id","x"*37,"GRADING.EVENT_PAYLOAD_INVALID"),("raw_score","NaN","GRADING.SCORE_INVALID")])
def test_invalid_event_values_are_rejected_before_mysql(client,field,value,code):
    data=envelope("quiz.completed","student_1","quiz_invalid",88);data["payload"][field]=value
    response=client.post("/api/v1/grading/events/consume",headers=SERVICE,json=data)
    assert response.status_code==422 and response.json()["code"]==code


def test_lab_submission_requires_release_id_and_keeps_source_for_trace(client):
    invalid=envelope("lab.submitted","student_1","submission_without_release",90)
    invalid["payload"].pop("lab_release_id")
    rejected=client.post("/api/v1/grading/events/consume",headers=SERVICE,json=invalid)
    assert rejected.status_code==422 and rejected.json()["details"]["missing"]==["lab_release_id"]
    valid=envelope("lab.submitted","student_1","submission_1",90,lab_release_id="lab_release_rsa")
    assert client.post("/api/v1/grading/events/consume",headers=SERVICE,json=valid).status_code==200
    assert client.post(f"/api/v1/grading/courses/{COURSE}/recalculate",headers=MANAGER,json={"class_id":CLASS}).status_code==200
    labs=client.get(f"/api/v1/analytics/courses/{COURSE}/labs/by-lab",headers=MANAGER,params={"class_id":CLASS}).json()["items"]
    trace=client.get(f"/api/v1/gradebook/courses/{COURSE}/trace/student_1",headers=MANAGER,params={"class_id":CLASS}).json()
    lab_source=next(x for x in trace["components"] if x["component"]=="LAB")["sources"][0]
    assert labs[0]["lab_release_id"]=="lab_release_rsa"
    assert lab_source["source_id"]=="submission_1" and lab_source["lab_release_id"]=="lab_release_rsa"


def test_lab_component_requires_submission_and_ignores_cumulative_checkpoint_snapshots(client):
    release_id="lab_release_rsa"
    for index,score in enumerate([20,40,70,90,100],start=1):
        checkpoint=envelope("lab.checkpoint.passed","student_1",f"checkpoint_{index}",score)
        checkpoint["payload"].update({"lab_release_id":release_id,"runtime_instance_id":"runtime_1","checkpoint_id":f"cp_{index}"})
        assert client.post("/api/v1/grading/events/consume",headers=SERVICE,json=checkpoint).status_code==200
    assert client.post(f"/api/v1/grading/courses/{COURSE}/recalculate",headers=MANAGER,json={"class_id":CLASS}).status_code==200
    trace=client.get(f"/api/v1/gradebook/courses/{COURSE}/trace/student_1",headers=MANAGER,params={"class_id":CLASS}).json()
    assert next(item for item in trace["components"] if item["component"]=="LAB")["score"]==0
    submitted=envelope("lab.submitted","student_1","submission_final",100,lab_release_id=release_id)
    assert client.post("/api/v1/grading/events/consume",headers=SERVICE,json=submitted).status_code==200
    rejudge=envelope("lab.checkpoint.passed","student_1","checkpoint_rejudge",50)
    rejudge["payload"].update({"lab_release_id":release_id,"runtime_instance_id":"runtime_1","checkpoint_id":"cp_3"})
    assert client.post("/api/v1/grading/events/consume",headers=SERVICE,json=rejudge).status_code==200
    assert client.post(f"/api/v1/grading/courses/{COURSE}/recalculate",headers=MANAGER,json={"class_id":CLASS}).status_code==200
    trace=client.get(f"/api/v1/gradebook/courses/{COURSE}/trace/student_1",headers=MANAGER,params={"class_id":CLASS}).json()
    assert next(item for item in trace["components"] if item["component"]=="LAB")["score"]==100


def test_fact_snapshot_blocks_stale_post_and_isolates_events_after_post(client,monkeypatch,tmp_path):
    seed_complete_facts(client);consume_freeze(client,"course.roster.frozen")
    assert client.post(f"/api/v1/grading/courses/{COURSE}/recalculate",headers=MANAGER,json={"class_id":CLASS}).status_code==200
    unapplied=envelope("quiz.completed","student_1","quiz_after_calculation",10)
    assert client.post("/api/v1/grading/events/consume",headers=SERVICE,json=unapplied).json()["status"]=="CONSUMED"
    trace=client.get(f"/api/v1/gradebook/courses/{COURSE}/trace/student_1",headers=MANAGER,params={"class_id":CLASS}).json()
    assert unapplied["event_id"] not in trace["source_snapshot_event_ids"]
    assert all(source["event_id"]!=unapplied["event_id"] for component in trace["components"] for source in component["sources"])
    stale=client.post(f"/api/v1/grading/courses/{COURSE}/post",headers=MANAGER,params={"class_id":CLASS})
    assert stale.status_code==409 and stale.json()["code"]=="GRADING.GRADEBOOK_STALE"
    assert unapplied["event_id"] in stale.json()["details"]["unapplied_source_event_ids"]
    assert client.post(f"/api/v1/grading/courses/{COURSE}/recalculate",headers=MANAGER,json={"class_id":CLASS}).status_code==200
    assert client.post(f"/api/v1/grading/courses/{COURSE}/post",headers=MANAGER,params={"class_id":CLASS}).status_code==200
    late=envelope("quiz.completed","student_1","quiz_after_post",100)
    late_response=client.post("/api/v1/grading/events/consume",headers=SERVICE,json=late)
    assert late_response.status_code==200 and late_response.json()["status"]=="LATE_IGNORED"
    consume_freeze(client,"resource.delivery.frozen")
    monkeypatch.setenv("YUEKE_RESOURCE_UPLOAD_DIR",str(tmp_path))
    frozen=client.post(f"/api/v1/archives/courses/{COURSE}/freeze",headers=MANAGER,json={"class_id":CLASS})
    assert frozen.status_code==200
    assert unapplied["event_id"] in frozen.json()["source_event_ids"]
    assert late["event_id"] not in frozen.json()["source_event_ids"]
    engine=create_engine(os.environ["YUEKE_DATABASE_URL"])
    with Session(engine) as session:
        stored=session.scalar(select(GradeEvent).where(GradeEvent.event_id==late["event_id"]))
        assert stored and stored.status=="LATE"
        assert session.scalar(select(AuditEvent).where(AuditEvent.action=="GRADE_EVENT_LATE_IGNORED",AuditEvent.resource_id==stored.grade_event_id))


def test_post_and_fact_ingestion_share_the_gradebook_scope_lock(client, monkeypatch):
    seed_complete_facts(client);consume_freeze(client,"course.roster.frozen")
    assert client.post(f"/api/v1/grading/courses/{COURSE}/recalculate",headers=MANAGER,json={"class_id":CLASS}).status_code==200
    post_holds_lock=Event();release_post=Event();consume_started=Event()
    original_integrity=GradingService.gradebook_integrity

    def paused_integrity(service, book):
        post_holds_lock.set()
        if not release_post.wait(5):
            raise RuntimeError("测试未释放成绩册范围锁")
        return original_integrity(service,book)

    monkeypatch.setattr(GradingService,"gradebook_integrity",paused_integrity)
    incoming=envelope("quiz.completed","student_1","quiz_racing_with_post",100)

    def post_gradebook():
        with TestClient(app) as local_client:
            return local_client.post(f"/api/v1/grading/courses/{COURSE}/post",headers=MANAGER,params={"class_id":CLASS})

    def consume_fact():
        consume_started.set()
        with TestClient(app) as local_client:
            return local_client.post("/api/v1/grading/events/consume",headers=SERVICE,json=incoming)

    with ThreadPoolExecutor(max_workers=2) as pool:
        post_future=pool.submit(post_gradebook)
        assert post_holds_lock.wait(5)
        consume_future=pool.submit(consume_fact)
        assert consume_started.wait(5)
        try:
            sleep(0.25)
            assert not consume_future.done(),"成绩事件未等待入账事务的范围锁"
        finally:
            release_post.set()
        posted=post_future.result(timeout=10)
        late=consume_future.result(timeout=10)

    assert posted.status_code==200 and posted.json()["status"]=="POSTED"
    assert late.status_code==200 and late.json()["status"]=="LATE_IGNORED"
    engine=create_engine(os.environ["YUEKE_DATABASE_URL"])
    with Session(engine) as session:
        stored=session.scalar(select(GradeEvent).where(GradeEvent.event_id==incoming["event_id"]))
        assert stored and stored.status=="LATE"


def test_recalculate_cannot_overwrite_a_concurrent_post(client, monkeypatch):
    seed_complete_facts(client);consume_freeze(client,"course.roster.frozen")
    assert client.post(f"/api/v1/grading/courses/{COURSE}/recalculate",headers=MANAGER,json={"class_id":CLASS}).status_code==200
    recalculate_holds_lock=Event();release_recalculate=Event();post_started=Event()
    original_build_risks=GradingService.build_risks

    def paused_build_risks(service, *args, **kwargs):
        if not recalculate_holds_lock.is_set():
            recalculate_holds_lock.set()
            if not release_recalculate.wait(5):
                raise RuntimeError("测试未释放重算范围锁")
        return original_build_risks(service,*args,**kwargs)

    monkeypatch.setattr(GradingService,"build_risks",paused_build_risks)

    def recalculate_gradebook():
        with TestClient(app) as local_client:
            return local_client.post(f"/api/v1/grading/courses/{COURSE}/recalculate",headers=MANAGER,json={"class_id":CLASS})

    def post_gradebook():
        post_started.set()
        with TestClient(app) as local_client:
            return local_client.post(f"/api/v1/grading/courses/{COURSE}/post",headers=MANAGER,params={"class_id":CLASS})

    with ThreadPoolExecutor(max_workers=2) as pool:
        recalculate_future=pool.submit(recalculate_gradebook)
        assert recalculate_holds_lock.wait(5)
        post_future=pool.submit(post_gradebook)
        assert post_started.wait(5)
        try:
            sleep(0.25)
            assert not post_future.done(),"成绩入账未等待正在执行的重算事务"
        finally:
            release_recalculate.set()
        recalculated=recalculate_future.result(timeout=10)
        posted=post_future.result(timeout=10)

    assert recalculated.status_code==200 and recalculated.json()["status"]=="READY"
    assert posted.status_code==200 and posted.json()["status"]=="POSTED"
    engine=create_engine(os.environ["YUEKE_DATABASE_URL"])
    with Session(engine) as session:
        book=session.scalar(select(Gradebook).where(Gradebook.course_id==COURSE,Gradebook.class_id==CLASS))
        assert book and book.status=="POSTED"


def test_policy_weight_recalculate_trace_post_analytics_and_student_scope(client):
    bad=client.put(f"/api/v1/grading/policies/{COURSE}",headers=MANAGER,json={"attendance":10,"assignment":20,"quiz":20,"lab":30,"interaction":10})
    assert bad.status_code==422 and bad.json()["code"]=="GRADING.POLICY_WEIGHT_INVALID"
    good=client.put(f"/api/v1/grading/policies/{COURSE}",headers=MANAGER,json={"attendance":10,"assignment":20,"quiz":20,"lab":40,"interaction":10})
    assert good.status_code==200 and sum(x["weight_percent"] for x in good.json()["items"])==100
    seed_complete_facts(client)
    calc=client.post(f"/api/v1/grading/courses/{COURSE}/recalculate",headers=MANAGER,json={"class_id":CLASS})
    assert calc.status_code==200 and calc.json()["status"]=="READY"
    hidden=client.get(f"/api/v1/gradebook/courses/{COURSE}/students/student_1",headers=STUDENT,params={"class_id":CLASS}).json()
    assert hidden["status"]=="PENDING" and hidden["items"]==[]
    assert client.get(f"/api/v1/gradebook/courses/{COURSE}",headers=MANAGER).status_code==422
    assert client.get(f"/api/v1/gradebook/courses/{COURSE}",headers=MANAGER,params={"class_id":"class_other"}).status_code==403
    book=client.get(f"/api/v1/gradebook/courses/{COURSE}",headers=MANAGER,params={"class_id":CLASS}).json()
    assert [x["total_score"] for x in book["items"]]==[91.0,68.0]
    trace=client.get(f"/api/v1/gradebook/courses/{COURSE}/trace/student_1",headers=MANAGER,params={"class_id":CLASS}).json()
    assert trace["total_score"]==91.0 and len(trace["components"])==5
    overview=client.get(f"/api/v1/analytics/courses/{COURSE}/overview",headers=MANAGER,params={"class_id":CLASS}).json()
    assert overview["avg_assignment"]==70.0 and overview["avg_quiz"]==80.0 and overview["attendance_rate"]==75.0 and overview["course_average"]==79.5
    section=client.get(f"/api/v1/analytics/courses/{COURSE}/sections/lesson_3_2",headers=MANAGER,params={"class_id":CLASS}).json()
    assert section["status"]=="READY" and sum(section["distribution"].values())==2
    labs=client.get(f"/api/v1/analytics/courses/{COURSE}/labs/by-lab",headers=MANAGER,params={"class_id":CLASS}).json()
    assert labs["items"][0]["lab_release_id"]=="lab_rsa" and labs["items"][0]["submitted_students"]==2 and labs["items"][0]["avg_score"]==90.0
    consume_freeze(client,"course.roster.frozen")
    posted=client.post(f"/api/v1/grading/courses/{COURSE}/post",headers=MANAGER,params={"class_id":CLASS})
    assert posted.status_code==200 and posted.json()["status"]=="POSTED"
    own=client.get(f"/api/v1/gradebook/courses/{COURSE}/students/student_1",headers=STUDENT,params={"class_id":CLASS})
    denied=client.get(f"/api/v1/gradebook/courses/{COURSE}/students/student_2",headers=STUDENT,params={"class_id":CLASS})
    assert own.status_code==200 and own.json()["items"][0]["total_score"]==91.0
    assert denied.status_code==403
    assert client.get(f"/api/v1/analytics/courses/{COURSE}/learning-summary",headers=STUDENT,params={"student_id":"student_1"}).status_code==422
    summary=client.get(f"/api/v1/analytics/courses/{COURSE}/learning-summary",headers=STUDENT,params={"class_id":CLASS,"student_id":"student_1"}).json()
    assert summary["status"]=="READY" and summary["grade"]["items"][0]["student_id"]=="student_1"
    gradebook_xlsx=client.get(f"/api/v1/gradebook/courses/{COURSE}/export.xlsx",headers=MANAGER,params={"class_id":CLASS}).content
    assert gradebook_xlsx.startswith(b"PK")
    gradebook_sheet=load_workbook(BytesIO(gradebook_xlsx),read_only=True)["成绩册"]
    assert [cell.value for cell in next(iter(gradebook_sheet.iter_rows(min_row=4,max_row=4)))][0:8]==["学生标识","签到","作业","测验","实验","互动","人工调整","总评"]
    analytics_xlsx=client.get(f"/api/v1/analytics/courses/{COURSE}/export.xlsx",headers=MANAGER,params={"class_id":CLASS}).content
    assert {"课程总览","课程排行","小节成绩","学生实验完成","实验完成统计","风险依据"}<=set(load_workbook(BytesIO(analytics_xlsx),read_only=True).sheetnames)


def test_trace_includes_manual_adjustment_that_explains_total(client):
    seed_complete_facts(client)
    adjustment=envelope("grade.manual.adjusted","student_1","manual_bonus",5)
    assert client.post("/api/v1/grading/events/consume",headers=SERVICE,json=adjustment).status_code==200
    assert client.post(f"/api/v1/grading/courses/{COURSE}/recalculate",headers=MANAGER,json={"class_id":CLASS}).status_code==200
    trace=client.get(f"/api/v1/gradebook/courses/{COURSE}/trace/student_1",headers=MANAGER,params={"class_id":CLASS}).json()
    manual=next(item for item in trace["components"] if item["component"]=="MANUAL_ADJUSTMENT")
    assert trace["total_score"]==96.0 and manual["score"]==5.0 and manual["sources"][0]["source_id"]=="manual_bonus"


def test_negative_manual_adjustment_is_audited_and_explains_total(client):
    seed_complete_facts(client)
    adjustment=envelope("grade.manual.adjusted","student_1","manual_penalty",-5)
    assert client.post("/api/v1/grading/events/consume",headers=SERVICE,json=adjustment).status_code==200
    assert client.post(f"/api/v1/grading/courses/{COURSE}/recalculate",headers=MANAGER,json={"class_id":CLASS}).status_code==200
    trace=client.get(f"/api/v1/gradebook/courses/{COURSE}/trace/student_1",headers=MANAGER,params={"class_id":CLASS}).json()
    manual=next(item for item in trace["components"] if item["component"]=="MANUAL_ADJUSTMENT")
    assert trace["total_score"]==86.0 and manual["score"]==-5.0
    audit=client.get("/api/v1/audit/events",headers=MANAGER,params={"action":"GRADE_MANUAL_ADJUSTMENT"}).json()
    assert audit["total"]==1 and audit["items"][0]["student_id"]=="student_1"


def test_lab_analytics_uses_latest_submission_and_keeps_non_submitter(client):
    checkpoint=envelope("lab.checkpoint.failed","student_1","checkpoint_only",0)
    checkpoint["payload"]["lab_release_id"]="lab_retry"
    assert client.post("/api/v1/grading/events/consume",headers=SERVICE,json=checkpoint).status_code==200
    for score in [40,80]:
        assert client.post("/api/v1/grading/events/consume",headers=SERVICE,json=envelope("lab.submitted","student_2",f"submission_{score}",score,lab_release_id="lab_retry")).status_code==200
    assert client.post(f"/api/v1/grading/courses/{COURSE}/recalculate",headers=MANAGER,json={"class_id":CLASS}).status_code==200
    students={item["student_id"]:item for item in client.get(f"/api/v1/analytics/courses/{COURSE}/labs/by-student",headers=MANAGER,params={"class_id":CLASS}).json()["items"]}
    assert students["student_1"]=={"student_id":"student_1","sum_lab_score":0.0,"submitted_count":0,"unsubmitted_count":1}
    assert students["student_2"]=={"student_id":"student_2","sum_lab_score":80.0,"submitted_count":1,"unsubmitted_count":0}
    lab=client.get(f"/api/v1/analytics/courses/{COURSE}/labs/by-lab",headers=MANAGER,params={"class_id":CLASS}).json()["items"][0]
    assert lab["avg_score"]==80.0 and lab["submitted_students"]==1 and lab["unsubmitted_students"]==1


def test_post_blocks_missing_enabled_component_but_allows_zero_weight_component(client):
    for student in ["student_1","student_2"]:
        for event_type,source in [("attendance.completed","att"),("assignment.submitted","ass"),("quiz.completed","quiz"),("lab.submitted","lab")]:
            data=envelope(event_type,student,f"{source}_{student}",80,lab_release_id="lab_rsa")
            assert client.post("/api/v1/grading/events/consume",headers=SERVICE,json=data).status_code==200
    consume_freeze(client,"course.roster.frozen")
    assert client.post(f"/api/v1/grading/courses/{COURSE}/recalculate",headers=MANAGER,json={"class_id":CLASS}).status_code==200
    blocked=client.post(f"/api/v1/grading/courses/{COURSE}/post",headers=MANAGER,params={"class_id":CLASS})
    assert blocked.status_code==409 and blocked.json()["code"]=="GRADING.GRADEBOOK_INCOMPLETE"
    assert all(item["missing_components"]==["INTERACTION"] for item in blocked.json()["details"]["incomplete_students"])
    policy=client.put(f"/api/v1/grading/policies/{COURSE}",headers=MANAGER,json={"attendance":10,"assignment":20,"quiz":20,"lab":50,"interaction":0})
    assert policy.status_code==200
    recalculated=client.post(f"/api/v1/grading/courses/{COURSE}/recalculate",headers=MANAGER,json={"class_id":CLASS})
    assert recalculated.status_code==200
    posted=client.post(f"/api/v1/grading/courses/{COURSE}/post",headers=MANAGER,params={"class_id":CLASS})
    assert posted.status_code==200 and posted.json()["status"]=="POSTED"


def test_post_blocks_when_frozen_roster_count_exceeds_gradebook(client):
    seed_complete_facts(client);consume_freeze(client,"course.roster.frozen",member_count=3)
    assert client.post(f"/api/v1/grading/courses/{COURSE}/recalculate",headers=MANAGER,json={"class_id":CLASS}).status_code==200
    blocked=client.post(f"/api/v1/grading/courses/{COURSE}/post",headers=MANAGER,params={"class_id":CLASS})
    assert blocked.status_code==409 and blocked.json()["code"]=="GRADING.ROSTER_COUNT_MISMATCH"
    assert blocked.json()["details"]=={"roster_member_count":3,"gradebook_student_count":2,"snapshot_hash":sha256(f"{CLASS}:3".encode()).hexdigest()}


def test_archive_blocks_until_verified_snapshots_then_persists_fixed_artifacts(client,monkeypatch,tmp_path):
    seed_complete_facts(client)
    client.post(f"/api/v1/grading/courses/{COURSE}/recalculate",headers=MANAGER,json={"class_id":CLASS})
    denied_post=client.post(f"/api/v1/grading/courses/{COURSE}/post",headers=MANAGER,params={"class_id":CLASS})
    assert denied_post.status_code==409 and denied_post.json()["code"]=="GRADING.ROSTER_FREEZE_REQUIRED"
    blocked=client.post(f"/api/v1/archives/courses/{COURSE}/precheck",headers=MANAGER,json={"class_id":CLASS}).json()
    assert set(blocked["blocking_items"])=={"roster_frozen","roster_student_count_matches","gradebook_posted","resources_frozen"}
    shell=freeze_envelope("resource.delivery.frozen");shell["payload"]={"course_id":COURSE}
    rejected=client.post("/api/v1/grading/events/consume",headers=SERVICE,json=shell)
    assert rejected.status_code==422 and rejected.json()["code"]=="GRADING.UPSTREAM_FREEZE_PAYLOAD_INVALID"
    consume_freeze(client,"course.roster.frozen",class_id="class_other")
    wrong_class=client.post(f"/api/v1/archives/courses/{COURSE}/precheck",headers=MANAGER,json={"class_id":CLASS}).json()
    assert {"roster_frozen","roster_student_count_matches","gradebook_posted","resources_frozen"}==set(wrong_class["blocking_items"])
    consume_freeze(client,"course.roster.frozen")
    posted=client.post(f"/api/v1/grading/courses/{COURSE}/post",headers=MANAGER,params={"class_id":CLASS})
    assert posted.status_code==200 and posted.json()["status"]=="POSTED"
    still_blocked=client.post(f"/api/v1/archives/courses/{COURSE}/precheck",headers=MANAGER,json={"class_id":CLASS}).json()
    assert still_blocked["blocking_items"]==["resources_frozen"]
    consume_freeze(client,"resource.delivery.frozen")
    ready=client.post(f"/api/v1/archives/courses/{COURSE}/precheck",headers=MANAGER,json={"class_id":CLASS}).json();assert ready["blocking"]==0
    assert ready["roster_snapshot"]["member_count"]==2 and ready["roster_snapshot"]["snapshot_hash"]
    assert ready["resource_manifest"]["manifest_id"]=="resource_manifest_1" and ready["resource_manifest"]["version_no"]==1
    monkeypatch.setenv("YUEKE_RESOURCE_UPLOAD_DIR",str(tmp_path))
    frozen=client.post(f"/api/v1/archives/courses/{COURSE}/freeze",headers=MANAGER,json={"class_id":CLASS})
    assert frozen.status_code==200 and len(frozen.json()["artifacts"])==4
    artifacts={item["type"]:item for item in frozen.json()["artifacts"]}
    assert all(artifacts[k].get("file_id") for k in ["GRADEBOOK_XLSX","ANALYTICS_XLSX","RESOURCE_VERSION_MANIFEST_JSON"])
    downloads={}
    for kind in ["GRADEBOOK_XLSX","ANALYTICS_XLSX","RESOURCE_VERSION_MANIFEST_JSON"]:
        response=client.get(artifacts[kind]["evidence_ref"],headers=MANAGER)
        assert response.status_code==200 and sha256(response.content).hexdigest()==artifacts[kind]["sha256"]
        assert response.headers["cache-control"]=="private, immutable"
        downloads[kind]=response.content
    resource_snapshot=json.loads(downloads["RESOURCE_VERSION_MANIFEST_JSON"])
    assert resource_snapshot["delivery_manifests"][0]["manifest_id"]=="resource_manifest_1"
    engine=create_engine(os.environ["YUEKE_DATABASE_URL"])
    with Session(engine) as session:
        files=list(session.scalars(select(FileObject).where(FileObject.bucket=="course-archives")))
        assert len(files)==3 and all(item.storage_provider=="local" and Path(item.object_key).is_file() for item in files)
        file_paths={item.file_id:Path(item.object_key) for item in files}
        score=session.scalar(select(StudentCourseScore).where(StudentCourseScore.student_id=="student_1"));score.total_score=1;session.commit()
    fixed=client.get(artifacts["GRADEBOOK_XLSX"]["evidence_ref"],headers=MANAGER)
    assert fixed.content==downloads["GRADEBOOK_XLSX"]
    tampered=file_paths[artifacts["RESOURCE_VERSION_MANIFEST_JSON"]["file_id"]];tampered.chmod(0o600);tampered.write_bytes(b"tampered")
    rejected_download=client.get(artifacts["RESOURCE_VERSION_MANIFEST_JSON"]["evidence_ref"],headers=MANAGER)
    assert rejected_download.status_code==409 and rejected_download.json()["code"]=="ARCHIVE.ARTIFACT_INTEGRITY_FAILED"
    locked=client.post(f"/api/v1/grading/courses/{COURSE}/recalculate",headers=MANAGER,json={"class_id":CLASS})
    assert locked.status_code==409 and locked.json()["code"]=="ARCHIVE.GRADEBOOK_LOCKED"
    audits=client.get("/api/v1/audit/events",headers=MANAGER,params={"course_id":COURSE}).json()
    actions={x["action"] for x in audits["items"]}
    assert {"GRADEBOOK_RECALCULATED","GRADEBOOK_POSTED","COURSE_ARCHIVE_PRECHECK","COURSE_ARCHIVED"}<=actions
    assert client.delete("/api/v1/audit/events",headers=MANAGER).status_code==405
    assert client.get("/api/v1/audit/events/export.xlsx",headers=MANAGER).content.startswith(b"PK")
    assert client.get("/api/v1/audit/events/export.csv",headers=MANAGER).content.startswith(b"\xef\xbb\xbf")


def test_cross_domain_high_risk_audit_is_idempotent_and_immutable(client):
    data={"source_event_id":"runtime-destroyed-1","actor_user_id":"admin_d","actor_role":"admin","action":"RUNTIME_INSTANCE_DESTROYED","resource_type":"runtime_instance","resource_id":"instance_1","course_id":COURSE,"class_id":CLASS,"student_id":"student_1","result":"SUCCESS","reason":"教师确认重建","occurred_at":datetime.now(timezone.utc).isoformat(),"details":{"node_id":"node_1"}}
    assert client.post("/api/v1/audit/events/ingest",headers=MANAGER,json=data).status_code==403
    assert client.post("/api/v1/audit/events/ingest",headers=BROWSER_SPOOF,json=data).status_code==403
    results=[client.post("/api/v1/audit/events/ingest",headers=SERVICE,json=data).json()["status"] for _ in range(3)]
    assert results==["RECORDED","DUPLICATE","DUPLICATE"]
    rows=client.get("/api/v1/audit/events",headers=MANAGER,params={"action":"RUNTIME_INSTANCE_DESTROYED"}).json()
    assert rows["total"]==1 and rows["items"][0]["reason"]=="教师确认重建"
    assert client.patch(f"/api/v1/audit/events/{rows['items'][0]['audit_event_id']}",headers=MANAGER,json={"reason":"篡改"}).status_code in {404,405}


def test_audit_query_and_exports_are_limited_to_authorized_courses(client):
    own={"source_event_id":"audit-own","actor_user_id":"admin_d","actor_role":"admin","action":"RUNTIME_INSTANCE_DESTROYED","resource_type":"runtime_instance","resource_id":"instance_1","course_id":COURSE,"class_id":CLASS,"student_id":"student_1","result":"SUCCESS","occurred_at":datetime.now(timezone.utc).isoformat(),"details":{}}
    foreign={**own,"source_event_id":"audit-foreign","resource_id":"instance_2","course_id":"course_foreign","class_id":"class_foreign"}
    assert client.post("/api/v1/audit/events/ingest",headers=SERVICE,json=own).status_code==200
    assert client.post("/api/v1/audit/events/ingest",headers=SERVICE,json=foreign).status_code==200
    scoped=client.get("/api/v1/audit/events",headers=MANAGER).json()
    assert scoped["total"]==1 and scoped["items"][0]["course_id"]==COURSE
    assert client.get("/api/v1/audit/events",headers=MANAGER,params={"course_id":"course_foreign"}).status_code==403
    sheet=load_workbook(BytesIO(client.get("/api/v1/audit/events/export.xlsx",headers=MANAGER).content),read_only=True)["审计"]
    exported=list(sheet.iter_rows(min_row=5,values_only=True))
    assert len(exported)==1 and exported[0][6]==COURSE
    admin={"X-User-Id":"admin_f","X-Role":"admin","X-Permissions":"audit:read,grading:all-courses,grading:all-classes"}
    assert client.get("/api/v1/audit/events",headers=admin).json()["total"]==2


def test_risk_flags_are_backed_by_checkpoint_and_submission_facts(client):
    for source in ["checkpoint_1","checkpoint_2"]:
        assert client.post("/api/v1/grading/events/consume",headers=SERVICE,json=envelope("lab.checkpoint.failed","student_1",source,0)).status_code==200
    for source in ["lab_1","lab_2","lab_3"]:
        assert client.post("/api/v1/grading/events/consume",headers=SERVICE,json=envelope("lab.submitted","student_2",f"submission_{source}",80,lab_release_id=source)).status_code==200
    assert client.post(f"/api/v1/grading/courses/{COURSE}/recalculate",headers=MANAGER,json={"class_id":CLASS}).status_code==200
    risks=client.get(f"/api/v1/analytics/courses/{COURSE}/risks",headers=MANAGER,params={"class_id":CLASS,"student_id":"student_1"}).json()["items"]
    by_type={x["risk_type"]:x["evidence"] for x in risks}
    assert by_type["CONSECUTIVE_CHECKPOINT_FAILURES"]["failure_count"]==2
    assert by_type["MULTIPLE_MISSING_SUBMISSIONS"]["missing_count"]==3


def test_learning_summary_is_pending_without_upstream_facts(client):
    result=client.get(f"/api/v1/analytics/courses/{COURSE}/learning-summary",headers=MANAGER,params={"class_id":CLASS}).json()
    assert result["status"]=="PENDING" and result["missing_upstream"]
