import os
from datetime import datetime
from io import BytesIO
from pathlib import Path
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine, delete, select
from sqlalchemy.orm import Session

from app.common.models import DomainEventOutbox
from app.main import app
from app.resources.models import (
    Question,
    QuestionBank,
    QuestionExplanation,
    QuestionImportJob,
    QuestionImportRow,
    QuestionLessonMap,
    QuestionOption,
    QuestionReview,
    ResourceQualityCheck,
)
from app.teaching.catalog import curriculum_rows
from app.teaching.models import Course, CourseChapter, CourseLesson


pytestmark = pytest.mark.skipif(not os.getenv("YUEKE_DATABASE_URL"), reason="需要专属 MySQL 集成库")

ROOT = Path(__file__).resolve().parents[2]
WORKBOOK_PATH = ROOT / "outputs" / "01a0c33f-d483-7ac0-aa95-632194b582d5" / "question-bank-196-v1.xlsx"


def headers(user_id: str, course_id: str) -> dict[str, str]:
    return {
        "X-User-Id": user_id,
        "X-Role": "teacher",
        "X-Teacher-Id": user_id,
        "X-Course-Ids": course_id,
        "X-Permissions": "resources:read,resources:write,resources:review,resources:freeze",
    }


def workbook_for_course(course_id: str) -> bytes:
    authority = curriculum_rows(course_id)
    lesson_ids = {row["lesson_code"]: row["lesson_id"] for row in authority["lessons"]}
    workbook = load_workbook(WORKBOOK_PATH)
    sheet = workbook["题目导入"]
    for row_number in range(2, 198):
        sheet.cell(row_number, 1).value = lesson_ids[str(sheet.cell(row_number, 2).value)]
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def test_formal_question_bank_real_mysql_import_and_independent_review():
    engine = create_engine(os.environ["YUEKE_DATABASE_URL"])
    client = TestClient(app)
    suffix = uuid4().hex[:8]
    course_id = f"course_formal_bank_{suffix}"
    author_headers = headers("content_author_b", course_id)
    reviewer_headers = headers("content_reviewer_b", course_id)
    job_id = bank_id = None
    question_ids: list[str] = []

    with Session(engine) as session:
        session.add(Course(course_id=course_id, name="数据安全技术基础", term="2026 秋季", owner_teacher_id="content_author_b", major="网络空间安全", description=None, status="ACTIVE", created_at=datetime.utcnow()))
        authority = curriculum_rows(course_id)
        session.add_all(CourseChapter(**row) for row in authority["chapters"])
        session.add_all(CourseLesson(**row) for row in authority["lessons"])
        session.commit()

    try:
        imported = client.post(
            "/api/v1/questions/import",
            headers={**author_headers, "Idempotency-Key": f"formal-bank-{suffix}"},
            data={"course_id": course_id},
            files={
                "file": (
                    WORKBOOK_PATH.name,
                    workbook_for_course(course_id),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        assert imported.status_code == 201
        import_job = imported.json()
        job_id = import_job["job_id"]
        assert import_job["status"] == "COMPLETED"
        assert import_job["total_count"] == 196
        assert import_job["imported_count"] == 196
        assert import_job["error_count"] == 0
        assert import_job["review_queue_count"] == 196

        queue = client.get(
            "/api/v1/questions/review-queue",
            headers=reviewer_headers,
            params={"course_id": course_id, "page_size": 200},
        )
        assert queue.status_code == 200
        items = queue.json()["items"]
        assert len(items) == 196
        assert all(item["can_review"] and item["created_by"] == "content_author_b" for item in items)

        self_review = client.post(
            f"/api/v1/questions/{items[0]['question_id']}/review",
            headers=author_headers,
            json={"decision": "APPROVED", "comment": "作者不得审核自己的题目"},
        )
        assert self_review.status_code == 409
        assert self_review.json()["code"] == "QUESTION.REVIEWER_MUST_BE_INDEPENDENT"

        for item in items:
            reviewed = client.post(
                f"/api/v1/questions/{item['question_id']}/review",
                headers=reviewer_headers,
                json={"decision": "APPROVED", "comment": "正式题库逐题字段、答案、解析与课时映射复核通过"},
            )
            assert reviewed.status_code == 200
            assert reviewed.json()["status"] == "PUBLISHED"

        queue_after = client.get(
            "/api/v1/questions/review-queue",
            headers=reviewer_headers,
            params={"course_id": course_id, "page_size": 200},
        )
        assert queue_after.status_code == 200
        assert queue_after.json()["total"] == 0

        coverage = client.get("/api/v1/questions/coverage", headers=reviewer_headers, params={"course_id": course_id})
        assert coverage.status_code == 200
        assert coverage.json()["total"] == coverage.json()["passed"] == 49

        readiness = client.get("/api/v1/resources/readiness", headers=reviewer_headers, params={"course_id": course_id})
        assert readiness.status_code == 200
        assert readiness.json()["question_lessons"] == {"ready": 49, "required": 49}
        assert readiness.json()["published_questions"] == {"ready": 196, "required": 196}

        audit = client.post("/api/v1/resources/audit/run", headers=reviewer_headers, json={"course_id": course_id})
        assert audit.status_code == 200
        assert audit.json()["total"] == 196
        assert audit.json()["pass"] == 49
        assert audit.json()["blocking"] == 147
        assert all(
            check["passed"]
            for check in audit.json()["checks"]
            if check["requirement"] == "QUESTION_BANK"
        )

        with Session(engine) as session:
            bank = session.scalar(select(QuestionBank).where(QuestionBank.course_id == course_id))
            bank_id = bank.question_bank_id
            question_ids = list(session.scalars(select(Question.question_id).where(Question.question_bank_id == bank_id)))
            assert len(question_ids) == 196
            assert len(list(session.scalars(select(QuestionReview).where(QuestionReview.question_id.in_(question_ids))))) == 196
    finally:
        with Session(engine) as session:
            if bank_id is None:
                bank = session.scalar(select(QuestionBank).where(QuestionBank.course_id == course_id))
                bank_id = bank.question_bank_id if bank else None
            if bank_id and not question_ids:
                question_ids = list(session.scalars(select(Question.question_id).where(Question.question_bank_id == bank_id)))
            job_ids = list(session.scalars(select(QuestionImportJob.import_job_id).where(QuestionImportJob.course_id == course_id)))
            if job_ids:
                session.execute(delete(QuestionImportRow).where(QuestionImportRow.import_job_id.in_(job_ids)))
            if question_ids:
                session.execute(delete(QuestionReview).where(QuestionReview.question_id.in_(question_ids)))
                session.execute(delete(QuestionOption).where(QuestionOption.question_id.in_(question_ids)))
                session.execute(delete(QuestionExplanation).where(QuestionExplanation.question_id.in_(question_ids)))
                session.execute(delete(QuestionLessonMap).where(QuestionLessonMap.question_id.in_(question_ids)))
                session.execute(delete(Question).where(Question.question_id.in_(question_ids)))
                session.execute(delete(DomainEventOutbox).where(DomainEventOutbox.aggregate_id.in_(question_ids)))
            if job_ids:
                session.execute(delete(DomainEventOutbox).where(DomainEventOutbox.aggregate_id.in_(job_ids)))
                session.execute(delete(QuestionImportJob).where(QuestionImportJob.import_job_id.in_(job_ids)))
            session.execute(delete(ResourceQualityCheck).where(ResourceQualityCheck.course_id == course_id))
            session.execute(delete(DomainEventOutbox).where(DomainEventOutbox.aggregate_id == course_id))
            if bank_id:
                session.execute(delete(QuestionBank).where(QuestionBank.question_bank_id == bank_id))
            session.execute(delete(CourseLesson).where(CourseLesson.course_id == course_id))
            session.execute(delete(CourseChapter).where(CourseChapter.course_id == course_id))
            session.execute(delete(Course).where(Course.course_id == course_id))
            session.commit()
