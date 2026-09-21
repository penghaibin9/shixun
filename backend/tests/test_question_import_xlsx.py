from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import load_workbook

from app.resources import question_xlsx
from app.resources.catalog import lesson_rows
from app.resources.question_xlsx import HEADERS, InvalidQuestionWorkbook, error_rows_bytes, parse_question_workbook, template_bytes


def completed_workbook() -> bytes:
    workbook = load_workbook(BytesIO(template_bytes(lesson_rows())))
    sheet = workbook["题目导入"]
    for row_number in range(2, 198):
        question_type = sheet.cell(row_number, 3).value
        sheet.cell(row_number, 4).value = f"第 {row_number - 1} 道{question_type}题"
        sheet.cell(row_number, 10).value = "这是逐题解析"
        if question_type == "填空":
            sheet.cell(row_number, 9).value = "参考答案"
        elif question_type == "单选":
            sheet.cell(row_number, 5).value = "正确选项"
            sheet.cell(row_number, 6).value = "干扰选项"
            sheet.cell(row_number, 9).value = "A"
        elif question_type == "多选":
            sheet.cell(row_number, 5).value = "正确选项一"
            sheet.cell(row_number, 6).value = "正确选项二"
            sheet.cell(row_number, 7).value = "干扰选项"
            sheet.cell(row_number, 9).value = "A,B"
        else:
            sheet.cell(row_number, 9).value = "A"
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def row_errors(rows: list[dict], row_number: int) -> set[str]:
    row = next(item for item in rows if item["row_number"] == row_number)
    return {error["code"] for error in row["errors"]}


def test_template_contains_exactly_49_by_4_prefilled_slots():
    workbook = load_workbook(BytesIO(template_bytes(lesson_rows())))
    sheet = workbook["题目导入"]

    assert [cell.value for cell in sheet[1]] == HEADERS
    assert sheet.max_row == 197
    assert sheet.freeze_panes == "A2"
    assert workbook.sheetnames == ["题目导入", "填写说明"]
    assert len({(sheet.cell(row, 1).value, sheet.cell(row, 3).value) for row in range(2, 198)}) == 196
    assert all(sheet.cell(row, 5).value == "正确" and sheet.cell(row, 6).value == "错误" for row in range(5, 198, 4))


def test_parser_accepts_a_complete_196_row_workbook():
    rows, total = parse_question_workbook(completed_workbook(), lesson_rows())

    assert total == 196
    assert len(rows) == 196
    assert all(row["status"] == "VALID" and not row["errors"] for row in rows)
    assert {row["normalized_data"]["question_type"] for row in rows} == {"FILL", "SINGLE", "MULTIPLE", "TRUE_FALSE"}


def test_parser_reports_formula_mismatch_and_answer_errors_by_excel_row():
    workbook = load_workbook(BytesIO(completed_workbook()))
    sheet = workbook["题目导入"]
    sheet["D2"] = "=1+1"
    sheet["B3"] = sheet["B7"].value
    sheet["I4"] = "A"
    stream = BytesIO()
    workbook.save(stream)

    rows, total = parse_question_workbook(stream.getvalue(), lesson_rows())

    assert total == 196
    assert "QUESTION_IMPORT.DANGEROUS_CELL" in row_errors(rows, 2)
    assert "QUESTION_IMPORT.LESSON_MISMATCH" in row_errors(rows, 3)
    assert "QUESTION_IMPORT.MULTIPLE_ANSWER_COUNT" in row_errors(rows, 4)

    exported = load_workbook(BytesIO(error_rows_bytes(rows)))['错误行']
    assert exported['E2'].data_type != 'f'
    assert exported['E2'].value == "'=1+1"


def test_parser_reports_file_level_count_and_missing_slot_errors():
    workbook = load_workbook(BytesIO(completed_workbook()))
    workbook["题目导入"].delete_rows(197)
    stream = BytesIO()
    workbook.save(stream)

    rows, total = parse_question_workbook(stream.getvalue(), lesson_rows())

    assert total == 195
    assert rows[0]["row_number"] == 0
    codes = {error["code"] for error in rows[0]["errors"]}
    assert codes == {"QUESTION_IMPORT.ROW_COUNT_INVALID", "QUESTION_IMPORT.SLOT_MISSING"}


def test_parser_bounds_declared_rows_without_iterating_the_entire_sheet():
    workbook = load_workbook(BytesIO(completed_workbook()))
    workbook["题目导入"].cell(100_000, 1).value = "超出模板范围"
    stream = BytesIO()
    workbook.save(stream)

    rows, total = parse_question_workbook(stream.getvalue(), lesson_rows())

    assert total == 99_999
    assert len(rows) == 197
    assert "QUESTION_IMPORT.ROW_COUNT_INVALID" in {error["code"] for error in rows[0]["errors"]}


def test_parser_rejects_an_xlsx_with_excessive_uncompressed_content(monkeypatch):
    source = BytesIO(template_bytes(lesson_rows()))
    expanded = BytesIO()
    with ZipFile(source) as original, ZipFile(expanded, "w", ZIP_DEFLATED) as rebuilt:
        for entry in original.infolist():
            rebuilt.writestr(entry, original.read(entry.filename))
        rebuilt.writestr("xl/oversized.bin", b"x" * 1024)
    monkeypatch.setattr(question_xlsx, "MAX_XLSX_UNCOMPRESSED_BYTES", 512)

    with pytest.raises(InvalidQuestionWorkbook) as captured:
        parse_question_workbook(expanded.getvalue(), lesson_rows())

    assert captured.value.code == "QUESTION_IMPORT.ARCHIVE_TOO_LARGE"
