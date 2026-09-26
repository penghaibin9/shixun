from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.common.context import UserContext
from app.common.models import Base
from app.main import app as _loaded_app  # noqa: F401 - register all model metadata
from app.resources.service import ResourceService
from app.teaching.catalog import catalog_metadata
from app.teaching.schemas import ClassCreate, CourseCreate
from app.teaching.service import TeachingService


def teacher_context(*, course_ids=frozenset(), class_ids=frozenset(), permissions=frozenset()) -> UserContext:
    return UserContext(
        user_id="teacher-multi-course",
        role="teacher",
        teacher_id="teacher-multi-course",
        student_id=None,
        permissions=permissions,
        course_ids=frozenset(course_ids),
        class_ids=frozenset(class_ids),
    )


def test_all_eleven_catalogs_create_course_class_resources_and_question_template():
    engine = create_engine(
        "sqlite+pysqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    Session = sessionmaker(bind=engine, expire_on_commit=False)

    expected = catalog_metadata()
    assert len(expected) == 11

    with Session() as session:
        for index, catalog in enumerate(expected, start=1):
            creator = TeachingService(
                session,
                teacher_context(permissions={"teaching.course.write"}),
            )
            created = creator.create_course(
                CourseCreate(
                    name=str(catalog["name"]),
                    term="2026 秋季",
                    major="网络空间安全",
                    catalog_key=str(catalog["catalog_key"]),
                )
            )
            course_id = created["course_id"]
            theory = int(catalog["theory_lessons"])
            labs = int(catalog["lab_lessons"])
            total = theory + labs

            assert created["catalog_key"] == catalog["catalog_key"]
            assert created["theory_lesson_count"] == theory
            assert created["lab_lesson_count"] == labs

            scoped = teacher_context(
                course_ids={course_id},
                permissions={"teaching.class.write", "teaching.course.read"},
            )
            teaching = TeachingService(session, scoped)
            classroom = teaching.create_class(
                ClassCreate(
                    name=f"多课程验收 {index:02d} 班",
                    term="2026 秋季",
                    course_id=course_id,
                )
            )
            assert classroom["course_id"] == course_id
            assert classroom["class_id"]

            resources = ResourceService(
                session,
                teacher_context(
                    course_ids={course_id},
                    permissions={"resources:read", "resources:write"},
                ),
            )
            resource_catalog = resources.lessons(course_id)
            assert resource_catalog["total"] == total
            assert sum(item["lesson_kind"] == "THEORY" for item in resource_catalog["items"]) == theory
            assert sum(item["lesson_kind"] == "LAB" for item in resource_catalog["items"]) == labs

            workbook = load_workbook(BytesIO(resources.question_template(course_id)), read_only=True)
            sheet = workbook["题目导入"]
            assert sheet.max_row == total * len(catalog["question_types"]) + 1
            lesson_codes = {
                str(sheet.cell(row, 2).value)
                for row in range(2, sheet.max_row + 1)
            }
            assert lesson_codes == {item["lesson_code"] for item in resource_catalog["items"]}

    Base.metadata.drop_all(engine)
    engine.dispose()
