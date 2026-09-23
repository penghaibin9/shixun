from datetime import datetime, timedelta
from decimal import Decimal
from io import BytesIO
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.common.models import Base, DomainEventOutbox
from app.common.errors import ApiError
from app.database import get_session
from app.main import app
from app.resources.models import Question, QuestionBank, QuestionLessonMap, QuestionOption
from app.teaching.models import AssignmentQuestionRef, AssignmentSubmission, ClassMembership, CourseLesson, ImportJob, TeachingScoreProof
from app.teaching.repository import TeachingRepository
from app.teaching.service import score_payload_sha256, sha256_json, verify_teaching_score_proof
from app.teaching.xlsx import HEADERS
from backend.tests.auth_profiles import seed_student_profile, seed_student_profiles


_active_sessions = None


@pytest.fixture()
def test_context():
    global _active_sessions
    engine = create_engine("sqlite+pysqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    sessions = sessionmaker(bind=engine, expire_on_commit=False)

    def override():
        with sessions() as session:
            yield session

    app.dependency_overrides[get_session] = override
    _active_sessions = sessions
    yield TestClient(app), sessions
    _active_sessions = None
    app.dependency_overrides.clear()
    Base.metadata.drop_all(engine)


def headers(*permissions: str, course_id: str = "", class_id: str = "", student_id: str = "", teacher_id: str = "teacher-a"):
    result = {"X-User-Id": student_id or teacher_id, "X-Role": "student" if student_id else "teacher", "X-Permissions": ",".join(permissions)}
    if teacher_id: result["X-Teacher-Id"] = teacher_id
    if student_id: result["X-Student-Id"] = student_id
    if course_id: result["X-Course-Ids"] = course_id
    if class_id: result["X-Class-Ids"] = class_id
    return result


def workbook_bytes(rows):
    # 名单导入只能解析既有学生档案；测试数据在生成 XLSX 时明确预置档案。
    if _active_sessions is not None:
        seed_student_profiles(_active_sessions, rows)
    book = Workbook(); sheet = book.active; sheet.append(HEADERS)
    for row in rows: sheet.append(row)
    stream = BytesIO(); book.save(stream); return stream.getvalue()


def build_course_class(client: TestClient):
    created = client.post("/api/v1/courses", headers=headers("teaching.course.write"), json={"name": "数据安全技术基础", "term": "2026 秋季", "major": "网络空间安全"})
    assert created.status_code == 201, created.text
    course_id = created.json()["course_id"]
    created_class = client.post("/api/v1/classes", headers=headers("teaching.class.write", course_id=course_id), json={"name": "网络安全 2301 班", "term": "2026 秋季", "course_id": course_id})
    assert created_class.status_code == 201, created_class.text
    return course_id, created_class.json()["class_id"]


def seed_published_question(sessions, course_id: str, *, question_id: str = "question-b-1", lesson_id: str | None = None, question_type: str = "SINGLE", answer: list | None = None):
    """Create a B-owned, independently reviewed question for A to freeze."""

    with sessions() as db:
        lesson_id = lesson_id or db.scalar(select(CourseLesson.lesson_id).where(CourseLesson.course_id == course_id).order_by(CourseLesson.sequence))
        bank = db.scalar(select(QuestionBank).where(QuestionBank.course_id == course_id))
        if not bank:
            bank = QuestionBank(question_bank_id=str(uuid4()), course_id=course_id, name="课程统一题库", status="PUBLISHED", created_by="resource-author", created_at=datetime.utcnow())
            db.add(bank)
            db.flush()
        values = answer or ["A"]
        db.add(
            Question(
                question_id=question_id,
                question_bank_id=bank.question_bank_id,
                import_job_id=None,
                source_row_number=None,
                question_type=question_type,
                stem="RSA 中公开的是哪一类密钥？",
                answer_json=values,
                status="PUBLISHED",
                created_by="resource-author",
                created_at=datetime.utcnow(),
                submitted_at=datetime.utcnow(),
                reviewed_by="resource-reviewer",
                reviewed_at=datetime.utcnow(),
            )
        )
        db.add(QuestionLessonMap(question_lesson_map_id=str(uuid4()), question_id=question_id, lesson_id=lesson_id))
        if question_type in {"SINGLE", "MULTIPLE", "TRUE_FALSE"}:
            db.add(QuestionOption(question_option_id=str(uuid4()), question_id=question_id, option_key="A", option_text="公钥", is_correct=True))
            db.add(QuestionOption(question_option_id=str(uuid4()), question_id=question_id, option_key="B", option_text="私钥", is_correct=False))
        db.commit()
    return question_id


def test_course_creation_builds_authoritative_49_lesson_catalog_and_rejects_cross_course_lesson(test_context):
    client, _ = test_context
    first_course, first_class = build_course_class(client)
    first_headers = headers("teaching.course.read", "teaching.attendance.write", course_id=first_course, class_id=first_class)
    catalog = client.get(f"/api/v1/courses/{first_course}/lessons", headers=first_headers)
    assert catalog.status_code == 200
    payload = catalog.json()
    assert (payload["total"], payload["theory_count"], payload["lab_count"]) == (49, 37, 12)
    assert [item["lesson_code"] for item in payload["items"] if item["chapter_sequence"] == 7] == ["7.1", "7.2", "7.3", "7.4"]
    assert payload["items"][-1]["lesson_code"] == "实验12"

    second_course, _ = build_course_class(client)
    foreign_lesson = client.get(
        f"/api/v1/courses/{second_course}/lessons",
        headers=headers("teaching.course.read", course_id=second_course),
    ).json()["items"][0]["lesson_id"]
    start, end = datetime.utcnow(), datetime.utcnow() + timedelta(minutes=30)
    denied = client.post(
        "/api/v1/attendance/tasks",
        headers=first_headers,
        json={
            "course_id": first_course,
            "class_id": first_class,
            "lesson_id": foreign_lesson,
            "task_type": "CLASSROOM",
            "title": "跨课程课时校验",
            "starts_at": start.isoformat(),
            "expires_at": end.isoformat(),
        },
    )
    assert denied.status_code == 422
    assert denied.json()["code"] == "COURSE.LESSON_SCOPE_MISMATCH"


def test_g1_course_class_xlsx_import_idempotency_and_error_rows(test_context):
    client, _ = test_context; course_id, class_id = build_course_class(client)
    scoped = headers("teaching.members.read", "teaching.members.import", course_id=course_id, class_id=class_id)
    template = client.get(f"/api/v1/classes/{class_id}/members/import-template", headers=scoped)
    assert template.status_code == 200
    assert load_workbook(BytesIO(template.content)).active["A1"].value == "学号*"
    rows = [[f"2301{i:03d}", f"学生{i}", "网络安全 2301 班", "", ""] for i in range(1, 44)]
    import_bytes = workbook_bytes(rows)
    imported = client.post(f"/api/v1/classes/{class_id}/members/import", headers={**scoped, "Idempotency-Key": "batch-43"}, files={"file": ("students.xlsx", import_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert imported.status_code == 201, imported.text
    assert imported.json()["success_count"] == 43
    replay = client.post(f"/api/v1/classes/{class_id}/members/import", headers={**scoped, "Idempotency-Key": "batch-43"}, files={"file": ("students.xlsx", import_bytes)})
    assert replay.json()["job_id"] == imported.json()["job_id"]
    conflict_rows = [["2399999", "不同文件", "网络安全 2301 班", "", ""]]
    conflict = client.post(f"/api/v1/classes/{class_id}/members/import", headers={**scoped, "Idempotency-Key": "batch-43"}, files={"file": ("students.xlsx", workbook_bytes(conflict_rows))})
    assert conflict.status_code == 409 and conflict.json()["code"] == "REQUEST.IDEMPOTENCY_CONFLICT"
    wrong_class = client.post(f"/api/v1/classes/{class_id}/members/import", headers={**scoped, "Idempotency-Key": "wrong-class"}, files={"file": ("students.xlsx", workbook_bytes([["2399998", "错班学生", "其他班级", "", ""]]))})
    assert wrong_class.status_code == 201 and wrong_class.json()["success_count"] == 0
    assert wrong_class.json()["error_rows_json"][0]["reason"] == "班级与当前导入目标不一致"
    bad = [["2301001", "重复学生", "", "", ""], ["", "无学号", "", "", ""], ["=2+2", "公式风险", "", "", ""]]
    failed = client.post(f"/api/v1/classes/{class_id}/members/import", headers={**scoped, "Idempotency-Key": "bad-batch"}, files={"file": ("bad.xlsx", workbook_bytes(bad))})
    assert failed.json()["failure_count"] == 3
    errors = client.get(f"/api/v1/import-jobs/{failed.json()['job_id']}/error-rows.xlsx", headers=scoped)
    assert errors.status_code == 200 and load_workbook(BytesIO(errors.content)).active.max_row == 4
    members = client.get(f"/api/v1/classes/{class_id}/members", headers=scoped)
    assert members.json()["total"] == 43


def test_member_import_rejects_non_xlsx_and_oversized_upload(test_context):
    client, _ = test_context; course_id, class_id = build_course_class(client)
    scoped = headers("teaching.members.import", course_id=course_id, class_id=class_id)
    wrong_type = client.post(f"/api/v1/classes/{class_id}/members/import", headers={**scoped, "Idempotency-Key":"wrong-type"}, files={"file":("students.csv",b"a,b")})
    assert wrong_type.status_code == 422 and wrong_type.json()["code"] == "IMPORT.FILE_TYPE_INVALID"
    oversized = client.post(f"/api/v1/classes/{class_id}/members/import", headers={**scoped, "Idempotency-Key":"too-large"}, files={"file":("students.xlsx",b"0"*(10*1024*1024+1))})
    assert oversized.status_code == 413 and oversized.json()["code"] == "IMPORT.FILE_TOO_LARGE"


def test_import_batch_duplicate_counts_only_rows_actually_written(test_context):
    client, _ = test_context; course_id, class_id = build_course_class(client)
    scoped = headers("teaching.members.read", "teaching.members.import", course_id=course_id, class_id=class_id)
    rows = [
        ["2301001", "张三", "网络安全 2301 班", "", ""],
        ["2301001", "同批次重复", "网络安全 2301 班", "", ""],
        ["2301002", "李四", "网络安全 2301 班", "", ""],
    ]
    imported = client.post(
        f"/api/v1/classes/{class_id}/members/import",
        headers={**scoped, "Idempotency-Key": "batch-with-duplicate"},
        files={"file": ("students.xlsx", workbook_bytes(rows))},
    )
    assert imported.status_code == 201, imported.text
    assert imported.json()["success_count"] == 2
    assert imported.json()["failure_count"] == 1
    assert imported.json()["duplicate_count"] == 1
    members = client.get(f"/api/v1/classes/{class_id}/members", headers=scoped).json()
    assert members["total"] == imported.json()["success_count"] == 2


def test_roster_mutations_and_freeze_use_the_same_class_lock(test_context, monkeypatch):
    client, sessions = test_context; course_id, class_id = build_course_class(client)
    scoped = headers(
        "teaching.members.import", "teaching.members.write", "teaching.roster.freeze",
        course_id=course_id, class_id=class_id,
    )
    lock_calls = []
    original = TeachingRepository.class_for_update

    def tracked_class_for_update(repo, locked_class_id):
        lock_calls.append(locked_class_id)
        return original(repo, locked_class_id)

    monkeypatch.setattr(TeachingRepository, "class_for_update", tracked_class_for_update)
    with sessions() as db:
        seed_student_profile(db, student_number="2301002", full_name="李四")
        db.commit()
    imported = client.post(
        f"/api/v1/classes/{class_id}/members/import",
        headers={**scoped, "Idempotency-Key": "freeze-lock-regression"},
        files={"file": ("one.xlsx", workbook_bytes([["2301001", "张三", "网络安全 2301 班", "", ""]]))},
    )
    assert imported.status_code == 201, imported.text
    added = client.post(
        f"/api/v1/classes/{class_id}/members",
        headers=scoped,
        json={"student_number": "2301002", "student_name": "李四"},
    )
    assert added.status_code == 201, added.text
    removed = client.delete(
        f"/api/v1/classes/{class_id}/members/{added.json()['class_membership_id']}",
        headers=scoped,
    )
    assert removed.status_code == 200, removed.text
    frozen = client.post(f"/api/v1/classes/{class_id}/roster/freeze", headers=scoped)
    assert frozen.status_code == 200, frozen.text
    assert frozen.json()["member_count"] == 1
    assert lock_calls == [class_id, class_id, class_id, class_id]


def test_legacy_import_without_request_digest_requires_a_new_idempotency_key(test_context):
    client, sessions = test_context; course_id, class_id = build_course_class(client)
    with sessions() as db:
        db.add(ImportJob(
            job_id="legacy-import-job",
            class_id=class_id,
            idempotency_key="legacy-key",
            request_sha256=None,
            status="COMPLETED",
            success_count=1,
            failure_count=0,
            duplicate_count=0,
            error_rows_json=[],
            created_by="teacher-a",
            created_at=datetime.utcnow(),
        ))
        db.commit()
    scoped = headers("teaching.members.import", course_id=course_id, class_id=class_id)
    replay = client.post(
        f"/api/v1/classes/{class_id}/members/import",
        headers={**scoped, "Idempotency-Key": "legacy-key"},
        files={"file": ("different.xlsx", workbook_bytes([["2301001", "张三", "网络安全 2301 班", "", ""]]))},
    )
    assert replay.status_code == 409
    assert replay.json()["code"] == "REQUEST.IDEMPOTENCY_DIGEST_MISSING"
    assert replay.json()["message"] == "历史导入任务缺少文件摘要，不能安全复用该幂等标识，请使用新的 Idempotency-Key"
    assert replay.json()["details"] == {"job_id": "legacy-import-job"}


def test_g2_attendance_real_sign_duplicate_close_summary_and_scope(test_context):
    client, sessions = test_context; course_id, class_id = build_course_class(client)
    teacher = headers("teaching.members.import", "teaching.members.read", "teaching.attendance.write", "teaching.attendance.read", course_id=course_id, class_id=class_id)
    rows = [["2301001", "张三", "网络安全 2301 班", "", ""]]
    client.post(f"/api/v1/classes/{class_id}/members/import", headers={**teacher, "Idempotency-Key": "one"}, files={"file": ("one.xlsx", workbook_bytes(rows))})
    with sessions() as db: student_id = db.scalar(select(ClassMembership.student_id))
    start, end = datetime.utcnow() - timedelta(minutes=1), datetime.utcnow() + timedelta(minutes=10)
    task = client.post("/api/v1/attendance/tasks", headers=teacher, json={"course_id": course_id, "class_id": class_id, "task_type": "CLASSROOM", "title": "RSA 数字签名课堂签到", "starts_at": start.isoformat(), "expires_at": end.isoformat()})
    assert task.status_code == 201
    published = client.post(f"/api/v1/attendance/tasks/{task.json()['task_id']}/publish", headers=teacher).json()
    student = headers("teaching.attendance.sign", student_id=student_id, teacher_id="")
    link = client.get(f"/api/v1/attendance/sign-links/{published['sign_token']}", headers=student)
    assert link.status_code == 200 and link.json()["task_id"] == task.json()["task_id"]
    signed = client.post(f"/api/v1/attendance/sign-links/{published['sign_token']}/sign", headers=student)
    assert signed.status_code == 200
    with sessions() as db:
        event = db.scalar(select(DomainEventOutbox).where(DomainEventOutbox.event_type == "attendance.completed"))
        assert event.payload_json["raw_score"] == event.payload_json["max_score"] == 1
    duplicate = client.post(f"/api/v1/attendance/{task.json()['task_id']}/sign", headers=student, params={"token": published["sign_token"]})
    assert duplicate.json()["record_id"] == signed.json()["record_id"]
    summary = client.get("/api/v1/attendance/section-summary", headers=teacher).json()["items"][0]
    assert summary | {"expected": 1, "present": 1, "absent": 0} == summary
    client.post(f"/api/v1/attendance/tasks/{task.json()['task_id']}/close", headers=teacher)
    closed = client.post(f"/api/v1/attendance/{task.json()['task_id']}/sign", headers=headers("teaching.attendance.sign", student_id="other", teacher_id="", course_id=course_id, class_id=class_id), params={"token": published["sign_token"]})
    assert closed.status_code in (403, 409)
    outsider = headers("teaching.attendance.read", class_id="class-b", teacher_id="teacher-b")
    assert client.get(f"/api/v1/attendance/tasks/{task.json()['task_id']}", headers=outsider).status_code == 403


def test_removed_member_cannot_use_attendance_poll_assignment_or_quiz(test_context):
    client, sessions = test_context; course_id, class_id = build_course_class(client)
    teacher = headers(
        "teaching.members.import", "teaching.members.write", "teaching.attendance.write",
        "teaching.poll.write", "teaching.assignment.write", "teaching.quiz.write",
        course_id=course_id, class_id=class_id,
    )
    client.post(
        f"/api/v1/classes/{class_id}/members/import",
        headers={**teacher, "Idempotency-Key": "removed-member"},
        files={"file": ("one.xlsx", workbook_bytes([["2301001", "张三", "", "", ""]]))},
    )
    with sessions() as db:
        membership = db.scalar(select(ClassMembership))
        student_id, membership_id = membership.student_id, membership.class_membership_id

    start, end = datetime.utcnow() - timedelta(minutes=1), datetime.utcnow() + timedelta(minutes=30)
    task = client.post("/api/v1/attendance/tasks", headers=teacher, json={"course_id": course_id, "class_id": class_id, "task_type": "CLASSROOM", "title": "成员状态校验", "starts_at": start.isoformat(), "expires_at": end.isoformat()}).json()
    token = client.post(f"/api/v1/attendance/tasks/{task['task_id']}/publish", headers=teacher).json()["sign_token"]
    poll = client.post("/api/v1/polls", headers=teacher, json={"course_id": course_id, "class_id": class_id, "poll_type": "UNDERSTANDING", "title": "成员状态校验", "options": ["掌握", "未掌握"]}).json()
    client.post(f"/api/v1/polls/{poll['poll_id']}/publish", headers=teacher)
    due = (datetime.utcnow() + timedelta(hours=1)).isoformat()
    question = {"question_id": seed_published_question(sessions, course_id)}
    assignment = client.post("/api/v1/assignments", headers=teacher, json={"course_id": course_id, "class_id": class_id, "title": "成员状态校验", "due_at": due, "questions": [question]}).json()
    client.post(f"/api/v1/assignments/{assignment['assignment_id']}/publish", headers=teacher)
    quiz = client.post("/api/v1/quizzes", headers=teacher, json={"course_id": course_id, "class_id": class_id, "title": "成员状态校验", "time_limit_minutes": 30, "questions": [question]}).json()
    client.post(f"/api/v1/quizzes/{quiz['quiz_id']}/publish", headers=teacher)

    student = headers("teaching.attendance.sign", "teaching.poll.answer", "teaching.assignment.submit", "teaching.quiz.submit", "teaching.student.read", student_id=student_id, teacher_id="")
    attempt = client.post(f"/api/v1/quizzes/{quiz['quiz_id']}/attempts", headers=student).json()
    removed = client.delete(f"/api/v1/classes/{class_id}/members/{membership_id}", headers=teacher)
    assert removed.status_code == 200

    requests = [
        client.get(f"/api/v1/attendance/sign-links/{token}", headers=student),
        client.post(f"/api/v1/attendance/sign-links/{token}/sign", headers=student),
        client.post(f"/api/v1/attendance/{task['task_id']}/sign", headers=student, params={"token": token}),
        client.post(f"/api/v1/polls/{poll['poll_id']}/answers", headers=student, json={"option_id": poll["options"][0]["option_id"]}),
        client.post(f"/api/v1/assignments/{assignment['assignment_id']}/submit", headers=student, json={"answers": {}}),
        client.post(f"/api/v1/quizzes/{quiz['quiz_id']}/attempts", headers=student),
        client.post(f"/api/v1/quizzes/{quiz['quiz_id']}/attempts/{attempt['attempt_id']}/submit", headers=student, json={"answers": {}}),
    ]
    assert all(response.status_code == 403 and response.json()["code"] == "AUTH.SCOPE_DENIED" for response in requests)
    read_model = client.get("/api/v1/teaching/student-read-model", headers=student)
    assert read_model.status_code == 200 and read_model.json() == {"class_count": 0, "open_attendance": []}


@pytest.mark.parametrize("poll_type", ["UNDERSTANDING", "ASSIGNMENT_COMPLETION", "TEACHING_FEEDBACK"])
def test_poll_modes_are_server_aggregated(test_context, poll_type):
    client, sessions = test_context; course_id, class_id = build_course_class(client)
    teacher = headers("teaching.members.import", "teaching.poll.write", "teaching.poll.read", course_id=course_id, class_id=class_id)
    client.post(f"/api/v1/classes/{class_id}/members/import", headers={**teacher, "Idempotency-Key": poll_type}, files={"file": ("one.xlsx", workbook_bytes([["2301001", "张三", "", "", ""]]))})
    with sessions() as db: student_id = db.scalar(select(ClassMembership.student_id))
    poll = client.post("/api/v1/polls", headers=teacher, json={"course_id": course_id, "class_id": class_id, "poll_type": poll_type, "title": "课堂反馈", "options": ["掌握", "需复习"]}).json()
    client.post(f"/api/v1/polls/{poll['poll_id']}/publish", headers=teacher)
    student = headers("teaching.poll.answer", student_id=student_id, teacher_id="", course_id=course_id, class_id=class_id)
    assert client.post(f"/api/v1/polls/{poll['poll_id']}/answers", headers=student, json={"option_id": poll["options"][0]["option_id"]}).status_code == 201
    with sessions() as db:
        event = db.scalar(select(DomainEventOutbox).where(DomainEventOutbox.event_type == "poll.completed"))
        assert event.payload_json["raw_score"] == event.payload_json["max_score"] == 1
    result = client.get(f"/api/v1/polls/{poll['poll_id']}/results", headers=teacher).json()
    assert result["poll_type"] == poll_type and result["items"][0]["count"] == 1


def test_assignment_quiz_events_and_cross_student_isolation(test_context):
    client, sessions = test_context; course_id, class_id = build_course_class(client)
    teacher = headers("teaching.members.import", "teaching.assignment.write", "teaching.quiz.write", course_id=course_id, class_id=class_id)
    client.post(f"/api/v1/classes/{class_id}/members/import", headers={**teacher, "Idempotency-Key": "students"}, files={"file": ("two.xlsx", workbook_bytes([["1", "甲", "", "", ""], ["2", "乙", "", "", ""]]))})
    with sessions() as db: students = list(db.scalars(select(ClassMembership.student_id).order_by(ClassMembership.student_number)))
    due = (datetime.utcnow() + timedelta(hours=1)).isoformat()
    question_id = seed_published_question(sessions, course_id)
    question = {"question_id": question_id}
    assignment = client.post("/api/v1/assignments", headers=teacher, json={"course_id": course_id, "class_id": class_id, "title": "RSA 作业", "due_at": due, "questions": [question]}).json()
    client.post(f"/api/v1/assignments/{assignment['assignment_id']}/publish", headers=teacher)
    student = headers("teaching.assignment.submit", "teaching.quiz.submit", student_id=students[0], teacher_id="", course_id=course_id, class_id=class_id)
    assignment_submission = client.post(f"/api/v1/assignments/{assignment['assignment_id']}/submit", headers=student, json={"answers": {question_id: "A"}})
    assert assignment_submission.status_code == 201
    assert assignment_submission.json()["raw_score"] == assignment_submission.json()["max_score"] == 10
    quiz = client.post("/api/v1/quizzes", headers=teacher, json={"course_id": course_id, "class_id": class_id, "title": "课堂小测", "time_limit_minutes": 10, "questions": [question]}).json()
    client.post(f"/api/v1/quizzes/{quiz['quiz_id']}/publish", headers=teacher)
    attempt = client.post(f"/api/v1/quizzes/{quiz['quiz_id']}/attempts", headers=student).json()
    other = headers("teaching.quiz.submit", student_id=students[1], teacher_id="", course_id=course_id, class_id=class_id)
    denied = client.post(f"/api/v1/quizzes/{quiz['quiz_id']}/attempts/{attempt['attempt_id']}/submit", headers=other, json={"answers": {}})
    assert denied.status_code == 403
    quiz_submission = client.post(f"/api/v1/quizzes/{quiz['quiz_id']}/attempts/{attempt['attempt_id']}/submit", headers=student, json={"answers": {question_id: "A"}})
    assert quiz_submission.status_code == 200
    assert quiz_submission.json()["raw_score"] == quiz_submission.json()["max_score"] == 10
    with sessions() as db:
        event_types = set(db.scalars(select(DomainEventOutbox.event_type)))
        proofs = list(db.scalars(select(TeachingScoreProof).order_by(TeachingScoreProof.source_type)))
        assert [proof.source_type for proof in proofs] == ["assignment_submission", "quiz_attempt"]
        for proof in proofs:
            verify_teaching_score_proof(db, proof.proof_id)
            proof_event = db.get(DomainEventOutbox, proof.score_proof_event_id)
            score_event = db.get(DomainEventOutbox, proof.score_event_id)
            assert proof_event.occurred_at < score_event.occurred_at
            assert score_event.payload_json["score_proof_event_id"] == proof_event.event_id
    assert {"assignment.submitted", "quiz.completed", "grading.score.proof.frozen", "teaching.audit"} <= event_types


def test_student_task_reads_are_scoped_and_do_not_expose_frozen_answer_evidence(test_context):
    client, sessions = test_context
    course_id, class_id = build_course_class(client)
    teacher = headers("teaching.members.import", "teaching.assignment.write", "teaching.quiz.write", course_id=course_id, class_id=class_id)
    client.post(
        f"/api/v1/classes/{class_id}/members/import",
        headers={**teacher, "Idempotency-Key": "student-task-read"},
        files={"file": ("one.xlsx", workbook_bytes([["2301001", "张三", "", "", ""]]))},
    )
    with sessions() as db:
        student_id = db.scalar(select(ClassMembership.student_id))
    question_id = seed_published_question(sessions, course_id)
    due_at = (datetime.utcnow() + timedelta(hours=1)).isoformat()
    assignment = client.post(
        "/api/v1/assignments",
        headers=teacher,
        json={"course_id": course_id, "class_id": class_id, "title": "真实作业", "due_at": due_at, "questions": [{"question_id": question_id}]},
    ).json()
    quiz = client.post(
        "/api/v1/quizzes",
        headers=teacher,
        json={"course_id": course_id, "class_id": class_id, "title": "真实测验", "time_limit_minutes": 10, "questions": [{"question_id": question_id}]},
    ).json()
    assert client.post(f"/api/v1/assignments/{assignment['assignment_id']}/publish", headers=teacher).status_code == 200
    assert client.post(f"/api/v1/quizzes/{quiz['quiz_id']}/publish", headers=teacher).status_code == 200
    student = headers("teaching.assignment.submit", "teaching.quiz.submit", student_id=student_id, teacher_id="", course_id=course_id, class_id=class_id)

    assignment_list = client.get("/api/v1/assignments/my", headers=student)
    assert assignment_list.status_code == 200
    assert [item["assignment_id"] for item in assignment_list.json()["items"]] == [assignment["assignment_id"]]
    assignment_task = client.get(f"/api/v1/assignments/{assignment['assignment_id']}/student-task", headers=student)
    assert assignment_task.status_code == 200
    assignment_question = assignment_task.json()["questions"][0]
    assert assignment_question["question_id"] == question_id
    assert assignment_question["options"] == [{"key": "A", "text": "公钥"}, {"key": "B", "text": "私钥"}]
    assert {"answer", "question_snapshot", "question_version", "max_score"}.isdisjoint(assignment_question)

    quiz_list = client.get("/api/v1/quizzes/my", headers=student)
    assert quiz_list.status_code == 200
    assert [item["quiz_id"] for item in quiz_list.json()["items"]] == [quiz["quiz_id"]]
    quiz_task = client.get(f"/api/v1/quizzes/{quiz['quiz_id']}/student-task", headers=student)
    assert quiz_task.status_code == 200
    quiz_question = quiz_task.json()["questions"][0]
    assert quiz_question["question_ref_id"] != question_id
    assert {"answer", "question_snapshot", "question_version", "max_score"}.isdisjoint(quiz_question)

    # Answers are keyed by the server-issued frozen reference, never a client
    # snapshot/version.  The original question ID remains an accepted legacy
    # alias on the write endpoint, but the browser uses the safer reference.
    submitted = client.post(
        f"/api/v1/assignments/{assignment['assignment_id']}/submit",
        headers=student,
        json={"answers": {assignment_question["question_ref_id"]: "A"}},
    )
    assert submitted.status_code == 201
    assert client.get("/api/v1/assignments/my", headers=student).json()["items"][0]["submission_status"] == "SUBMITTED"

    intruder = headers("teaching.assignment.submit", "teaching.quiz.submit", student_id="not-in-class", teacher_id="", course_id=course_id, class_id=class_id)
    assert client.get(f"/api/v1/assignments/{assignment['assignment_id']}/student-task", headers=intruder).status_code == 403
    assert client.get(f"/api/v1/quizzes/{quiz['quiz_id']}/student-task", headers=intruder).status_code == 403


def test_server_scored_submission_freezes_question_and_emits_paired_proof_events(test_context):
    client, sessions = test_context
    course_id, class_id = build_course_class(client)
    teacher = headers("teaching.members.import", "teaching.assignment.write", course_id=course_id, class_id=class_id)
    client.post(
        f"/api/v1/classes/{class_id}/members/import",
        headers={**teacher, "Idempotency-Key": "score-proof-student"},
        files={"file": ("one.xlsx", workbook_bytes([["2301001", "张三", "", "", ""]]))},
    )
    with sessions() as db:
        student_id = db.scalar(select(ClassMembership.student_id))
    question_id = seed_published_question(sessions, course_id)
    due_at = (datetime.utcnow() + timedelta(hours=1)).isoformat()

    unsupported_answer_question = seed_published_question(
        sessions,
        course_id,
        question_id="question-b-invalid-answer",
        answer=[{"not": "a scoreable answer"}],
    )
    invalid_source = client.post(
        "/api/v1/assignments",
        headers=teacher,
        json={
            "course_id": course_id,
            "class_id": class_id,
            "title": "不可判分的题库事实",
            "due_at": due_at,
            "questions": [{"question_id": unsupported_answer_question}],
        },
    )
    assert invalid_source.status_code == 422
    assert invalid_source.json()["code"] == "QUESTION.FROZEN_ANSWER_INVALID"

    forged_task = client.post(
        "/api/v1/assignments",
        headers=teacher,
        json={
            "course_id": course_id,
            "class_id": class_id,
            "title": "伪造快照作业",
            "due_at": due_at,
            "questions": [{"question_id": question_id, "question_version": "browser-v1", "question_snapshot": {"answer": ["B"]}, "max_score": 999}],
        },
    )
    assert forged_task.status_code == 422

    assignment = client.post(
        "/api/v1/assignments",
        headers=teacher,
        json={"course_id": course_id, "class_id": class_id, "title": "服务端评分作业", "due_at": due_at, "questions": [{"question_id": question_id}]},
    )
    assert assignment.status_code == 201, assignment.text
    assignment_id = assignment.json()["assignment_id"]
    with sessions() as db:
        frozen_ref = db.scalar(select(AssignmentQuestionRef).where(AssignmentQuestionRef.assignment_id == assignment_id))
        original_snapshot = dict(frozen_ref.question_snapshot)
        assert frozen_ref.question_id == question_id
        assert len(frozen_ref.question_version) == 64
        assert frozen_ref.question_snapshot["contract"] == "teaching-frozen-question/v1"
        assert frozen_ref.question_snapshot["max_score"] == 10
    assert client.post(f"/api/v1/assignments/{assignment_id}/publish", headers=teacher).status_code == 200
    student = headers("teaching.assignment.submit", student_id=student_id, teacher_id="", course_id=course_id, class_id=class_id)

    forged_submission = client.post(
        f"/api/v1/assignments/{assignment_id}/submit",
        headers=student,
        json={"answers": {question_id: "A"}, "raw_score": 100, "max_score": 100, "source_proof": {"issuer": "browser"}},
    )
    assert forged_submission.status_code == 422
    with sessions() as db:
        assert not list(db.scalars(select(TeachingScoreProof)))

        # The submission path must also distrust a database-side edit of the
        # frozen answer/snapshot when its version hash no longer matches.
        frozen_ref = db.scalar(select(AssignmentQuestionRef).where(AssignmentQuestionRef.assignment_id == assignment_id))
        frozen_ref.question_snapshot = {**frozen_ref.question_snapshot, "answer": ["B"]}
        db.commit()

    tampered_snapshot = client.post(
        f"/api/v1/assignments/{assignment_id}/submit",
        headers=student,
        json={"answers": {question_id: "A"}},
    )
    assert tampered_snapshot.status_code == 409
    assert tampered_snapshot.json()["code"] == "TEACHING.FROZEN_QUESTION_INVALID"
    with sessions() as db:
        frozen_ref = db.scalar(select(AssignmentQuestionRef).where(AssignmentQuestionRef.assignment_id == assignment_id))
        frozen_ref.question_snapshot = original_snapshot
        db.commit()

    submitted = client.post(f"/api/v1/assignments/{assignment_id}/submit", headers=student, json={"answers": {question_id: "A"}})
    assert submitted.status_code == 201, submitted.text
    assert submitted.json()["raw_score"] == submitted.json()["max_score"] == 10
    with sessions() as db:
        proof_row = db.scalar(select(TeachingScoreProof))
        verify_teaching_score_proof(db, proof_row.proof_id)
        proof_event = db.get(DomainEventOutbox, proof_row.score_proof_event_id)
        score_event = db.get(DomainEventOutbox, proof_row.score_event_id)
        assert proof_event.event_type == "grading.score.proof.frozen"
        assert proof_event.actor_user_id == "service_teaching_score_prover"
        assert (proof_event.aggregate_type, proof_event.aggregate_id) == ("assignment_submission", submitted.json()["submission_id"])
        assert score_event.event_type == "assignment.submitted"
        assert proof_event.occurred_at + timedelta(seconds=1) == score_event.occurred_at
        assert score_event.payload_json["score_proof_event_id"] == proof_event.event_id
        assert "source_proof" not in score_event.payload_json
        proof = proof_event.payload_json["source_proof"]
        assert proof["contract"] == "grading-score-proof/v1"
        assert proof["issuer"] == "teaching-core"
        assert proof["origin"] == "SERVER_GRADED"
        assert proof["evidence_type"] == "ASSIGNMENT_FROZEN_QUESTION_SET"
        assert proof["source_event_id"] == score_event.event_id
        assert proof["source_fact_id"] == submitted.json()["submission_id"]
        assert proof["frozen_question_sha256"] == sha256_json(proof_row.frozen_question_evidence_json)
        assert proof["answer_evidence_sha256"] == sha256_json(proof_row.answer_evidence_json)
        assert proof["scoring_evidence_sha256"] == sha256_json(proof_row.scoring_evidence_json)
        assert proof["score_payload_sha256"] == score_payload_sha256(
            event_id=score_event.event_id,
            event_type="assignment.submitted",
            aggregate_id=assignment_id,
            payload=score_event.payload_json,
            proof=proof,
            raw=Decimal(str(score_event.payload_json["raw_score"])),
            maximum=Decimal(str(score_event.payload_json["max_score"])),
        )

        source_submission = db.get(AssignmentSubmission, submitted.json()["submission_id"])
        original_answers = dict(source_submission.answers_json)
        source_submission.answers_json = {proof_row.answer_evidence_json[0]["question_ref_id"]: "B"}
        db.commit()
        with pytest.raises(ApiError) as tampered_saved_answer:
            verify_teaching_score_proof(db, proof_row.proof_id)
        assert tampered_saved_answer.value.code == "TEACHING.SCORE_PROOF_TAMPERED"

        source_submission.answers_json = original_answers
        db.commit()

        original_hash = proof_row.answer_evidence_sha256
        proof_row.answer_evidence_sha256 = "0" * 64
        db.commit()
        with pytest.raises(ApiError) as tampered_proof:
            verify_teaching_score_proof(db, proof_row.proof_id)
        assert tampered_proof.value.code == "TEACHING.SCORE_PROOF_TAMPERED"

        proof_row.answer_evidence_sha256 = original_hash
        original_proof_payload = proof_event.payload_json
        proof_event.payload_json = {**proof_event.payload_json, "course_id": "tampered-course"}
        db.commit()
        with pytest.raises(ApiError) as tampered_proof_event:
            verify_teaching_score_proof(db, proof_row.proof_id)
        assert tampered_proof_event.value.code == "TEACHING.SCORE_PROOF_TAMPERED"

        proof_event.payload_json = original_proof_payload
        score_event.payload_json = {**score_event.payload_json, "raw_score": 0.0}
        db.commit()
        with pytest.raises(ApiError) as tampered_outbox:
            verify_teaching_score_proof(db, proof_row.proof_id)
        assert tampered_outbox.value.code == "TEACHING.SCORE_PROOF_TAMPERED"

        # A retry must not turn a legacy/no-proof submission back into an
        # accepted client-visible score.  This is the fail-closed migration
        # boundary for old insecure submission rows.
        db.delete(proof_row)
        db.commit()

    legacy_replay = client.post(
        f"/api/v1/assignments/{assignment_id}/submit",
        headers=student,
        json={"answers": {question_id: "A"}},
    )
    assert legacy_replay.status_code == 409
    assert legacy_replay.json()["code"] == "TEACHING.LEGACY_SCORE_UNVERIFIED"


def test_teacher_cannot_access_another_class_and_student_cannot_publish(test_context):
    client, _ = test_context; course_id, class_id = build_course_class(client)
    teacher_b = headers("teaching.members.read", class_id="another-class", teacher_id="teacher-b")
    assert client.get(f"/api/v1/classes/{class_id}/members", headers=teacher_b).status_code == 403
    student = headers("teaching.attendance.sign", student_id="student-a", teacher_id="", class_id=class_id, course_id=course_id)
    body = {"course_id": course_id, "class_id": class_id, "task_type": "CLASSROOM", "title": "越权", "starts_at": datetime.utcnow().isoformat(), "expires_at": (datetime.utcnow() + timedelta(minutes=5)).isoformat()}
    assert client.post("/api/v1/attendance/tasks", headers=student, json=body).status_code == 403


def test_teacher_student_management_search_detail_summary_remove_and_permissions(test_context):
    client, sessions = test_context; course_id, class_id = build_course_class(client)
    teacher = headers("teaching.members.read", "teaching.members.write", course_id=course_id, class_id=class_id)
    with sessions() as db:
        seed_student_profile(db, student_id="existing-student-1", student_number="2301002", full_name="李四", email="li@example.edu.cn")
        seed_student_profile(db, student_id="existing-student-2", student_number="2301001", full_name="张三")
        db.commit()
    first = client.post(f"/api/v1/classes/{class_id}/members", headers=teacher, json={"student_number":"2301002","student_name":"李四"})
    assert first.status_code == 201, first.text
    second = client.post(f"/api/v1/classes/{class_id}/members", headers=teacher, json={"student_number":"2301001","student_name":"张三"})
    assert second.status_code == 201
    listing = client.get(f"/api/v1/classes/{class_id}/members", headers=teacher, params={"search":"张","sort":"student_name","direction":"desc","page":1,"page_size":10}).json()
    assert listing["total"] == 1 and listing["items"][0]["student_name"] == "张三"
    membership_id = first.json()["class_membership_id"]
    assert client.get(f"/api/v1/classes/{class_id}/members/{membership_id}", headers=teacher).json()["email"] == "li@example.edu.cn"
    summary = client.get(f"/api/v1/classes/{class_id}/members/{membership_id}/learning-summary", headers=teacher).json()
    assert summary["experiment"]["label"] == "数据待汇总" and summary["grade"]["label"] == "数据待汇总"
    student_summary = client.get(f"/api/v1/classes/{class_id}/students/existing-student-1/learning-summary", headers=teacher)
    assert student_summary.status_code == 200 and student_summary.json()["student_id"] == "existing-student-1"
    exported = client.get(f"/api/v1/classes/{class_id}/members/export.xlsx", headers=teacher)
    assert exported.status_code == 200 and load_workbook(BytesIO(exported.content)).active.max_row == 3
    student = headers("teaching.members.read", student_id="existing-student-1", teacher_id="", class_id=class_id, course_id=course_id)
    assert client.delete(f"/api/v1/classes/{class_id}/members/{membership_id}", headers=student).status_code == 403
    other_teacher = headers("teaching.members.read", "teaching.members.write", teacher_id="teacher-b", class_id="class-b")
    assert client.delete(f"/api/v1/classes/{class_id}/members/{membership_id}", headers=other_teacher).status_code == 403
    removed = client.delete(f"/api/v1/classes/{class_id}/members/{membership_id}", headers=teacher)
    assert removed.status_code == 200 and removed.json()["status"] == "REMOVED"
    assert client.get(f"/api/v1/classes/{class_id}/members", headers=teacher).json()["total"] == 1
