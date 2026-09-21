from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from app.common.models import Base, DomainEventOutbox
from app.database import get_session
from app.main import app
from app.resources.catalog import COURSE_ID, lesson_rows
from app.resources.models import LessonResource, Question, QuestionImportJob, QuestionImportRow, QuestionReview


AUTHOR_HEADERS = {
    "X-User-Id": "question_author",
    "X-Role": "teacher",
    "X-Teacher-Id": "question_author",
    "X-Course-Ids": COURSE_ID,
    "X-Permissions": "resources:read,resources:write,resources:review",
}
REVIEWER_HEADERS = {**AUTHOR_HEADERS, "X-User-Id": "question_reviewer", "X-Teacher-Id": "question_reviewer"}
READ_ONLY_HEADERS = {**AUTHOR_HEADERS, "X-Permissions": "resources:read"}
OUT_OF_SCOPE_HEADERS = {**AUTHOR_HEADERS, "X-Course-Ids": "course_other"}


@pytest.fixture()
def question_client():
    engine = create_engine(
        "sqlite+pysqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    sessions = sessionmaker(bind=engine, expire_on_commit=False)
    with sessions() as session:
        session.add_all(LessonResource(**row) for row in lesson_rows())
        session.commit()

    def override_session():
        with sessions() as session:
            yield session

    app.dependency_overrides[get_session] = override_session
    try:
        yield TestClient(app), sessions
    finally:
        app.dependency_overrides.pop(get_session, None)
        Base.metadata.drop_all(engine)
        engine.dispose()


def complete_template(content: bytes) -> bytes:
    workbook = load_workbook(BytesIO(content))
    sheet = workbook["题目导入"]
    for row_number in range(2, 198):
        question_type = sheet.cell(row_number, 3).value
        sheet.cell(row_number, 4).value = f"导入题 {row_number - 1}"
        sheet.cell(row_number, 10).value = "导入题解析"
        if question_type == "填空":
            sheet.cell(row_number, 9).value = "答案"
        elif question_type == "单选":
            sheet.cell(row_number, 5).value = "正确"
            sheet.cell(row_number, 6).value = "错误"
            sheet.cell(row_number, 9).value = "A"
        elif question_type == "多选":
            sheet.cell(row_number, 5).value = "正确一"
            sheet.cell(row_number, 6).value = "正确二"
            sheet.cell(row_number, 7).value = "错误"
            sheet.cell(row_number, 9).value = "A,B"
        else:
            sheet.cell(row_number, 9).value = "A"
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def post_import(client: TestClient, content: bytes, key: str):
    return client.post(
        "/api/v1/questions/import",
        headers={**AUTHOR_HEADERS, "Idempotency-Key": key},
        data={"course_id": COURSE_ID},
        files={"file": ("questions.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )


def test_import_is_all_or_nothing_idempotent_and_uses_independent_review(question_client):
    client, sessions = question_client
    template = client.get("/api/v1/questions/import-template.xlsx", headers=AUTHOR_HEADERS, params={"course_id": COURSE_ID})
    assert template.status_code == 200
    valid_content = complete_template(template.content)

    invalid_workbook = load_workbook(BytesIO(valid_content))
    invalid_workbook["题目导入"]["D2"] = "@unsafe"
    invalid_stream = BytesIO()
    invalid_workbook.save(invalid_stream)
    invalid = post_import(client, invalid_stream.getvalue(), "invalid-import")
    assert invalid.status_code == 201
    assert invalid.json()["status"] == "VALIDATION_FAILED"
    assert invalid.json()["imported_count"] == 0
    assert any(error["row_number"] == 2 and error["code"] == "QUESTION_IMPORT.DANGEROUS_CELL" for error in invalid.json()["error_rows"])
    with sessions() as session:
        assert session.scalar(select(func.count()).select_from(Question)) == 0
        assert session.scalar(select(func.count()).select_from(QuestionImportRow)) == 196

    completed = post_import(client, valid_content, "valid-import")
    assert completed.status_code == 201
    result = completed.json()
    assert result["status"] == "COMPLETED"
    assert result["total_count"] == result["imported_count"] == result["review_queue_count"] == 196
    replay = post_import(client, valid_content, "valid-import")
    assert replay.status_code == 201 and replay.json()["job_id"] == result["job_id"]

    changed_workbook = load_workbook(BytesIO(valid_content))
    changed_workbook["题目导入"]["J2"] = "改动后的解析"
    changed_stream = BytesIO()
    changed_workbook.save(changed_stream)
    conflict = post_import(client, changed_stream.getvalue(), "valid-import")
    assert conflict.status_code == 409 and conflict.json()["code"] == "REQUEST.IDEMPOTENCY_CONFLICT"

    queue = client.get("/api/v1/questions/review-queue", headers=REVIEWER_HEADERS, params={"course_id": COURSE_ID, "page_size": 200})
    assert queue.status_code == 200 and queue.json()["total"] == 196
    first = queue.json()["items"][0]
    assert first["can_review"] is True and first["answer"] and first["explanation"]
    own_queue = client.get("/api/v1/questions/review-queue", headers=AUTHOR_HEADERS, params={"course_id": COURSE_ID, "page_size": 200})
    assert own_queue.status_code == 200 and own_queue.json()["items"][0]["can_review"] is False
    assert client.post(f"/api/v1/questions/{first['question_id']}/review", headers=AUTHOR_HEADERS).status_code == 409
    approved = client.post(f"/api/v1/questions/{first['question_id']}/review", headers=REVIEWER_HEADERS)
    assert approved.status_code == 200 and approved.json()["status"] == "PUBLISHED"

    second = queue.json()["items"][1]
    missing_comment = client.post(f"/api/v1/questions/{second['question_id']}/review", headers=REVIEWER_HEADERS, json={"decision": "REJECTED"})
    assert missing_comment.status_code == 422
    rejected = client.post(
        f"/api/v1/questions/{second['question_id']}/review",
        headers=REVIEWER_HEADERS,
        json={"decision": "REJECTED", "comment": "请修正题干"},
    )
    assert rejected.status_code == 200 and rejected.json()["status"] == "REJECTED"

    with sessions() as session:
        assert session.scalar(select(func.count()).select_from(Question)) == 196
        assert session.scalar(select(func.count()).select_from(QuestionImportJob)) == 2
        assert session.scalar(select(func.count()).select_from(QuestionReview)) == 2
        assert session.scalar(select(func.count()).select_from(DomainEventOutbox)) >= 200


def test_question_import_and_review_queue_enforce_permission_and_course_scope(question_client):
    client, _ = question_client
    denied_template = client.get("/api/v1/questions/import-template.xlsx", headers=READ_ONLY_HEADERS, params={"course_id": COURSE_ID})
    assert denied_template.status_code == 403 and denied_template.json()["code"] == "AUTH.PERMISSION_DENIED"

    denied_scope = client.get("/api/v1/questions/import-template.xlsx", headers=OUT_OF_SCOPE_HEADERS, params={"course_id": COURSE_ID})
    assert denied_scope.status_code == 403 and denied_scope.json()["code"] == "RESOURCE.COURSE_SCOPE_DENIED"

    template = client.get("/api/v1/questions/import-template.xlsx", headers=AUTHOR_HEADERS, params={"course_id": COURSE_ID})
    missing_key = client.post(
        "/api/v1/questions/import",
        headers=AUTHOR_HEADERS,
        data={"course_id": COURSE_ID},
        files={"file": ("questions.xlsx", template.content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert missing_key.status_code == 422
    assert missing_key.json()["code"] == "REQUEST.VALIDATION_FAILED"

    oversized_key = client.post(
        "/api/v1/questions/import",
        headers={**AUTHOR_HEADERS, "Idempotency-Key": "x" * 129},
        data={"course_id": COURSE_ID},
        files={"file": ("questions.xlsx", template.content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert oversized_key.status_code == 422

    oversized_file = client.post(
        "/api/v1/questions/import",
        headers={**AUTHOR_HEADERS, "Idempotency-Key": "oversized-file"},
        data={"course_id": COURSE_ID},
        files={"file": ("questions.xlsx", b"PK" + b"0" * (10 * 1024 * 1024), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert oversized_file.status_code == 413 and oversized_file.json()["code"] == "QUESTION_IMPORT.FILE_TOO_LARGE"

    denied_queue = client.get("/api/v1/questions/review-queue", headers=READ_ONLY_HEADERS, params={"course_id": COURSE_ID})
    assert denied_queue.status_code == 403 and denied_queue.json()["code"] == "AUTH.PERMISSION_DENIED"
