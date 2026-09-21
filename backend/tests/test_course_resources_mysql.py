import os
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from io import BytesIO
from threading import Barrier
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine, delete, select
from sqlalchemy.orm import Session

from app.common.models import DomainEventOutbox, FileObject
from app.common.context import UserContext
from app.common.errors import ApiError
from app.main import app
from app.resources.catalog import COURSE_ID, lesson_rows
from app.resources.models import LessonResource, PptAsset, Question, QuestionBank, QuestionExplanation, QuestionImportJob, QuestionImportRow, QuestionLessonMap, QuestionOption, QuestionReview, Resource, ResourceDeliveryManifest, ResourceVersion, VideoAsset
from app.resources.repository import ResourceRepository
from app.resources.service import ResourceService
from app.resources.schemas import QuestionReviewDecision

pytestmark = pytest.mark.skipif(not os.getenv("YUEKE_DATABASE_URL"), reason="需要专属 MySQL 集成库")

HEADERS = {
    "X-User-Id": "teacher_b",
    "X-Role": "teacher",
    "X-Teacher-Id": "teacher_b",
    "X-Course-Ids": COURSE_ID,
    "X-Permissions": "resources:read,resources:write,resources:review,resources:freeze",
}
REVIEW_HEADERS = {**HEADERS, "X-User-Id": "reviewer_b", "X-Teacher-Id": "reviewer_b"}


@pytest.fixture()
def client():
    return TestClient(app)


def complete_question_template(content: bytes) -> bytes:
    workbook = load_workbook(BytesIO(content))
    sheet = workbook["题目导入"]
    for row_number in range(2, 198):
        question_type = sheet.cell(row_number, 3).value
        sheet.cell(row_number, 4).value = f"MySQL 导入验证题 {row_number - 1}"
        sheet.cell(row_number, 10).value = "MySQL 真实链路解析"
        if question_type == "填空":
            sheet.cell(row_number, 9).value = "参考答案"
        elif question_type == "单选":
            sheet.cell(row_number, 5).value = "正确选项"
            sheet.cell(row_number, 6).value = "干扰选项"
            sheet.cell(row_number, 9).value = "A"
        elif question_type == "多选":
            sheet.cell(row_number, 5).value = "正确选项一"
            sheet.cell(row_number, 6).value = "正确选项二"
            sheet.cell(row_number, 7).value = "干扰选项"
            sheet.cell(row_number, 9).value = "A,B"
        else:
            sheet.cell(row_number, 9).value = "A"
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def test_mysql_catalog_audit_manifest_and_freeze_blocker(client):
    theory = client.get("/api/v1/resources/theory-lessons", headers=HEADERS).json()
    labs = client.get("/api/v1/resources/lab-lessons", headers=HEADERS).json()
    assert theory["total"] == 37
    assert labs["total"] == 12
    assert all(row["purpose"] and row["environment"] and row["principle"] for row in labs["items"])
    audit = client.post("/api/v1/resources/audit/run", headers=HEADERS, json={"course_id": COURSE_ID})
    assert audit.status_code == 200
    assert audit.json()["total"] == 196
    assert audit.json()["blocking"] > 0
    manifest = client.get("/api/v1/resources/delivery/manifest.json", headers=HEADERS).json()
    assert manifest["status"] == "BLOCKED"
    assert manifest["theory_lessons"] == 37 and manifest["lab_lessons"] == 12
    xlsx = client.get("/api/v1/resources/delivery/manifest.xlsx", headers=HEADERS)
    assert xlsx.status_code == 200 and xlsx.content.startswith(b"PK")
    frozen = client.post("/api/v1/resources/delivery/freeze", headers=HEADERS, json={"course_id": COURSE_ID})
    assert frozen.status_code == 409
    assert frozen.json()["code"] == "RESOURCE.DELIVERY_BLOCKED"


def test_real_file_upload_version_download_and_readiness(client, tmp_path, monkeypatch):
    monkeypatch.setenv("YUEKE_RESOURCE_UPLOAD_DIR", str(tmp_path))
    engine = create_engine(os.environ["YUEKE_DATABASE_URL"])
    suffix = uuid4().hex[:8]
    resource_id = file_id = None
    content = b"PK\x03\x04real-pptx-upload-" + suffix.encode()
    try:
        uploaded = client.post(
            "/api/v1/resources/files",
            headers=HEADERS,
            data={"course_id": COURSE_ID},
            files={"file": (f"lesson-{suffix}.pptx", content, "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
        )
        assert uploaded.status_code == 201
        file_data = uploaded.json(); file_id = file_data["file_id"]
        assert file_data["size_bytes"] == len(content)
        assert len(file_data["sha256"]) == 64

        duplicate = client.post(
            "/api/v1/resources/files",
            headers=HEADERS,
            data={"course_id": COURSE_ID},
            files={"file": (f"copy-{suffix}.pptx", content, "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
        )
        assert duplicate.status_code == 201 and duplicate.json()["file_id"] == file_id

        resource = client.post("/api/v1/resources", headers=HEADERS, json={"course_id": COURSE_ID, "lesson_id": "lesson_theory_3_2", "name": f"真实上传讲义 {suffix}", "resource_type": "PPT"})
        assert resource.status_code == 201
        resource_id = resource.json()["resource_id"]
        version = client.post(f"/api/v1/resources/{resource_id}/versions", headers=HEADERS, json={"file_id": file_id, "sha256": file_data["sha256"]})
        assert version.status_code == 201

        listed = client.get("/api/v1/resources", headers=HEADERS, params={"course_id": COURSE_ID, "name": suffix}).json()
        assert listed["total"] == 1 and listed["items"][0]["latest_version"]["file_id"] == file_id
        downloaded = client.get(f"/api/v1/resources/{resource_id}/download", headers=HEADERS)
        assert downloaded.status_code == 200 and downloaded.content == content
        student_headers = {"X-User-Id": "student_b", "X-Role": "student", "X-Student-Id": "student_b", "X-Course-Ids": COURSE_ID, "X-Permissions": "resources:read"}
        student_download = client.get(f"/api/v1/resources/{resource_id}/download", headers=student_headers)
        assert student_download.status_code == 404

        readiness = client.get("/api/v1/resources/readiness", headers=HEADERS)
        assert readiness.status_code == 200
        assert readiness.json()["published_questions"]["required"] == 196
        assert readiness.json()["ppt"]["required"] == 37

        mismatch = client.post("/api/v1/resources", headers=HEADERS, json={"course_id": COURSE_ID, "lesson_id": "lesson_theory_3_2", "name": f"类型不符 {suffix}", "resource_type": "VIDEO"})
        mismatch_id = mismatch.json()["resource_id"]
        mismatch_version = client.post(f"/api/v1/resources/{mismatch_id}/versions", headers=HEADERS, json={"file_id": file_id, "sha256": file_data["sha256"]})
        assert mismatch_version.status_code == 422 and mismatch_version.json()["code"] == "RESOURCE.FILE_TYPE_MISMATCH"
        with Session(engine) as session:
            session.execute(delete(DomainEventOutbox).where(DomainEventOutbox.aggregate_id == mismatch_id))
            session.execute(delete(Resource).where(Resource.resource_id == mismatch_id))
            session.commit()

        unsupported = client.post("/api/v1/resources/files", headers=HEADERS, data={"course_id": COURSE_ID}, files={"file": ("bad.exe", b"bad", "application/octet-stream")})
        assert unsupported.status_code == 422 and unsupported.json()["code"] == "RESOURCE.FILE_TYPE_UNSUPPORTED"
    finally:
        with Session(engine) as session:
            if resource_id:
                version_ids = list(session.scalars(select(ResourceVersion.resource_version_id).where(ResourceVersion.resource_id == resource_id)))
                if version_ids:
                    session.execute(delete(PptAsset).where(PptAsset.resource_version_id.in_(version_ids)))
                session.execute(delete(ResourceVersion).where(ResourceVersion.resource_id == resource_id))
                session.execute(delete(DomainEventOutbox).where(DomainEventOutbox.aggregate_id == resource_id))
                session.execute(delete(Resource).where(Resource.resource_id == resource_id))
            if file_id:
                session.execute(delete(FileObject).where(FileObject.file_id == file_id))
            session.commit()


def test_three_dimension_filter_sha_duration_and_frozen_immutability(client, tmp_path):
    engine = create_engine(os.environ["YUEKE_DATABASE_URL"])
    suffix = uuid4().hex[:8]
    created_ids = []
    file_id = f"file_{suffix}"
    sha = "a" * 64
    video_path = tmp_path / f"video-{suffix}.mp4"
    video_path.write_bytes(b"not-a-real-video")
    with Session(engine) as session:
        session.add(FileObject(file_id=file_id, storage_provider="local", bucket="course-resources", object_key=str(video_path), original_name="真实解析样例.mp4", mime_type="video/mp4", size_bytes=123, sha256=sha, created_by="teacher_b", created_at=datetime.utcnow()))
        session.commit()
    try:
        for name, resource_type in [(f"RSA 视频 {suffix}", "VIDEO"), (f"RSA 讲义 {suffix}", "PPT")]:
            response = client.post("/api/v1/resources", headers=HEADERS, json={"course_id": COURSE_ID, "lesson_id": "lesson_theory_3_2", "name": name, "resource_type": resource_type})
            assert response.status_code == 201
            created_ids.append(response.json()["resource_id"])
        filtered = client.get("/api/v1/resources", headers=HEADERS, params={"course_id": COURSE_ID, "status": "DRAFT", "name": suffix, "resource_type": "VIDEO"}).json()
        assert filtered["total"] == 1 and filtered["items"][0]["resource_type"] == "VIDEO"
        video_id = filtered["items"][0]["resource_id"]
        bad_sha = client.post(f"/api/v1/resources/{video_id}/versions", headers=HEADERS, json={"file_id": file_id, "sha256": "b" * 64})
        assert bad_sha.status_code == 422 and bad_sha.json()["code"] == "RESOURCE.FILE_SHA256_MISMATCH"
        probe = client.post(f"/api/v1/resources/{video_id}/versions", headers=HEADERS, json={"file_id": file_id, "sha256": sha})
        assert probe.status_code in {422, 503}
        assert probe.json()["code"] in {"RESOURCE.MEDIA_PROBE_UNAVAILABLE", "RESOURCE.VIDEO_PROBE_FAILED"}
        with Session(engine) as session:
            resource = session.get(Resource, video_id); resource.status = "FROZEN"; session.commit()
        immutable = client.post(f"/api/v1/resources/{video_id}/versions", headers=HEADERS, json={"file_id": file_id, "sha256": sha})
        assert immutable.status_code == 409 and immutable.json()["code"] == "RESOURCE.VERSION_FROZEN"
    finally:
        with Session(engine) as session:
            for resource_id in created_ids:
                versions = list(session.scalars(select(ResourceVersion).where(ResourceVersion.resource_id == resource_id)))
                for item in versions:
                    session.execute(delete(VideoAsset).where(VideoAsset.resource_version_id == item.resource_version_id))
                session.execute(delete(ResourceVersion).where(ResourceVersion.resource_id == resource_id))
                session.execute(delete(Resource).where(Resource.resource_id == resource_id))
            session.execute(delete(FileObject).where(FileObject.file_id == file_id))
            session.commit()


def test_question_coverage_requires_four_published_types(client):
    engine = create_engine(os.environ["YUEKE_DATABASE_URL"])
    lesson_id = "lesson_theory_3_2"
    question_ids = []
    try:
        for question_type in ["FILL", "SINGLE", "MULTIPLE", "TRUE_FALSE"]:
            response = client.post("/api/v1/questions", headers=HEADERS, json={"course_id": COURSE_ID, "lesson_id": lesson_id, "question_type": question_type, "stem": f"{question_type} 测试题", "answer": ["A"], "explanation": "用于验证覆盖门禁", "options": [{"key": "A", "text": "正确答案", "is_correct": True}]})
            assert response.status_code == 201
            question_id = response.json()["question_id"]; question_ids.append(question_id)
            reviewed = client.post(f"/api/v1/questions/{question_id}/review", headers=REVIEW_HEADERS)
            assert reviewed.status_code == 200
        coverage = client.get("/api/v1/questions/coverage", headers=HEADERS).json()
        item = next(row for row in coverage["items"] if row["lesson_id"] == lesson_id)
        assert item["passed"] is True
        assert set(item["types"]) == {"FILL", "SINGLE", "MULTIPLE", "TRUE_FALSE"}
        student_headers = {"X-User-Id": "student_b", "X-Role": "student", "X-Student-Id": "student_b", "X-Course-Ids": COURSE_ID, "X-Permissions": "resources:read"}
        student_questions = client.get("/api/v1/questions", headers=student_headers).json()["items"]
        own_questions = [row for row in student_questions if row["question_id"] in question_ids]
        assert len(own_questions) == 4
        assert all("answer" not in row and "explanation" not in row for row in own_questions)
    finally:
        with Session(engine) as session:
            session.execute(delete(QuestionReview).where(QuestionReview.question_id.in_(question_ids)))
            session.execute(delete(QuestionOption).where(QuestionOption.question_id.in_(question_ids)))
            session.execute(delete(QuestionExplanation).where(QuestionExplanation.question_id.in_(question_ids)))
            session.execute(delete(QuestionLessonMap).where(QuestionLessonMap.question_id.in_(question_ids)))
            session.execute(delete(Question).where(Question.question_id.in_(question_ids)))
            remaining = session.scalar(select(Question).where(Question.question_bank_id == f"qb_{COURSE_ID}").limit(1))
            if not remaining: session.execute(delete(QuestionBank).where(QuestionBank.question_bank_id == f"qb_{COURSE_ID}"))
            session.commit()


def test_question_xlsx_import_idempotency_error_rows_and_independent_review(client):
    engine = create_engine(os.environ["YUEKE_DATABASE_URL"])
    suffix = uuid4().hex[:8]
    course_id = f"course_import_{suffix}"
    author_headers = {**HEADERS, "X-Course-Ids": course_id}
    reviewer_headers = {**REVIEW_HEADERS, "X-Course-Ids": course_id}
    bank_id = None
    job_ids: list[str] = []
    question_ids: list[str] = []

    with Session(engine) as session:
        for source in lesson_rows():
            values = dict(source)
            values["lesson_resource_id"] = str(uuid4())
            values["course_id"] = course_id
            session.add(LessonResource(**values))
        session.commit()

    try:
        template = client.get("/api/v1/questions/import-template.xlsx", headers=author_headers, params={"course_id": course_id})
        assert template.status_code == 200
        workbook = load_workbook(BytesIO(template.content))
        assert workbook["题目导入"].max_row == 197
        valid_content = complete_question_template(template.content)

        invalid_workbook = load_workbook(BytesIO(valid_content))
        invalid_workbook["题目导入"]["D2"] = ""
        invalid_stream = BytesIO()
        invalid_workbook.save(invalid_stream)
        invalid_content = invalid_stream.getvalue()
        invalid_key = f"invalid-{suffix}"
        invalid = client.post(
            "/api/v1/questions/import",
            headers={**author_headers, "Idempotency-Key": invalid_key},
            data={"course_id": course_id},
            files={"file": ("questions-invalid.xlsx", invalid_content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert invalid.status_code == 201
        invalid_job = invalid.json()
        job_ids.append(invalid_job["job_id"])
        assert invalid_job["status"] == "VALIDATION_FAILED"
        assert invalid_job["imported_count"] == 0
        assert invalid_job["error_count"] == 1
        assert {error["code"] for error in invalid_job["error_rows"]} == {"QUESTION_IMPORT.STEM_REQUIRED"}

        errors = client.get(f"/api/v1/questions/import-jobs/{invalid_job['job_id']}/error-rows.xlsx", headers=author_headers)
        assert errors.status_code == 200 and errors.content.startswith(b"PK")
        error_sheet = load_workbook(BytesIO(errors.content))["错误行"]
        assert error_sheet.cell(2, 1).value == 2
        assert "QUESTION_IMPORT.STEM_REQUIRED" in error_sheet.cell(2, 13).value

        with Session(engine) as session:
            bank = session.scalar(select(QuestionBank).where(QuestionBank.course_id == course_id))
            bank_id = bank.question_bank_id
            assert session.scalar(select(Question).where(Question.question_bank_id == bank_id).limit(1)) is None

        valid_key = f"valid-{suffix}"
        valid = client.post(
            "/api/v1/questions/import",
            headers={**author_headers, "Idempotency-Key": valid_key},
            data={"course_id": course_id},
            files={"file": ("questions.xlsx", valid_content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert valid.status_code == 201
        completed_job = valid.json()
        job_ids.append(completed_job["job_id"])
        assert completed_job["status"] == "COMPLETED"
        assert completed_job["total_count"] == completed_job["imported_count"] == completed_job["review_queue_count"] == 196
        assert completed_job["error_count"] == 0

        replay = client.post(
            "/api/v1/questions/import",
            headers={**author_headers, "Idempotency-Key": valid_key},
            data={"course_id": course_id},
            files={"file": ("questions.xlsx", valid_content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert replay.status_code == 201 and replay.json()["job_id"] == completed_job["job_id"]

        changed_workbook = load_workbook(BytesIO(valid_content))
        changed_workbook["题目导入"]["J2"] = "另一份解析"
        changed_stream = BytesIO()
        changed_workbook.save(changed_stream)
        conflict = client.post(
            "/api/v1/questions/import",
            headers={**author_headers, "Idempotency-Key": valid_key},
            data={"course_id": course_id},
            files={"file": ("questions.xlsx", changed_stream.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert conflict.status_code == 409 and conflict.json()["code"] == "REQUEST.IDEMPOTENCY_CONFLICT"

        fetched = client.get(f"/api/v1/questions/import-jobs/{completed_job['job_id']}", headers=author_headers)
        assert fetched.status_code == 200 and fetched.json()["imported_count"] == 196
        queue = client.get("/api/v1/questions/review-queue", headers=reviewer_headers, params={"course_id": course_id, "page_size": 200})
        assert queue.status_code == 200 and queue.json()["total"] == 196
        first = queue.json()["items"][0]
        assert first["can_review"] is True
        assert first["stem"] and first["answer"] and first["explanation"]
        assert first["lesson_id"] and first["lesson_code"] and first["source_row_number"]

        self_review = client.post(f"/api/v1/questions/{first['question_id']}/review", headers=author_headers)
        assert self_review.status_code == 409 and self_review.json()["code"] == "QUESTION.REVIEWER_MUST_BE_INDEPENDENT"
        approved = client.post(f"/api/v1/questions/{first['question_id']}/review", headers=reviewer_headers)
        assert approved.status_code == 200 and approved.json()["status"] == "PUBLISHED"

        second = queue.json()["items"][1]
        missing_reason = client.post(f"/api/v1/questions/{second['question_id']}/review", headers=reviewer_headers, json={"decision": "REJECTED"})
        assert missing_reason.status_code == 422 and missing_reason.json()["code"] == "QUESTION.REJECTION_COMMENT_REQUIRED"
        rejected = client.post(
            f"/api/v1/questions/{second['question_id']}/review",
            headers=reviewer_headers,
            json={"decision": "REJECTED", "comment": "题干需要补充限定条件"},
        )
        assert rejected.status_code == 200 and rejected.json()["status"] == "REJECTED"

        queue_after = client.get("/api/v1/questions/review-queue", headers=reviewer_headers, params={"course_id": course_id, "page_size": 200})
        assert queue_after.status_code == 200 and queue_after.json()["total"] == 194

        with Session(engine) as session:
            question_ids = list(session.scalars(select(Question.question_id).where(Question.question_bank_id == bank_id)))
            assert len(question_ids) == 196
            assert session.scalar(select(QuestionReview).where(QuestionReview.question_id == first["question_id"])) is not None
            assert session.scalar(select(DomainEventOutbox).where(DomainEventOutbox.aggregate_id == completed_job["job_id"])) is not None
    finally:
        with Session(engine) as session:
            if bank_id:
                question_ids = list(session.scalars(select(Question.question_id).where(Question.question_bank_id == bank_id)))
            if job_ids:
                session.execute(delete(QuestionImportRow).where(QuestionImportRow.import_job_id.in_(job_ids)))
            if question_ids:
                session.execute(delete(QuestionReview).where(QuestionReview.question_id.in_(question_ids)))
                session.execute(delete(QuestionOption).where(QuestionOption.question_id.in_(question_ids)))
                session.execute(delete(QuestionExplanation).where(QuestionExplanation.question_id.in_(question_ids)))
                session.execute(delete(QuestionLessonMap).where(QuestionLessonMap.question_id.in_(question_ids)))
                session.execute(delete(Question).where(Question.question_id.in_(question_ids)))
                session.execute(delete(DomainEventOutbox).where(DomainEventOutbox.aggregate_id.in_(question_ids)))
            if job_ids:
                session.execute(delete(DomainEventOutbox).where(DomainEventOutbox.aggregate_id.in_(job_ids)))
                session.execute(delete(QuestionImportJob).where(QuestionImportJob.import_job_id.in_(job_ids)))
            if bank_id:
                session.execute(delete(QuestionBank).where(QuestionBank.question_bank_id == bank_id))
            session.execute(delete(LessonResource).where(LessonResource.course_id == course_id))
            session.commit()


def test_question_import_and_review_are_serialized_under_concurrency(client, monkeypatch):
    engine = create_engine(os.environ["YUEKE_DATABASE_URL"])
    suffix = uuid4().hex[:8]
    replay_course = f"course_replay_{suffix}"
    competing_course = f"course_competing_{suffix}"
    courses = [replay_course, competing_course]

    def context(user_id: str, course_id: str) -> UserContext:
        return UserContext(
            user_id,
            "teacher",
            user_id,
            None,
            frozenset({"resources:read", "resources:write", "resources:review"}),
            frozenset({course_id}),
            frozenset(),
        )

    def run_import(course_id: str, key: str, gate: Barrier, content: bytes) -> dict:
        with Session(engine) as session:
            gate.wait(timeout=30)
            return ResourceService(session, context("concurrent_author", course_id)).import_questions(
                course_id, content, "questions.xlsx", key
            )

    with Session(engine) as session:
        for course_id in courses:
            for source in lesson_rows():
                values = dict(source)
                values["lesson_resource_id"] = str(uuid4())
                values["course_id"] = course_id
                session.add(LessonResource(**values))
        session.commit()

    try:
        template = client.get(
            "/api/v1/questions/import-template.xlsx",
            headers={**HEADERS, "X-Course-Ids": replay_course},
            params={"course_id": replay_course},
        )
        assert template.status_code == 200
        content = complete_question_template(template.content)

        replay_gate = Barrier(2)
        with ThreadPoolExecutor(max_workers=2) as executor:
            replay_results = list(
                executor.map(
                    lambda _: run_import(replay_course, f"same-{suffix}", replay_gate, content),
                    range(2),
                )
            )
        assert len({item["job_id"] for item in replay_results}) == 1

        competing_gate = Barrier(2)
        with ThreadPoolExecutor(max_workers=2) as executor:
            futures = [
                executor.submit(run_import, competing_course, f"key-{index}-{suffix}", competing_gate, content)
                for index in range(2)
            ]
            competing_results = [future.result(timeout=60) for future in futures]
        assert sorted(item["status"] for item in competing_results) == ["COMPLETED", "VALIDATION_FAILED"]

        with Session(engine) as session:
            replay_bank = session.scalar(select(QuestionBank).where(QuestionBank.course_id == replay_course))
            competing_bank = session.scalar(select(QuestionBank).where(QuestionBank.course_id == competing_course))
            assert len(list(session.scalars(select(Question).where(Question.question_bank_id == replay_bank.question_bank_id)))) == 196
            assert len(list(session.scalars(select(Question).where(Question.question_bank_id == competing_bank.question_bank_id)))) == 196
            question_id = session.scalar(
                select(Question.question_id).where(Question.question_bank_id == replay_bank.question_bank_id).limit(1)
            )

        review_gate = Barrier(2)
        lock_gate = Barrier(2)
        original_lock_question = ResourceRepository.lock_question

        def synchronized_lock_question(repository: ResourceRepository, target_question_id: str):
            lock_gate.wait(timeout=30)
            return original_lock_question(repository, target_question_id)

        monkeypatch.setattr(ResourceRepository, "lock_question", synchronized_lock_question)

        def review(user_id: str, decision: str):
            with Session(engine) as session:
                review_gate.wait(timeout=30)
                try:
                    result = ResourceService(session, context(user_id, replay_course)).review_question(
                        question_id,
                        QuestionReviewDecision(decision=decision, comment="并发审核验证"),
                    )
                    return "ok", result["status"]
                except ApiError as exc:
                    return "error", exc.code

        with ThreadPoolExecutor(max_workers=2) as executor:
            review_results = [
                executor.submit(review, "concurrent_reviewer_a", "APPROVED"),
                executor.submit(review, "concurrent_reviewer_b", "REJECTED"),
            ]
            review_results = [future.result(timeout=60) for future in review_results]
        assert sorted(item[0] for item in review_results) == ["error", "ok"]
        assert next(item[1] for item in review_results if item[0] == "error") == "QUESTION.INVALID_REVIEW_STATE"
        with Session(engine) as session:
            reviews = list(session.scalars(select(QuestionReview).where(QuestionReview.question_id == question_id)))
            assert len(reviews) == 1
    finally:
        with Session(engine) as session:
            for course_id in courses:
                bank_ids = list(session.scalars(select(QuestionBank.question_bank_id).where(QuestionBank.course_id == course_id)))
                question_ids = list(session.scalars(select(Question.question_id).where(Question.question_bank_id.in_(bank_ids)))) if bank_ids else []
                job_ids = list(session.scalars(select(QuestionImportJob.import_job_id).where(QuestionImportJob.course_id == course_id)))
                if job_ids:
                    session.execute(delete(QuestionImportRow).where(QuestionImportRow.import_job_id.in_(job_ids)))
                if question_ids:
                    session.execute(delete(QuestionReview).where(QuestionReview.question_id.in_(question_ids)))
                    session.execute(delete(QuestionOption).where(QuestionOption.question_id.in_(question_ids)))
                    session.execute(delete(QuestionExplanation).where(QuestionExplanation.question_id.in_(question_ids)))
                    session.execute(delete(QuestionLessonMap).where(QuestionLessonMap.question_id.in_(question_ids)))
                    session.execute(delete(Question).where(Question.question_id.in_(question_ids)))
                    session.execute(delete(DomainEventOutbox).where(DomainEventOutbox.aggregate_id.in_(question_ids)))
                if job_ids:
                    session.execute(delete(DomainEventOutbox).where(DomainEventOutbox.aggregate_id.in_(job_ids)))
                    session.execute(delete(QuestionImportJob).where(QuestionImportJob.import_job_id.in_(job_ids)))
                if bank_ids:
                    session.execute(delete(QuestionBank).where(QuestionBank.question_bank_id.in_(bank_ids)))
                session.execute(delete(LessonResource).where(LessonResource.course_id == course_id))
            session.commit()


def test_freeze_pass_creates_manifest_when_dynamic_audit_has_no_blockers(monkeypatch):
    engine = create_engine(os.environ["YUEKE_DATABASE_URL"])
    context = UserContext("reviewer_b", "teacher", "reviewer_b", None, frozenset({"resources:read", "resources:freeze"}), frozenset({COURSE_ID}), frozenset())
    with Session(engine) as session:
        session.execute(delete(ResourceDeliveryManifest).where(ResourceDeliveryManifest.course_id == COURSE_ID))
        session.execute(delete(DomainEventOutbox).where(DomainEventOutbox.event_type == "resource.delivery.frozen", DomainEventOutbox.aggregate_id == COURSE_ID))
        session.commit()
        svc = ResourceService(session, context)
        clean = {"course_id": COURSE_ID, "total": 196, "pass": 196, "warning": 0, "blocking": 0, "blocking_items": [], "procurement_mapping": svc.procurement_mapping(), "checked_at": datetime.utcnow().isoformat() + "Z"}
        monkeypatch.setattr(svc, "audit", lambda course_id, persist=False: clean)
        manifest = svc.freeze(COURSE_ID)
        assert manifest["status"] == "FROZEN"
        assert session.scalar(select(ResourceDeliveryManifest).where(ResourceDeliveryManifest.course_id == COURSE_ID)) is not None
        event = session.scalar(select(DomainEventOutbox).where(DomainEventOutbox.event_type == "resource.delivery.frozen", DomainEventOutbox.aggregate_id == COURSE_ID))
        assert event.payload_json["course_id"] == COURSE_ID
        session.execute(delete(ResourceDeliveryManifest).where(ResourceDeliveryManifest.course_id == COURSE_ID))
        session.execute(delete(DomainEventOutbox).where(DomainEventOutbox.event_type == "resource.delivery.frozen", DomainEventOutbox.aggregate_id == COURSE_ID))
        session.commit()
