import json
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

from app.resources.catalog import lesson_rows
from app.resources.question_xlsx import parse_question_workbook


ROOT = Path(__file__).resolve().parents[2]
CONTENT_PATH = ROOT / "backend" / "app" / "resources" / "content" / "question-bank-v1.json"
WORKBOOK_PATH = ROOT / "outputs" / "01a0c33f-d483-7ac0-aa95-632194b582d5" / "question-bank-196-v1.xlsx"


def test_formal_question_bank_matches_catalog_and_import_contract():
    content = json.loads(CONTENT_PATH.read_text(encoding="utf-8"))
    catalog = lesson_rows()

    assert content["courseId"] == "course_data_security"
    assert len(content["lessons"]) == len(catalog) == 49
    assert {(lesson["lessonId"], lesson["lessonCode"]) for lesson in content["lessons"]} == {
        (lesson["lesson_id"], lesson["lesson_code"]) for lesson in catalog
    }
    assert len(content["sources"]) >= 5
    assert any("2026年1月1日" in source["title"] for source in content["sources"])

    parsed, total = parse_question_workbook(WORKBOOK_PATH.read_bytes(), catalog)
    assert total == 196
    assert len(parsed) == 196
    assert all(row["status"] == "VALID" and not row["errors"] for row in parsed)

    questions = [row["normalized_data"] for row in parsed]
    assert Counter(question["question_type"] for question in questions) == {
        "FILL": 49,
        "SINGLE": 49,
        "MULTIPLE": 49,
        "TRUE_FALSE": 49,
    }
    assert len({question["stem"] for question in questions}) == 196
    assert all(question["answer"] and question["explanation"] for question in questions)
    assert not any("示例题目" in question["stem"] or "待补充" in question["stem"] for question in questions)

    single_answers = Counter(question["answer"][0] for question in questions if question["question_type"] == "SINGLE")
    assert set(single_answers) == {"A", "B", "C", "D"}
    assert max(single_answers.values()) - min(single_answers.values()) <= 1


def test_formal_question_bank_workbook_is_safe_and_auditable():
    workbook = load_workbook(WORKBOOK_PATH, read_only=True, data_only=False)
    assert workbook.sheetnames == ["题目导入", "填写说明"]
    questions = workbook["题目导入"]
    question_rows = list(questions.iter_rows(min_col=1, max_col=10))
    assert len(question_rows) == 197
    assert not any(cell.data_type == "f" for row in question_rows for cell in row)

    guide_values = "\n".join(str(cell.value or "") for row in workbook["填写说明"].iter_rows() for cell in row)
    assert "待独立审核" in guide_values
    assert "网络安全法" in guide_values
    assert "FIPS（美国联邦信息处理标准）197" in guide_values
