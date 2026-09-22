import os
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta
from io import BytesIO
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import text

from app.auth import models as auth
from app.common.context import RequestStateTrustedIdentityResolver, TrustedIdentity, configure_trusted_identity_resolver
from app.database import SessionLocal, engine
from app.lab_classroom.gateway import ContractClient, GatewayBundle, get_gateways
from app.main import app
from app.teaching import models as teaching_models
from app.teaching.xlsx import HEADERS
from backend.tests.auth_profiles import seed_student_profile


pytestmark = pytest.mark.skipif(not os.getenv("YUEKE_DATABASE_URL"), reason="需要专属 MySQL 集成库")


def identity(*permissions: str, course_id: str = "", class_id: str = "", student_id: str = ""):
    headers = {"X-User-Id": student_id or "mysql-teacher", "X-Role": "student" if student_id else "teacher", "X-Permissions": ",".join(permissions), "X-Course-Ids": course_id, "X-Class-Ids": class_id}
    headers["X-Student-Id" if student_id else "X-Teacher-Id"] = student_id or "mysql-teacher"
    return headers


def xlsx_43(class_name: str) -> tuple[bytes, list[list[str]]]:
    book = Workbook(); sheet = book.active; sheet.append(HEADERS)
    run = uuid4().hex[:6]
    rows = [[f"{run}{number:03d}", f"验收学生{number}", class_name, "", ""] for number in range(1, 44)]
    for row in rows: sheet.append(row)
    stream = BytesIO(); book.save(stream); return stream.getvalue(), rows


def test_mysql_84_g1_g2_main_chain():
    assert engine.dialect.name == "mysql"
    with engine.connect() as connection:
        assert connection.scalar(text("select version_num from alembic_version"))
        assert connection.scalar(text("select count(*) from information_schema.tables where table_schema = database() and table_name = 'class_membership'")) == 1
        assert connection.scalar(text("select count(*) from information_schema.tables where table_schema = database() and table_name = 'auth_user'")) == 1
        assert connection.scalar(text("select count(*) from information_schema.tables where table_schema = database() and table_name = 'student_profile'")) == 1
        assert connection.scalar(text("select count(*) from information_schema.tables where table_schema = database() and table_name = 'identity_reconciliation_case'")) == 1
        assert connection.scalar(text("select version()"))
    client = TestClient(app)
    course = client.post("/api/v1/courses", headers=identity("teaching.course.write"), json={"name":"MySQL 验收课程","term":"2026 秋季"})
    assert course.status_code == 201, course.text
    course_id = course.json()["course_id"]
    teaching = identity("teaching.class.write","teaching.members.import","teaching.members.read","teaching.members.write","teaching.roster.freeze","teaching.attendance.write","teaching.attendance.read", course_id=course_id)
    created_class = client.post("/api/v1/classes", headers=teaching, json={"name":"MySQL 验收班","term":"2026 秋季","course_id":course_id})
    assert created_class.status_code == 201, created_class.text
    class_id = created_class.json()["class_id"]; teaching["X-Class-Ids"] = class_id
    import_key, (import_file, rows) = uuid4().hex, xlsx_43("MySQL 验收班")
    admin = {
        "X-User-Id": "mysql-admin",
        "X-Role": "admin",
        "X-Permissions": "auth.accounts.write,auth.accounts.read",
    }
    created_account = client.post(
        "/api/v1/auth/users",
        headers=admin,
        json={
            "display_name": rows[0][1],
            "role": "student",
            "login_name": f"mysql-{rows[0][0]}",
            "student_number": rows[0][0],
        },
    )
    assert created_account.status_code == 201, created_account.text
    assert created_account.json()["student_number"] == rows[0][0]
    assert "password" not in created_account.json()
    expected_student_ids = {rows[0][0]: created_account.json()["student_id"]}
    with SessionLocal() as db:
        for row in rows[1:]:
            expected_student_ids[row[0]] = seed_student_profile(db, student_number=row[0], full_name=row[1]).student_id
        db.commit()
    def import_once():
        with TestClient(app) as concurrent_client:
            return concurrent_client.post(f"/api/v1/classes/{class_id}/members/import", headers={**teaching,"Idempotency-Key":import_key}, files={"file":("students.xlsx",import_file)})
    with ThreadPoolExecutor(max_workers=2) as pool:
        imports = list(pool.map(lambda _: import_once(), range(2)))
    assert all(response.status_code == 201 and response.json()["success_count"] == 43 for response in imports)
    assert len({response.json()["job_id"] for response in imports}) == 1
    members = client.get(f"/api/v1/classes/{class_id}/members?page_size=100", headers=teaching).json()
    assert members["total"] == 43
    assert {item["student_number"]: item["student_id"] for item in members["items"]} == expected_student_ids
    start, end = datetime.utcnow()-timedelta(minutes=1), datetime.utcnow()+timedelta(minutes=5)
    task = client.post("/api/v1/attendance/tasks", headers=teaching, json={"course_id":course_id,"class_id":class_id,"task_type":"CLASSROOM","title":"真实 MySQL 签到","starts_at":start.isoformat(),"expires_at":end.isoformat()}).json()
    published = client.post(f"/api/v1/attendance/tasks/{task['task_id']}/publish", headers=teaching).json()
    student_id = members["items"][0]["student_id"]
    student = identity("teaching.attendance.sign", student_id=student_id)
    link = client.get(f"/api/v1/attendance/sign-links/{published['sign_token']}", headers=student)
    assert link.status_code == 200 and link.json()["task_id"] == task["task_id"]
    signed = client.post(f"/api/v1/attendance/sign-links/{published['sign_token']}/sign", headers=student)
    assert signed.status_code == 200
    summary = client.get("/api/v1/attendance/section-summary", headers=teaching).json()["items"][0]
    assert summary["expected"] == 43 and summary["present"] == 1
    frozen = client.post(f"/api/v1/classes/{class_id}/roster/freeze", headers=teaching)
    assert frozen.status_code == 200 and frozen.json()["member_count"] == 43
    blocked = client.post(f"/api/v1/classes/{class_id}/members", headers=teaching, json={"student_number":"late-001","student_name":"迟到名单"})
    assert blocked.status_code == 409 and blocked.json()["code"] == "CLASS.ROSTER_FROZEN"
    dispatcher = {"X-User-Id":"service_contract_dispatcher","X-Role":"admin","X-Permissions":"integration:dispatch"}
    dispatched = client.post("/api/v1/integration/outbox/dispatch", headers=dispatcher, params={"limit":500})
    assert dispatched.status_code == 200
    results = dispatched.json()["results"]
    assert any(item["event_type"] == "course.roster.frozen" and item["status"] == "PUBLISHED" and item["targets"] == ["grading_facts"] for item in results)
    teaching_audits = [item for item in results if item["event_type"] == "teaching.audit"]
    assert teaching_audits and all(item["status"] == "PUBLISHED" and item["targets"] == ["audit_event"] for item in teaching_audits)


def test_mysql_trusted_identity_uses_0018_role_mappings_for_interactive_modules():
    """真实 MySQL 上的可信身份路径：不得由开发身份头掩盖角色权限遗漏。"""

    suffix = uuid4().hex[:12]
    teacher_user_id, student_user_id = f"mysql-teacher-{suffix}", f"mysql-student-{suffix}"
    teacher_id, student_id = f"mysql-teacher-profile-{suffix}", f"mysql-student-profile-{suffix}"
    course_id, class_id = f"mysql-course-{suffix}", f"mysql-class-{suffix}"
    stamp = datetime.utcnow()
    with SessionLocal() as db:
        role_ids = dict(db.execute(text("SELECT code, role_id FROM auth_role WHERE code IN ('teacher', 'student')")).all())
        assert set(role_ids) == {"teacher", "student"}
        db.add_all([
            auth.AuthUser(user_id=teacher_user_id, external_subject=f"sso:mysql-teacher-{suffix}", login_name=f"mysql-teacher-{suffix}", display_name="可信教师", status="ACTIVE", primary_role_code="teacher", created_at=stamp, updated_at=stamp),
            auth.AuthUser(user_id=student_user_id, external_subject=f"sso:mysql-student-{suffix}", login_name=f"mysql-student-{suffix}", display_name="可信学生", status="ACTIVE", primary_role_code="student", created_at=stamp, updated_at=stamp),
            auth.AuthUserRole(auth_user_role_id=f"mysql-role-teacher-{suffix}", user_id=teacher_user_id, role_id=role_ids["teacher"], assigned_at=stamp),
            auth.AuthUserRole(auth_user_role_id=f"mysql-role-student-{suffix}", user_id=student_user_id, role_id=role_ids["student"], assigned_at=stamp),
            auth.TeacherProfile(teacher_id=teacher_id, user_id=teacher_user_id, staff_number=f"T-{suffix}", display_name="可信教师", status="ACTIVE", created_at=stamp, updated_at=stamp),
            auth.StudentProfile(student_id=student_id, user_id=student_user_id, student_number=f"2026{suffix[:8]}", full_name="可信学生", phone=None, email=None, status="ACTIVE", created_at=stamp, updated_at=stamp),
            teaching_models.Course(course_id=course_id, name="可信身份权限课程", term="2026 秋季", owner_teacher_id=teacher_id, major=None, description=None, status="ACTIVE", created_at=stamp),
            teaching_models.TeachingClass(class_id=class_id, name="可信身份权限班", term="2026 秋季", owner_teacher_id=teacher_id, created_at=stamp, roster_frozen_at=None, roster_frozen_by=None, roster_snapshot_hash=None),
        ])
        # 这些教学模型没有 ORM 关系；真实 MySQL 必须先落父表，不能靠 SQLite 的宽松行为掩盖外键顺序。
        db.flush()
        db.add_all([
            teaching_models.ClassCourse(class_course_id=f"mysql-class-course-{suffix}", class_id=class_id, course_id=course_id),
            teaching_models.TeachingTeacherAssignment(assignment_id=f"mysql-assignment-{suffix}", teacher_id=teacher_id, class_id=class_id, course_id=course_id),
            teaching_models.ClassMembership(class_membership_id=f"mysql-membership-{suffix}", class_id=class_id, student_id=student_id, student_number=f"2026{suffix[:8]}", student_name="可信学生", phone=None, email=None, status="ACTIVE", joined_at=stamp),
        ])
        db.commit()

    class Resolver:
        subject = f"sso:mysql-teacher-{suffix}"

        def resolve(self, request):
            return TrustedIdentity(external_subject=self.subject)

    resolver = Resolver()
    configure_trusted_identity_resolver(app, resolver)
    app.dependency_overrides[get_gateways] = lambda: GatewayBundle(
        runtime=ContractClient("D 实验运行", None),
        teaching=ContractClient("A 教学核心", None),
        grading=ContractClient("F 成绩审计", None),
    )
    try:
        client = TestClient(app)
        teacher_context = client.get("/api/v1/auth/context", headers={"X-Role": "admin", "X-Permissions": "infrastructure.write"})
        assert teacher_context.status_code == 200, teacher_context.text
        teacher_permissions = set(teacher_context.json()["permissions"])
        assert {"resources:read", "labs.read", "classroom.release.read", "runtime.read", "grading:read", "archives:read", "audit:read"}.issubset(teacher_permissions)
        assert client.get("/api/v1/resources", params={"course_id": course_id}).status_code == 200
        assert client.get("/api/v1/labs").status_code == 200

        resolver.subject = f"sso:mysql-student-{suffix}"
        student_context = client.get("/api/v1/auth/context")
        assert student_context.status_code == 200, student_context.text
        student_permissions = set(student_context.json()["permissions"])
        assert {"classroom.lab.start", "classroom.lab.read", "classroom.lab.submit", "runtime.start", "runtime.read", "runtime.submit", "runtime.terminal", "grading:read", "analytics:read"}.issubset(student_permissions)
        classroom = client.post("/api/v1/classroom/my/lab-releases/missing-release/start")
        assert classroom.status_code == 503 and classroom.json()["code"] == "DEPENDENCY.PENDING"
        runtime = client.get("/api/v1/runtime/requests/missing-request")
        assert runtime.status_code == 404 and runtime.json()["code"] == "RUNTIME.REQUEST_NOT_FOUND"
    finally:
        app.dependency_overrides.pop(get_gateways, None)
        configure_trusted_identity_resolver(app, RequestStateTrustedIdentityResolver())
