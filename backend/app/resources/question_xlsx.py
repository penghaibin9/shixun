from __future__ import annotations

import re
from io import BytesIO
from itertools import islice
from typing import Iterable
from zipfile import BadZipFile, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


HEADERS = [
    "课时标识*",
    "课时编号",
    "题型*",
    "题干*",
    "选项A",
    "选项B",
    "选项C",
    "选项D",
    "正确答案*",
    "解析*",
]

QUESTION_TYPE_LABELS = {
    "填空": "FILL",
    "填空题": "FILL",
    "单选": "SINGLE",
    "单选题": "SINGLE",
    "多选": "MULTIPLE",
    "多选题": "MULTIPLE",
    "判断": "TRUE_FALSE",
    "判断题": "TRUE_FALSE",
}
QUESTION_TYPE_NAMES = {
    "FILL": "填空",
    "SINGLE": "单选",
    "MULTIPLE": "多选",
    "TRUE_FALSE": "判断",
}
QUESTION_TYPE_ORDER = ("FILL", "SINGLE", "MULTIPLE", "TRUE_FALSE")
MAX_IMPORT_ROWS = 196
MAX_XLSX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_XLSX_ARCHIVE_ENTRIES = 256
MAX_XLSX_UNCOMPRESSED_BYTES = 64 * 1024 * 1024


class InvalidQuestionWorkbook(ValueError):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


def _lesson_value(lesson, name: str):
    if isinstance(lesson, dict):
        return lesson[name]
    return getattr(lesson, name)


def template_bytes(lessons: Iterable) -> bytes:
    lesson_items = list(lessons)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "题目导入"
    sheet.append(HEADERS)

    header_fill = PatternFill("solid", fgColor="4267D5")
    readonly_fill = PatternFill("solid", fgColor="E9EEF8")
    editable_fill = PatternFill("solid", fgColor="FFF2CC")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for lesson in lesson_items:
        lesson_id = _lesson_value(lesson, "lesson_id")
        lesson_code = _lesson_value(lesson, "lesson_code")
        for question_type in QUESTION_TYPE_ORDER:
            row = [lesson_id, lesson_code, QUESTION_TYPE_NAMES[question_type], "", "", "", "", "", "", ""]
            if question_type == "TRUE_FALSE":
                row[4], row[5] = "正确", "错误"
            sheet.append(row)

    last_row = sheet.max_row
    for row in sheet.iter_rows(min_row=2, max_row=last_row, min_col=1, max_col=len(HEADERS)):
        for cell in row[:3]:
            cell.fill = readonly_fill
        for cell in row[3:]:
            cell.fill = editable_fill
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    type_validation = DataValidation(type="list", formula1='"填空,单选,多选,判断"', allow_blank=False)
    type_validation.error = "题型只能选择填空、单选、多选或判断"
    type_validation.errorTitle = "题型无效"
    sheet.add_data_validation(type_validation)
    type_validation.add(f"C2:C{last_row}")

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:J{last_row}"
    widths = {"A": 26, "B": 14, "C": 12, "D": 42, "E": 24, "F": 24, "G": 24, "H": 24, "I": 24, "J": 42}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    guide = workbook.create_sheet("填写说明")
    guide_rows = [
        ["题库批量导入说明"],
        ["1", "模板已按 49 个课时、每课时四种题型预置 196 行，请勿增删行或修改课时标识、课时编号和题型。"],
        ["2", "单选题正确答案填写一个选项字母，例如 A。"],
        ["3", "多选题正确答案使用英文逗号分隔，例如 A,B。"],
        ["4", "判断题固定选项 A=正确、B=错误，正确答案填写 A 或 B。"],
        ["5", "填空题不要填写选项；多个可接受答案使用竖线分隔，例如 密钥|key。"],
        ["6", "题干、正确答案和解析均不能为空；任何包含公式或以 =、+、-、@ 开头的文本都会被拒绝。"],
    ]
    for row in guide_rows:
        guide.append(row)
    guide["A1"].font = Font(bold=True, size=14)
    guide.column_dimensions["A"].width = 8
    guide.column_dimensions["B"].width = 100
    for row in guide.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _text(value) -> str:
    return str(value if value is not None else "").strip()


def _error(field: str, code: str, message: str) -> dict:
    return {"field": field, "code": code, "message": message}


def _answer_tokens(value: str) -> list[str]:
    return [part.strip().upper() for part in re.split(r"[,，]", value) if part.strip()]


def _fill_answers(value: str) -> list[str]:
    return [part.strip() for part in re.split(r"[|｜]", value) if part.strip()]


def _safe_export_text(value) -> str:
    text = _text(value)
    if text.lstrip().startswith(("=", "+", "-", "@")):
        return f"'{text}"
    return text


def _validate_xlsx_archive(data: bytes) -> None:
    if len(data) > MAX_XLSX_UPLOAD_BYTES:
        raise InvalidQuestionWorkbook("QUESTION_IMPORT.FILE_TOO_LARGE", "题库导入文件不能超过 10 MB")
    try:
        with ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_XLSX_ARCHIVE_ENTRIES:
                raise InvalidQuestionWorkbook("QUESTION_IMPORT.ARCHIVE_TOO_COMPLEX", "XLSX（电子表格）内部文件数超出限制")
            if any(entry.flag_bits & 0x1 for entry in entries):
                raise InvalidQuestionWorkbook("QUESTION_IMPORT.ENCRYPTED_XLSX", "不支持加密的 XLSX（电子表格）")
            if sum(entry.file_size for entry in entries) > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise InvalidQuestionWorkbook("QUESTION_IMPORT.ARCHIVE_TOO_LARGE", "XLSX（电子表格）解压后内容超出限制")
    except BadZipFile as exc:
        raise InvalidQuestionWorkbook("QUESTION_IMPORT.INVALID_XLSX", "文件不是有效的 XLSX（电子表格）") from exc


def parse_question_workbook(data: bytes, lessons: Iterable) -> tuple[list[dict], int]:
    _validate_xlsx_archive(data)
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=False)
    except Exception as exc:
        raise InvalidQuestionWorkbook("QUESTION_IMPORT.INVALID_XLSX", "文件不是有效的 XLSX（电子表格）") from exc

    if "题目导入" not in workbook.sheetnames:
        raise InvalidQuestionWorkbook("QUESTION_IMPORT.SHEET_MISSING", "缺少“题目导入”工作表")
    sheet = workbook["题目导入"]
    header = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1, min_col=1, max_col=len(HEADERS)))]
    if header != HEADERS:
        raise InvalidQuestionWorkbook("QUESTION_IMPORT.HEADER_MISMATCH", "题目导入表头与模板不一致")

    lesson_items = list(lessons)
    by_id = {_lesson_value(item, "lesson_id"): item for item in lesson_items}
    by_code = {_lesson_value(item, "lesson_code"): item for item in lesson_items}

    rows = list(islice(sheet.iter_rows(min_row=2, min_col=1, max_col=len(HEADERS)), MAX_IMPORT_ROWS + 1))
    has_extra_row = len(rows) > MAX_IMPORT_ROWS
    if has_extra_row:
        rows = rows[:MAX_IMPORT_ROWS]
        actual_count = max((sheet.max_row or (MAX_IMPORT_ROWS + 2)) - 1, MAX_IMPORT_ROWS + 1)
    else:
        while rows and all(cell.value in (None, "") for cell in rows[-1]):
            rows.pop()
        actual_count = len(rows)
    results: list[dict] = []
    seen_slots: dict[tuple[str, str], int] = {}

    for excel_row_number, cells in enumerate(rows, start=2):
        raw = {HEADERS[index]: cells[index].value for index in range(len(HEADERS))}
        errors: list[dict] = []
        for index, cell in enumerate(cells):
            value = cell.value
            if cell.data_type == "f" or (isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@"))):
                errors.append(_error(HEADERS[index], "QUESTION_IMPORT.DANGEROUS_CELL", "单元格不能包含公式或以 =、+、-、@ 开头"))

        lesson_id = _text(raw["课时标识*"])
        lesson_code = _text(raw["课时编号"])
        type_label = _text(raw["题型*"])
        question_type = QUESTION_TYPE_LABELS.get(type_label)
        stem = _text(raw["题干*"])
        explanation = _text(raw["解析*"])
        answer_value = _text(raw["正确答案*"])
        option_texts = {key: _text(raw[f"选项{key}"]) for key in "ABCD"}

        lesson_by_id = by_id.get(lesson_id)
        lesson_by_code = by_code.get(lesson_code)
        if not lesson_id:
            errors.append(_error("课时标识*", "QUESTION_IMPORT.LESSON_ID_REQUIRED", "课时标识不能为空"))
        elif not lesson_by_id:
            errors.append(_error("课时标识*", "QUESTION_IMPORT.LESSON_NOT_FOUND", "课时标识不属于当前课程"))
        if not lesson_code:
            errors.append(_error("课时编号", "QUESTION_IMPORT.LESSON_CODE_REQUIRED", "课时编号不能为空"))
        elif not lesson_by_code:
            errors.append(_error("课时编号", "QUESTION_IMPORT.LESSON_CODE_NOT_FOUND", "课时编号不属于当前课程"))
        if lesson_by_id and lesson_by_code and _lesson_value(lesson_by_id, "lesson_id") != _lesson_value(lesson_by_code, "lesson_id"):
            errors.append(_error("课时编号", "QUESTION_IMPORT.LESSON_MISMATCH", "课时标识与课时编号不匹配"))
        if not question_type:
            errors.append(_error("题型*", "QUESTION_IMPORT.QUESTION_TYPE_INVALID", "题型只能是填空、单选、多选或判断"))
        if not stem:
            errors.append(_error("题干*", "QUESTION_IMPORT.STEM_REQUIRED", "题干不能为空"))
        if not explanation:
            errors.append(_error("解析*", "QUESTION_IMPORT.EXPLANATION_REQUIRED", "解析不能为空"))
        if not answer_value:
            errors.append(_error("正确答案*", "QUESTION_IMPORT.ANSWER_REQUIRED", "正确答案不能为空"))

        answer: list[str] = []
        options: list[dict] = []
        if question_type == "FILL":
            answer = _fill_answers(answer_value)
            if answer_value and not answer:
                errors.append(_error("正确答案*", "QUESTION_IMPORT.FILL_ANSWER_INVALID", "填空题答案格式无效"))
            if any(option_texts.values()):
                errors.append(_error("选项A", "QUESTION_IMPORT.FILL_OPTIONS_NOT_ALLOWED", "填空题不能填写选项"))
        elif question_type in {"SINGLE", "MULTIPLE"}:
            present_options = {key: value for key, value in option_texts.items() if value}
            if len(present_options) < 2:
                errors.append(_error("选项A", "QUESTION_IMPORT.OPTIONS_INSUFFICIENT", "单选题和多选题至少需要两个非空选项"))
            answer = _answer_tokens(answer_value)
            if question_type == "SINGLE" and answer_value and len(answer) != 1:
                errors.append(_error("正确答案*", "QUESTION_IMPORT.SINGLE_ANSWER_COUNT", "单选题必须且只能填写一个正确选项"))
            if question_type == "MULTIPLE" and answer_value and len(set(answer)) < 2:
                errors.append(_error("正确答案*", "QUESTION_IMPORT.MULTIPLE_ANSWER_COUNT", "多选题至少需要两个不同的正确选项"))
            if len(answer) != len(set(answer)):
                errors.append(_error("正确答案*", "QUESTION_IMPORT.ANSWER_DUPLICATE", "正确答案中不能重复填写同一选项"))
            missing_answers = [key for key in answer if key not in present_options]
            if missing_answers:
                errors.append(_error("正确答案*", "QUESTION_IMPORT.ANSWER_OPTION_NOT_FOUND", f"正确答案引用了不存在的选项：{','.join(missing_answers)}"))
            options = [{"key": key, "text": value, "is_correct": key in answer} for key, value in present_options.items()]
        elif question_type == "TRUE_FALSE":
            if option_texts["A"] != "正确" or option_texts["B"] != "错误" or option_texts["C"] or option_texts["D"]:
                errors.append(_error("选项A", "QUESTION_IMPORT.TRUE_FALSE_OPTIONS_INVALID", "判断题必须固定为选项 A=正确、B=错误，且 C、D 留空"))
            normalized_answer = {"正确": "A", "错误": "B", "A": "A", "B": "B"}.get(answer_value.upper() if answer_value.upper() in {"A", "B"} else answer_value)
            if answer_value and not normalized_answer:
                errors.append(_error("正确答案*", "QUESTION_IMPORT.TRUE_FALSE_ANSWER_INVALID", "判断题正确答案只能填写 A、B、正确或错误"))
            answer = [normalized_answer] if normalized_answer else []
            options = [
                {"key": "A", "text": "正确", "is_correct": normalized_answer == "A"},
                {"key": "B", "text": "错误", "is_correct": normalized_answer == "B"},
            ]

        normalized = None
        if lesson_by_id and lesson_by_code and question_type:
            normalized = {
                "lesson_id": lesson_id,
                "lesson_code": lesson_code,
                "question_type": question_type,
                "stem": stem,
                "answer": answer,
                "explanation": explanation,
                "options": options,
            }
            slot = (lesson_id, question_type)
            if slot in seen_slots:
                errors.append(_error("题型*", "QUESTION_IMPORT.SLOT_DUPLICATE", f"与第 {seen_slots[slot]} 行重复占用同一课时题型"))
            else:
                seen_slots[slot] = excel_row_number

        results.append(
            {
                "row_number": excel_row_number,
                "status": "ERROR" if errors else "VALID",
                "raw_data": {key: _text(value) for key, value in raw.items()},
                "normalized_data": normalized,
                "errors": errors,
            }
        )

    file_errors: list[dict] = []
    if actual_count != MAX_IMPORT_ROWS:
        file_errors.append(_error("文件", "QUESTION_IMPORT.ROW_COUNT_INVALID", f"题目数据必须正好 196 行，当前为 {actual_count} 行"))

    expected_slots = {(_lesson_value(lesson, "lesson_id"), question_type) for lesson in lesson_items for question_type in QUESTION_TYPE_ORDER}
    actual_slots = set(seen_slots)
    missing_slots = sorted(expected_slots - actual_slots)
    for lesson_id, question_type in missing_slots:
        lesson = by_id[lesson_id]
        file_errors.append(
            _error(
                "题型*",
                "QUESTION_IMPORT.SLOT_MISSING",
                f"缺少课时 {_lesson_value(lesson, 'lesson_code')} 的{QUESTION_TYPE_NAMES[question_type]}题",
            )
        )
    if file_errors:
        results.insert(
            0,
            {
                "row_number": 0,
                "status": "ERROR",
                "raw_data": {},
                "normalized_data": None,
                "errors": file_errors,
            },
        )

    return results, actual_count


def error_rows_bytes(rows: Iterable[dict]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "错误行"
    sheet.append(["原行号", *HEADERS, "错误字段", "错误代码", "错误原因"])
    for item in rows:
        errors = item.get("errors") or []
        if not errors:
            continue
        raw = item.get("raw_data") or {}
        sheet.append(
            [
                item.get("row_number", ""),
                *[_safe_export_text(raw.get(header, "")) for header in HEADERS],
                "\n".join(str(error.get("field", "")) for error in errors),
                "\n".join(str(error.get("code", "")) for error in errors),
                "\n".join(str(error.get("message", "")) for error in errors),
            ]
        )
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="B91C1C")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in range(1, sheet.max_column + 1):
        sheet.column_dimensions[chr(64 + column)].width = 22 if column != 1 else 10
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()
