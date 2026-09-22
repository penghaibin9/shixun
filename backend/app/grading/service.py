from collections import defaultdict
from datetime import datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from hashlib import sha256
from io import BytesIO
import json
import os
from pathlib import Path
import re
from stat import S_IRGRP, S_IROTH, S_IRUSR
from uuid import uuid4

from openpyxl import Workbook
from sqlalchemy import delete, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.common.context import UserContext
from app.common.errors import ApiError
from app.common.models import DomainEventOutbox, FileObject
from app.common.outbox import enqueue_event
from app.resources.storage import upload_root

from .models import (AnalyticsCourseSummary, AnalyticsLabSummary, AnalyticsSectionSummary, AnalyticsStudentLabSummary,
                     AuditEvent, CourseArchive, CourseArchiveArtifact, GradeEvent, Gradebook, GradebookItem,
                     GradeScoreProof, GradingPolicy, GradingPolicyItem, StudentCourseScore, StudentRiskFlag)
from .repository import GradingRepository

COMPONENTS = ["ATTENDANCE", "ASSIGNMENT", "QUIZ", "LAB", "INTERACTION"]
DEFAULT_WEIGHTS = {"ATTENDANCE": Decimal("10"), "ASSIGNMENT": Decimal("20"), "QUIZ": Decimal("20"), "LAB": Decimal("40"), "INTERACTION": Decimal("10")}
EVENT_TYPES = {
    "attendance.completed": "ATTENDANCE", "assignment.submitted": "ASSIGNMENT", "quiz.completed": "QUIZ",
    "lab.checkpoint.passed": "LAB_CHECKPOINT", "lab.checkpoint.failed": "LAB_CHECKPOINT",
    "lab.submitted": "LAB_SUBMISSION", "poll.completed": "INTERACTION", "grade.manual.adjusted": "MANUAL_ADJUSTMENT",
}
ARCHIVE_FILE_TYPES = {"GRADEBOOK_XLSX", "ANALYTICS_XLSX", "RESOURCE_VERSION_MANIFEST_JSON"}
SHA256_PATTERN = re.compile(r"^[0-9a-fA-F]{64}$")
SOURCE_AGGREGATE_TYPES = {
    "attendance.completed": "attendance_task",
    "assignment.submitted": "assignment",
    "quiz.completed": "quiz",
    "poll.completed": "poll",
    "lab.checkpoint.passed": "runtime_instance",
    "lab.checkpoint.failed": "runtime_instance",
    "lab.submitted": "runtime_instance",
    "grade.manual.adjusted": "grade_event",
}
FREEZE_AGGREGATE_TYPES = {
    "course.roster.frozen": "class_roster",
    "resource.delivery.frozen": "course_resource",
}
SCORE_PROOF_REQUIREMENTS = {
    "assignment.submitted": ("teaching-core", "ASSIGNMENT_FROZEN_QUESTION_SET"),
    "quiz.completed": ("teaching-core", "QUIZ_FROZEN_QUESTION_SET"),
}
SCORE_PROOF_EVENT_TYPE = "grading.score.proof.frozen"
SCORE_PROOF_PRODUCER_ID = "service_teaching_score_prover"
SCORE_PROOF_SOURCE_AGGREGATE_TYPES = {
    "assignment.submitted": "assignment_submission",
    "quiz.completed": "quiz_attempt",
}
VERIFIED_OUTBOX = "VERIFIED_OUTBOX"
VERIFIED_SCORE_PROOF = "VERIFIED_SCORE_PROOF"
SCORE_PROOF_CONTRACT = "grading-score-proof/v1"
SCORE_PROOF_ORIGIN = "SERVER_GRADED"


def utcnow(): return datetime.utcnow()
def number(value): return float(value) if value is not None else None


def canonical_score(value: Decimal) -> str:
    """稳定表示数值，防止 80 和 80.0 绕过同一来源证明绑定。"""

    normalized = value.normalize()
    return format(normalized, "f") if normalized != 0 else "0"


def source_proof_digest(*, event_id: str, event_type: str, aggregate_id: str, payload: dict, proof: dict, raw: Decimal, maximum: Decimal) -> str:
    """冻结成绩证明与事件、范围和分数的可重算绑定摘要。"""

    binding = {
        "aggregate_id": aggregate_id,
        "answer_evidence_sha256": proof["answer_evidence_sha256"],
        "contract": proof["contract"],
        "course_id": payload["course_id"],
        "event_id": event_id,
        "event_type": event_type,
        "evidence_type": proof["evidence_type"],
        "frozen_question_count": proof["frozen_question_count"],
        "frozen_question_sha256": proof["frozen_question_sha256"],
        "issuer": proof["issuer"],
        "max_score": canonical_score(maximum),
        "origin": proof["origin"],
        "raw_score": canonical_score(raw),
        "scoring_evidence_sha256": proof["scoring_evidence_sha256"],
        "source_event_id": proof["source_event_id"],
        "source_fact_id": proof["source_fact_id"],
        "student_id": payload["student_id"],
        "class_id": payload["class_id"],
        "lesson_id": payload.get("lesson_id"),
    }
    return sha256(json.dumps(binding, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False).encode("utf-8")).hexdigest()


def canonical_json(value: object) -> str:
    """用于比较不可变事件箱与消费者收到的事件，不依赖字典插入顺序。"""

    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False)


def normalized_event_time(value: datetime) -> datetime:
    """按 MySQL `DATETIME(0)` 的秒级舍入规则比较事件时间。"""

    if value.tzinfo is not None:
        value = value.astimezone(timezone.utc).replace(tzinfo=None)
    # 当前冻结表使用 DATETIME（无 fsp），MySQL 会将 .5 秒及以上写入下一秒。
    # 消费端必须比对数据库实际持久化表示，不能因合法精度折损误拒绝同一事件箱事实。
    if value.microsecond >= 500_000:
        value += timedelta(seconds=1)
    return value.replace(microsecond=0)


def validate_freeze_contract(event_type: str, aggregate_id: str, payload: dict) -> dict:
    """Validate the frozen cross-domain identifiers F is allowed to retain."""
    if event_type == "course.roster.frozen":
        required = ["course_id", "class_id", "member_count", "snapshot_hash", "frozen_at"]
        missing = [key for key in required if payload.get(key) is None or payload.get(key) == ""]
        if missing:
            raise ApiError("GRADING.UPSTREAM_FREEZE_PAYLOAD_INVALID", "班级名单冻结事件缺少快照字段", 422, {"missing": missing})
        member_count = payload["member_count"]
        if isinstance(member_count, bool) or not isinstance(member_count, int) or member_count < 1:
            raise ApiError("GRADING.ROSTER_SNAPSHOT_INVALID", "班级名单冻结人数必须为正整数", 422, {"member_count": member_count})
        if not SHA256_PATTERN.fullmatch(str(payload["snapshot_hash"])):
            raise ApiError("GRADING.ROSTER_SNAPSHOT_INVALID", "班级名单快照摘要无效", 422, {"snapshot_hash": payload["snapshot_hash"]})
        try:
            datetime.fromisoformat(str(payload["frozen_at"]).replace("Z", "+00:00"))
        except ValueError as exc:
            raise ApiError("GRADING.ROSTER_SNAPSHOT_INVALID", "班级名单冻结时间无效", 422) from exc
        if aggregate_id != payload["class_id"]:
            raise ApiError("GRADING.ROSTER_SNAPSHOT_INVALID", "名单快照标识与班级不一致", 422, {"aggregate_id": aggregate_id, "class_id": payload["class_id"]})
    elif event_type == "resource.delivery.frozen":
        required = ["course_id", "manifest_id", "version_no"]
        missing = [key for key in required if payload.get(key) is None or payload.get(key) == ""]
        if missing:
            raise ApiError("GRADING.UPSTREAM_FREEZE_PAYLOAD_INVALID", "资源冻结事件缺少清单字段", 422, {"missing": missing})
        version_no = payload["version_no"]
        if isinstance(version_no, bool) or not isinstance(version_no, int) or version_no < 1:
            raise ApiError("GRADING.RESOURCE_MANIFEST_INVALID", "资源清单版本必须为正整数", 422, {"version_no": version_no})
        if not isinstance(payload["manifest_id"], str) or not payload["manifest_id"].strip():
            raise ApiError("GRADING.RESOURCE_MANIFEST_INVALID", "资源清单标识无效", 422)
        if aggregate_id != payload["course_id"]:
            raise ApiError("GRADING.RESOURCE_MANIFEST_INVALID", "资源清单标识与课程不一致", 422, {"aggregate_id": aggregate_id, "course_id": payload["course_id"]})
    else:
        raise ApiError("GRADING.EVENT_TYPE_UNSUPPORTED", "该事件不是冻结事实", 422)
    return dict(payload)


def archive_storage_root() -> Path:
    root = (upload_root() / "course_archives").resolve()
    root.mkdir(parents=True, exist_ok=True)
    return root


class GradingService:
    def __init__(self, session: Session, user: UserContext, request_id: str = "system", ip: str = "local"):
        self.session, self.user, self.repo = session, user, GradingRepository(session)
        self.request_id, self.ip = request_id, ip

    def authorize(self, course_id: str, permission: str, class_id: str | None = None, student_id: str | None = None):
        if course_id not in self.user.course_ids and "grading:all-courses" not in self.user.permissions:
            raise ApiError("GRADING.COURSE_SCOPE_DENIED", "无权访问该课程成绩", 403)
        if class_id and class_id not in self.user.class_ids and "grading:all-classes" not in self.user.permissions:
            raise ApiError("GRADING.CLASS_SCOPE_DENIED", "无权访问该班级成绩", 403)
        if student_id and self.user.role == "student" and student_id != self.user.student_id:
            raise ApiError("GRADING.STUDENT_SCOPE_DENIED", "学生只能查看本人数据", 403)
        if permission not in self.user.permissions:
            raise ApiError("AUTH.PERMISSION_DENIED", "缺少成绩操作权限", 403)

    def require_internal_service(self, permission: str):
        if self.user.role != "admin" or not self.user.user_id.startswith("service_") or permission not in self.user.permissions:
            raise ApiError("AUTH.INTERNAL_SERVICE_REQUIRED", "该入口仅允许受信任的内部服务调用", 403)

    def audit(self, action: str, resource_type: str, resource_id: str, *, result="SUCCESS", reason=None, course_id=None, class_id=None, student_id=None, details=None):
        event = AuditEvent(audit_event_id=str(uuid4()), source_event_id=None, actor_user_id=self.user.user_id, actor_role=self.user.role, action=action, resource_type=resource_type, resource_id=resource_id, course_id=course_id, class_id=class_id, student_id=student_id, request_id=self.request_id, ip=self.ip, result=result, reason=reason, occurred_at=utcnow(), details_json=details or {})
        self.repo.add(event); return event

    def ingest_audit(self, data) -> dict:
        self.require_internal_service("audit:ingest")
        limits={"source_event_id":64,"actor_user_id":36,"actor_role":16,"action":64,"resource_type":64,"resource_id":36,"course_id":36,"class_id":36,"student_id":36,"result":16}
        invalid=[field for field,limit in limits.items() if (value:=getattr(data,field)) is not None and (not isinstance(value,str) or not value.strip() or len(value)>limit)]
        if invalid:raise ApiError("AUDIT.EVENT_INVALID","审计事件标识或状态字段无效",422,{"invalid":invalid})
        if data.actor_role not in {"teacher","student","admin","service"}: raise ApiError("AUDIT.ACTOR_ROLE_INVALID", "审计操作者角色无效", 422)
        existing=self.session.scalar(select(AuditEvent).where(AuditEvent.source_event_id==data.source_event_id))
        if existing:return {"status":"DUPLICATE","audit_event_id":existing.audit_event_id}
        row=AuditEvent(audit_event_id=str(uuid4()),source_event_id=data.source_event_id,actor_user_id=data.actor_user_id,actor_role=data.actor_role,action=data.action,resource_type=data.resource_type,resource_id=data.resource_id,course_id=data.course_id,class_id=data.class_id,student_id=data.student_id,request_id=self.request_id,ip=self.ip,result=data.result,reason=data.reason,occurred_at=data.occurred_at.replace(tzinfo=None),details_json=data.details)
        self.repo.add(row)
        try:self.session.commit()
        except IntegrityError:
            self.session.rollback();existing=self.session.scalar(select(AuditEvent).where(AuditEvent.source_event_id==data.source_event_id))
            if existing:return {"status":"DUPLICATE","audit_event_id":existing.audit_event_id}
            raise
        return {"status":"RECORDED","audit_event_id":row.audit_event_id}

    def get_policy(self, course_id: str) -> dict:
        self.authorize(course_id, "grading:read")
        policy = self.repo.policy(course_id)
        if not policy: return {"course_id": course_id, "status": "PENDING", "version_no": None, "items": []}
        return self.policy_dict(policy)

    def put_policy(self, course_id: str, values: dict, effective_at=None) -> dict:
        self.authorize(course_id, "grading:policy")
        weights = {key.upper(): Decimal(str(value)) for key, value in values.items()}
        if set(weights) != set(COMPONENTS) or sum(weights.values()) != Decimal("100"):
            raise ApiError("GRADING.POLICY_WEIGHT_INVALID", "成绩权重合计必须为 100%", 422, {"total": float(sum(weights.values()))})
        latest = self.repo.policy(course_id); version = latest.version_no + 1 if latest else 1
        policy = GradingPolicy(grading_policy_id=str(uuid4()), course_id=course_id, version_no=version, status="PUBLISHED", effective_at=effective_at or utcnow(), created_by=self.user.user_id, created_at=utcnow())
        self.repo.add(policy)
        for component, weight in weights.items(): self.repo.add(GradingPolicyItem(grading_policy_item_id=str(uuid4()), grading_policy_id=policy.grading_policy_id, component=component, weight_percent=weight))
        self.audit("GRADING_POLICY_PUBLISHED", "grading_policy", policy.grading_policy_id, course_id=course_id, details={"version_no": version, "weights": {k: float(v) for k,v in weights.items()}})
        self.session.commit(); return self.policy_dict(policy)

    def ensure_policy(self, course_id: str) -> GradingPolicy:
        policy = self.repo.policy(course_id)
        if policy: return policy
        policy = GradingPolicy(grading_policy_id=str(uuid4()), course_id=course_id, version_no=1, status="PUBLISHED", effective_at=utcnow(), created_by=self.user.user_id, created_at=utcnow())
        self.repo.add(policy)
        for component, weight in DEFAULT_WEIGHTS.items(): self.repo.add(GradingPolicyItem(grading_policy_item_id=str(uuid4()), grading_policy_id=policy.grading_policy_id, component=component, weight_percent=weight))
        self.session.flush(); return policy

    def verify_source_outbox(self, envelope, expected_aggregate_type: str) -> DomainEventOutbox:
        """只接受同一事务事件箱中原封不动投递到 F 的冻结来源。"""

        source = self.session.get(DomainEventOutbox, envelope.event_id)
        if not source:
            raise ApiError(
                "GRADING.SOURCE_EVENT_UNVERIFIED",
                "成绩事件未找到受控事件箱来源",
                422,
                {"reason": "OUTBOX_EVENT_NOT_FOUND"},
            )
        mismatches: list[str] = []
        scalar_fields = ("event_type", "aggregate_type", "aggregate_id", "actor_user_id", "idempotency_key")
        for field in scalar_fields:
            if getattr(source, field) != getattr(envelope, field):
                mismatches.append(field)
        if normalized_event_time(source.occurred_at) != normalized_event_time(envelope.occurred_at):
            mismatches.append("occurred_at")
        try:
            if canonical_json(source.payload_json or {}) != canonical_json(envelope.payload):
                mismatches.append("payload")
        except (TypeError, ValueError):
            mismatches.append("payload")
        if mismatches:
            raise ApiError(
                "GRADING.SOURCE_EVENT_UNVERIFIED",
                "成绩事件与受控事件箱来源不一致",
                422,
                {"reason": "OUTBOX_ENVELOPE_MISMATCH", "fields": sorted(mismatches)},
            )
        if source.aggregate_type != expected_aggregate_type:
            raise ApiError(
                "GRADING.SOURCE_SCOPE_INVALID",
                "成绩事件来源聚合范围不符合冻结契约",
                422,
                {"event_type": envelope.event_type, "expected_aggregate_type": expected_aggregate_type},
            )
        return source

    @staticmethod
    def verify_runtime_evidence(envelope, payload: dict) -> None:
        """D 的运行事件必须继续绑定到实例、发布与服务端状态，而非裸分数。"""

        required = ["lab_release_id", "runtime_instance_id"]
        missing = [field for field in required if not isinstance(payload.get(field), str) or not payload[field].strip()]
        if payload.get("runtime_instance_id") != envelope.aggregate_id:
            missing.append("runtime_instance_id")
        if envelope.event_type == "lab.submitted":
            if payload.get("submission_status") != "SUBMITTED":
                missing.append("submission_status")
        else:
            expected_status = "PASSED" if envelope.event_type == "lab.checkpoint.passed" else "FAILED"
            if not isinstance(payload.get("checkpoint_id"), str) or not payload["checkpoint_id"].strip():
                missing.append("checkpoint_id")
            if payload.get("checkpoint_status") != expected_status:
                missing.append("checkpoint_status")
        if missing:
            raise ApiError(
                "GRADING.SOURCE_EVIDENCE_INVALID",
                "实验成绩事件缺少受控运行证据",
                422,
                {"event_type": envelope.event_type, "invalid": sorted(set(missing))},
            )

    @staticmethod
    def validate_score_proof_fields(*, event_id: str, event_type: str, aggregate_id: str, payload: dict, proof: object, raw: Decimal, maximum: Decimal) -> tuple[str, str]:
        """验证冻结题集、服务端来源标记与分数绑定摘要；调用方必须先证明确实来自持久化证明事实。"""

        expected_issuer, expected_evidence_type = SCORE_PROOF_REQUIREMENTS[event_type]
        if not isinstance(proof, dict):
            raise ApiError("GRADING.SCORE_PROOF_REQUIRED", "评分证明缺少冻结题集和服务端判分摘要", 422, {"event_type": event_type, "missing": ["source_proof"]})
        required = ("contract", "issuer", "origin", "evidence_type", "source_event_id", "source_fact_id", "frozen_question_count", "frozen_question_sha256", "answer_evidence_sha256", "scoring_evidence_sha256", "score_payload_sha256")
        missing = [field for field in required if proof.get(field) is None or proof.get(field) == ""]
        if missing:
            raise ApiError("GRADING.SCORE_PROOF_REQUIRED", "评分证明字段不完整", 422, {"event_type": event_type, "missing": missing})
        invalid: list[str] = []
        if proof.get("contract") != SCORE_PROOF_CONTRACT: invalid.append("contract")
        if proof.get("issuer") != expected_issuer: invalid.append("issuer")
        if proof.get("origin") != SCORE_PROOF_ORIGIN: invalid.append("origin")
        if proof.get("evidence_type") != expected_evidence_type: invalid.append("evidence_type")
        if proof.get("source_event_id") != event_id: invalid.append("source_event_id")
        if proof.get("source_fact_id") != payload.get("source_id"): invalid.append("source_fact_id")
        count = proof.get("frozen_question_count")
        if isinstance(count, bool) or not isinstance(count, int) or not 1 <= count <= 10000: invalid.append("frozen_question_count")
        for field in ("frozen_question_sha256", "answer_evidence_sha256", "scoring_evidence_sha256", "score_payload_sha256"):
            if not isinstance(proof.get(field), str) or not SHA256_PATTERN.fullmatch(proof[field]): invalid.append(field)
        if invalid:
            raise ApiError("GRADING.SCORE_PROOF_INVALID", "评分证明格式或服务端来源标记无效", 422, {"event_type": event_type, "invalid": sorted(set(invalid))})
        digest = source_proof_digest(event_id=event_id, event_type=event_type, aggregate_id=aggregate_id, payload=payload, proof=proof, raw=raw, maximum=maximum)
        if proof["score_payload_sha256"].lower() != digest:
            raise ApiError("GRADING.SCORE_PROOF_INVALID", "评分证明与冻结分数事实不一致", 422, {"event_type": event_type, "invalid": ["score_payload_sha256"]})
        return expected_issuer, digest

    def consume_score_proof(self, envelope) -> dict:
        """只由独立的受控事件箱事件写入 F 的服务端评分证明事实。"""

        payload = envelope.payload
        if envelope.actor_user_id != SCORE_PROOF_PRODUCER_ID:
            self.reject_envelope(
                envelope,
                "GRADING.SCORE_PROOF_PRODUCER_INVALID",
                "评分证明事件不是受信任的服务端判分产物",
                {"expected_actor_user_id": SCORE_PROOF_PRODUCER_ID, "actor_user_id": envelope.actor_user_id},
            )
        score_event_type = payload.get("score_event_type")
        if score_event_type not in SCORE_PROOF_REQUIREMENTS:
            self.reject_envelope(envelope, "GRADING.SCORE_PROOF_EVENT_INVALID", "评分证明未声明受支持的成绩事件类型", {"invalid": ["score_event_type"]})
        required = ("score_event_id", "score_aggregate_id", "source_fact_id", "course_id", "class_id", "student_id", "raw_score", "max_score")
        missing = [field for field in required if payload.get(field) is None or payload.get(field) == ""]
        if missing:
            self.reject_envelope(envelope, "GRADING.SCORE_PROOF_EVENT_INVALID", "评分证明事件缺少冻结成绩字段", {"missing": missing})
        identifiers = ("score_event_id", "score_aggregate_id", "source_fact_id", "course_id", "class_id", "student_id", "lesson_id")
        invalid_ids = [field for field in identifiers if payload.get(field) is not None and (not isinstance(payload[field], str) or not payload[field].strip() or len(payload[field]) > 36)]
        if invalid_ids:
            self.reject_envelope(envelope, "GRADING.SCORE_PROOF_EVENT_INVALID", "评分证明事件标识字段无效", {"invalid": sorted(set(invalid_ids))})
        if envelope.aggregate_id != payload["source_fact_id"]:
            self.reject_envelope(envelope, "GRADING.SOURCE_SCOPE_INVALID", "评分证明聚合标识与提交事实不一致", {"invalid": ["aggregate_id"]})
        try:
            raw, maximum = Decimal(str(payload["raw_score"])), Decimal(str(payload["max_score"]))
        except (InvalidOperation, ValueError):
            self.reject_envelope(envelope, "GRADING.SCORE_PROOF_EVENT_INVALID", "评分证明分数必须为有限数值", {})
        if not raw.is_finite() or not maximum.is_finite() or maximum <= 0 or maximum > Decimal("999999.99") or raw < 0 or raw > maximum:
            self.reject_envelope(envelope, "GRADING.SCORE_PROOF_EVENT_INVALID", "评分证明分数不在有效范围内", {})
        score_payload = {"course_id": payload["course_id"], "class_id": payload["class_id"], "student_id": payload["student_id"], "lesson_id": payload.get("lesson_id"), "source_id": payload["source_fact_id"], "raw_score": payload["raw_score"], "max_score": payload["max_score"]}
        try:
            self.verify_source_outbox(envelope, SCORE_PROOF_SOURCE_AGGREGATE_TYPES[score_event_type])
            issuer, digest = self.validate_score_proof_fields(event_id=payload["score_event_id"], event_type=score_event_type, aggregate_id=payload["score_aggregate_id"], payload=score_payload, proof=payload.get("source_proof"), raw=raw, maximum=maximum)
        except ApiError as exc:
            self.reject_envelope(envelope, exc.code, exc.message, exc.details)
        existing = self.session.get(GradeScoreProof, envelope.event_id)
        if existing:
            return {"status": "DUPLICATE", "source_proof_event_id": envelope.event_id}
        same_score_event = self.session.scalar(select(GradeScoreProof).where(GradeScoreProof.score_event_id == payload["score_event_id"]))
        if same_score_event:
            self.reject_envelope(envelope, "GRADING.SCORE_PROOF_CONFLICT", "同一成绩事件已存在另一份服务端评分证明", {"score_event_id": payload["score_event_id"]})
        proof = payload["source_proof"]
        row = GradeScoreProof(source_proof_event_id=envelope.event_id, score_event_id=payload["score_event_id"], score_event_type=score_event_type, score_aggregate_id=payload["score_aggregate_id"], source_fact_id=payload["source_fact_id"], course_id=payload["course_id"], class_id=payload["class_id"], lesson_id=payload.get("lesson_id"), student_id=payload["student_id"], raw_score=raw, max_score=maximum, contract=proof["contract"], issuer=issuer, origin=proof["origin"], evidence_type=proof["evidence_type"], frozen_question_count=proof["frozen_question_count"], frozen_question_sha256=proof["frozen_question_sha256"], answer_evidence_sha256=proof["answer_evidence_sha256"], scoring_evidence_sha256=proof["scoring_evidence_sha256"], score_payload_sha256=digest, recorded_at=envelope.occurred_at.replace(tzinfo=None))
        self.repo.add(row)
        self.audit("GRADE_SCORE_PROOF_FROZEN", "grade_score_proof", row.source_proof_event_id, course_id=row.course_id, class_id=row.class_id, student_id=row.student_id, details={"score_event_id": row.score_event_id, "score_event_type": row.score_event_type, "source_fact_id": row.source_fact_id, "digest": row.score_payload_sha256})
        try:
            self.session.commit()
        except IntegrityError:
            self.session.rollback()
            if self.session.get(GradeScoreProof, envelope.event_id):
                return {"status": "DUPLICATE", "source_proof_event_id": envelope.event_id}
            raise
        return {"status": "RECORDED", "source_proof_event_id": row.source_proof_event_id, "score_event_id": row.score_event_id}

    def defer_envelope(self, envelope, code: str, message: str, details: dict) -> None:
        """证明尚未投递时留下可审计待重试事实，不能用永久拒绝堵死同批重试。"""

        existing = self.session.scalar(select(AuditEvent).where(AuditEvent.source_event_id == envelope.event_id, AuditEvent.action == "GRADE_EVENT_DEFERRED"))
        if not existing:
            payload = envelope.payload
            self.repo.add(AuditEvent(audit_event_id=str(uuid4()), source_event_id=envelope.event_id, actor_user_id=envelope.actor_user_id, actor_role="service", action="GRADE_EVENT_DEFERRED", resource_type=envelope.aggregate_type, resource_id=envelope.aggregate_id, course_id=payload.get("course_id"), class_id=payload.get("class_id"), student_id=payload.get("student_id"), request_id=self.request_id, ip=self.ip, result="PENDING", reason=message, occurred_at=utcnow(), details_json={"code": code, "details": details, "event_type": envelope.event_type}))
            self.session.commit()
        raise ApiError(code, message, 422, details)

    def verify_score_proof(self, envelope, payload: dict, raw: Decimal, maximum: Decimal) -> tuple[str, str]:
        """分数事件只引用 F 已从独立受控证明事件持久化的来源事实。"""

        proof_event_id = payload.get("score_proof_event_id")
        if not isinstance(proof_event_id, str) or not proof_event_id.strip() or len(proof_event_id) > 36:
            raise ApiError("GRADING.SCORE_PROOF_REQUIRED", "作业或测验缺少独立冻结的服务端评分证明引用", 422, {"event_type": envelope.event_type, "reason": "SERVER_PERSISTED_PROOF_REQUIRED", "missing": ["score_proof_event_id"]})
        stored = self.session.get(GradeScoreProof, proof_event_id)
        if not stored:
            self.defer_envelope(envelope, "GRADING.SERVER_PROOF_PENDING", "作业或测验引用的服务端评分证明尚未入库", {"event_type": envelope.event_type, "score_proof_event_id": proof_event_id})
        expected = {"score_event_id": envelope.event_id, "score_event_type": envelope.event_type, "score_aggregate_id": envelope.aggregate_id, "source_fact_id": payload.get("source_id"), "course_id": payload.get("course_id"), "class_id": payload.get("class_id"), "lesson_id": payload.get("lesson_id"), "student_id": payload.get("student_id")}
        invalid = [field for field, value in expected.items() if getattr(stored, field) != value]
        if stored.raw_score != raw: invalid.append("raw_score")
        if stored.max_score != maximum: invalid.append("max_score")
        if invalid:
            raise ApiError("GRADING.SCORE_PROOF_INVALID", "服务端评分证明与成绩事件范围或分数不一致", 422, {"event_type": envelope.event_type, "invalid": sorted(invalid)})
        proof = {"contract": stored.contract, "issuer": stored.issuer, "origin": stored.origin, "evidence_type": stored.evidence_type, "source_event_id": stored.score_event_id, "source_fact_id": stored.source_fact_id, "frozen_question_count": stored.frozen_question_count, "frozen_question_sha256": stored.frozen_question_sha256, "answer_evidence_sha256": stored.answer_evidence_sha256, "scoring_evidence_sha256": stored.scoring_evidence_sha256, "score_payload_sha256": stored.score_payload_sha256}
        return self.validate_score_proof_fields(event_id=envelope.event_id, event_type=envelope.event_type, aggregate_id=envelope.aggregate_id, payload=payload, proof=proof, raw=raw, maximum=maximum)

    def consume(self, envelope) -> dict:
        self.require_internal_service("grading:consume")
        envelope_limits={"event_id":64,"event_type":64,"aggregate_type":64,"aggregate_id":36,"actor_user_id":36,"idempotency_key":128}
        invalid_envelope=[field for field,limit in envelope_limits.items() if not isinstance((value:=getattr(envelope,field)),str) or not value.strip() or len(value)>limit]
        if invalid_envelope:raise ApiError("GRADING.EVENT_ENVELOPE_INVALID","上游事件信封字段无效",422,{"invalid":invalid_envelope})
        if envelope.event_type == SCORE_PROOF_EVENT_TYPE:
            rejected = self.session.scalar(select(AuditEvent).where(AuditEvent.source_event_id == envelope.event_id, AuditEvent.action == "GRADE_EVENT_REJECTED"))
            if rejected:
                raise ApiError(rejected.details_json["code"], rejected.reason or "评分证明事件已拒绝", 422, rejected.details_json.get("details", {}))
            return self.consume_score_proof(envelope)
        existing = self.session.scalar(select(GradeEvent).where(GradeEvent.event_id == envelope.event_id))
        if existing: return {"status": "DUPLICATE", "grade_event_id": existing.grade_event_id, "event_id": envelope.event_id}
        rejected=self.session.scalar(select(AuditEvent).where(AuditEvent.source_event_id==envelope.event_id,AuditEvent.action=="GRADE_EVENT_REJECTED"))
        if rejected: raise ApiError(rejected.details_json["code"],rejected.reason or "成绩事件已拒绝",422,rejected.details_json.get("details",{}))
        payload = envelope.payload
        if envelope.event_type not in EVENT_TYPES:
            if envelope.event_type in {"resource.delivery.frozen", "course.roster.frozen"}:
                try:
                    self.verify_source_outbox(envelope, FREEZE_AGGREGATE_TYPES[envelope.event_type])
                    frozen_payload = validate_freeze_contract(envelope.event_type, envelope.aggregate_id, payload)
                except ApiError as exc:
                    self.reject_envelope(envelope, exc.code, exc.message, exc.details)
                action = "UPSTREAM.RESOURCE_DELIVERY_FROZEN" if envelope.event_type.startswith("resource") else "UPSTREAM.COURSE_ROSTER_FROZEN"
                existing_audit=self.session.scalar(select(AuditEvent).where(AuditEvent.source_event_id==envelope.event_id))
                if existing_audit:return {"status":"DUPLICATE","event_id":envelope.event_id}
                self.repo.add(AuditEvent(audit_event_id=str(uuid4()),source_event_id=envelope.event_id,actor_user_id=envelope.actor_user_id,actor_role="service",action=action,resource_type=envelope.aggregate_type,resource_id=envelope.aggregate_id,course_id=frozen_payload["course_id"],class_id=frozen_payload.get("class_id"),student_id=None,request_id=self.request_id,ip=self.ip,result="SUCCESS",reason=None,occurred_at=envelope.occurred_at.replace(tzinfo=None),details_json={"payload":frozen_payload}))
                try:
                    self.session.commit()
                except IntegrityError:
                    self.session.rollback()
                    if self.session.scalar(select(AuditEvent).where(AuditEvent.source_event_id == envelope.event_id)):
                        return {"status": "DUPLICATE", "event_id": envelope.event_id}
                    raise
                return {"status": "RECORDED", "event_id": envelope.event_id}
            self.reject_envelope(envelope,"GRADING.EVENT_TYPE_UNSUPPORTED","该事件不属于成绩事实来源",{})
        required = ["course_id", "class_id", "student_id", "source_id", "raw_score", "max_score"]
        missing = [key for key in required if payload.get(key) is None]
        if envelope.event_type == "lab.submitted" and not payload.get("lab_release_id"): missing.append("lab_release_id")
        if missing:self.reject_envelope(envelope,"GRADING.EVENT_PAYLOAD_INVALID","上游事件缺少成绩字段",{"missing":missing})
        invalid_ids = [key for key in ["course_id", "class_id", "student_id", "source_id", "lesson_id"] if payload.get(key) is not None and (not isinstance(payload[key], str) or not payload[key].strip() or len(payload[key]) > 36)]
        if envelope.event_type == "lab.submitted" and (not isinstance(payload.get("lab_release_id"), str) or len(payload["lab_release_id"]) > 36): invalid_ids.append("lab_release_id")
        if invalid_ids:self.reject_envelope(envelope,"GRADING.EVENT_PAYLOAD_INVALID","上游事件标识字段无效",{"invalid":sorted(set(invalid_ids))})
        if len(envelope.event_id) > 36:self.reject_envelope(envelope,"GRADING.EVENT_PAYLOAD_INVALID","成绩事件标识超过存储上限",{"invalid":["event_id"]})
        try:
            raw, maximum = Decimal(str(payload["raw_score"])), Decimal(str(payload["max_score"]))
        except (InvalidOperation, ValueError):
            self.reject_envelope(envelope,"GRADING.SCORE_INVALID","原始分数必须为有限数值",{})
        if not raw.is_finite() or not maximum.is_finite():self.reject_envelope(envelope,"GRADING.SCORE_INVALID","原始分数必须为有限数值",{})
        manual_adjustment=envelope.event_type=="grade.manual.adjusted"
        if maximum <= 0 or maximum > Decimal("999999.99") or raw > maximum or (raw < 0 and (not manual_adjustment or raw < -maximum)):
            self.reject_envelope(envelope,"GRADING.SCORE_INVALID","原始分数必须在有效范围内",{})
        try:
            self.verify_source_outbox(envelope, SOURCE_AGGREGATE_TYPES[envelope.event_type])
            if envelope.event_type.startswith("lab."):
                self.verify_runtime_evidence(envelope, payload)
            if envelope.event_type in SCORE_PROOF_REQUIREMENTS:
                source_proof_issuer, source_proof_digest_value = self.verify_score_proof(envelope, payload, raw, maximum)
                source_verification_status = VERIFIED_SCORE_PROOF
            else:
                source_proof_issuer, source_proof_digest_value = None, None
                source_verification_status = VERIFIED_OUTBOX
        except ApiError as exc:
            if exc.code == "GRADING.SERVER_PROOF_PENDING":
                raise
            self.reject_envelope(envelope, exc.code, exc.message, exc.details)
        # All fact-window transitions use the latest gradebook row as their
        # scope lock. If posting wins, this event becomes LATE; if ingestion
        # wins, posting observes the new fact and rejects a stale snapshot.
        book = self.repo.gradebook(payload["course_id"], payload["class_id"], lock=True)
        archive = self.repo.archive(payload["course_id"], payload["class_id"], lock=True)
        fact_window_closed = bool(book and book.status in {"POSTED", "LOCKED"}) or bool(archive and archive.status == "ARCHIVED")
        grade = GradeEvent(grade_event_id=str(uuid4()), course_id=payload["course_id"], class_id=payload["class_id"], lesson_id=payload.get("lesson_id"), student_id=payload["student_id"], source_type=EVENT_TYPES[envelope.event_type], source_id=payload["source_id"], raw_score=raw, max_score=maximum, normalized_score=(raw/maximum*100).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP), occurred_at=envelope.occurred_at.replace(tzinfo=None), event_id=envelope.event_id, status="LATE" if fact_window_closed else "CONSUMED", source_verification_status=source_verification_status, source_proof_issuer=source_proof_issuer, source_proof_digest=source_proof_digest_value, payload_json={**payload,"source_event_type":envelope.event_type,"source_occurred_at":envelope.occurred_at.isoformat()})
        self.repo.add(grade)
        if fact_window_closed:
            self.audit("GRADE_EVENT_LATE_IGNORED", "grade_event", grade.grade_event_id, result="BLOCKED", reason="成绩事实窗口已关闭", course_id=grade.course_id, class_id=grade.class_id, student_id=grade.student_id, details={"source_id": grade.source_id, "event_id": grade.event_id, "gradebook_status": book.status if book else None, "archive_status": archive.status if archive else None})
            self.session.commit(); return {"status": "LATE_IGNORED", "grade_event_id": grade.grade_event_id, "event_id": envelope.event_id}
        if grade.source_type == "MANUAL_ADJUSTMENT": self.audit("GRADE_MANUAL_ADJUSTMENT", "grade_event", grade.grade_event_id, course_id=grade.course_id, class_id=grade.class_id, student_id=grade.student_id, details={"source_id": grade.source_id, "normalized_score": number(grade.normalized_score)})
        enqueue_event(self.session, event_type="grade.event.created", aggregate_type="grade_event", aggregate_id=grade.grade_event_id, actor_user_id=self.user.user_id, idempotency_key=f"grade-event:{envelope.event_id}", payload={"course_id":grade.course_id,"class_id":grade.class_id,"student_id":grade.student_id,"source_type":grade.source_type})
        try:
            self.session.commit()
        except IntegrityError:
            self.session.rollback()
            existing = self.session.scalar(select(GradeEvent).where(GradeEvent.event_id == envelope.event_id))
            if existing:
                return {"status": "DUPLICATE", "grade_event_id": existing.grade_event_id, "event_id": envelope.event_id}
            raise
        return {"status": "CONSUMED", "grade_event_id": grade.grade_event_id, "event_id": envelope.event_id}

    def reject_envelope(self,envelope,code,message,details):
        payload=envelope.payload
        self.repo.add(AuditEvent(audit_event_id=str(uuid4()),source_event_id=envelope.event_id,actor_user_id=envelope.actor_user_id,actor_role="service",action="GRADE_EVENT_REJECTED",resource_type=envelope.aggregate_type,resource_id=envelope.aggregate_id,course_id=payload.get("course_id"),class_id=payload.get("class_id"),student_id=payload.get("student_id"),request_id=self.request_id,ip=self.ip,result="ERROR",reason=message,occurred_at=utcnow(),details_json={"code":code,"details":details,"event_type":envelope.event_type}))
        try:
            self.session.commit()
        except IntegrityError:
            self.session.rollback()
            rejected=self.session.scalar(select(AuditEvent).where(AuditEvent.source_event_id==envelope.event_id,AuditEvent.action=="GRADE_EVENT_REJECTED"))
            if rejected:
                raise ApiError(rejected.details_json["code"],rejected.reason or message,422,rejected.details_json.get("details",{}))
            raise
        raise ApiError(code,message,422,details)

    def recalculate(self, course_id: str, class_id: str) -> dict:
        self.authorize(course_id, "grading:recalculate", class_id)
        # Recalculation participates in the same gradebook -> archive lock
        # order as ingestion and archival, so it cannot overwrite POSTED or
        # LOCKED state after reading an earlier snapshot.
        book = self.repo.gradebook(course_id, class_id, lock=True)
        archived = self.repo.archive(course_id, class_id, lock=True)
        if archived and archived.status == "ARCHIVED": raise ApiError("ARCHIVE.GRADEBOOK_LOCKED", "课程已归档，不能静默重算成绩", 409)
        policy = self.ensure_policy(course_id); weights = {x.component: Decimal(x.weight_percent) for x in self.repo.policy_items(policy.grading_policy_id)}
        enabled_components = {component for component, weight in weights.items() if weight > 0}
        events = self.repo.grade_events(course_id, class_id)
        if not events: raise ApiError("GRADING.SOURCE_FACTS_PENDING", "尚未收到 A/D/E 上游成绩事实", 409)
        if book and book.policy_version == policy.version_no and book.status in {"POSTED","LOCKED"}: raise ApiError("GRADING.GRADEBOOK_LOCKED", "成绩已入账或归档；请发布新规则版本后重算", 409)
        if book and book.policy_version == policy.version_no:
            for model in [StudentRiskFlag, AnalyticsLabSummary, AnalyticsStudentLabSummary, AnalyticsSectionSummary, AnalyticsCourseSummary, StudentCourseScore, GradebookItem]: self.session.execute(delete(model).where(model.gradebook_id == book.gradebook_id))
            book.status, book.calculated_at = "CALCULATING", utcnow()
        else:
            book = Gradebook(gradebook_id=str(uuid4()), course_id=course_id, class_id=class_id, policy_version=policy.version_no, status="CALCULATING", calculated_at=utcnow(), posted_at=None, locked_at=None); self.repo.add(book); self.session.flush()
        by_student = defaultdict(list)
        for event in events: by_student[event.student_id].append(event)
        rankings=[]
        for student_id, student_events in by_student.items():
            grouped=self.authoritative_component_scores(student_events)
            total=Decimal("0"); present=[]
            for component in COMPONENTS:
                scores=grouped.get(component,[]); score=(sum(scores)/len(scores)).quantize(Decimal("0.01")) if scores else Decimal("0")
                if scores: present.append(component)
                weighted=(score*weights[component]/100).quantize(Decimal("0.01")); total+=weighted
                self.repo.add(GradebookItem(gradebook_item_id=str(uuid4()),gradebook_id=book.gradebook_id,student_id=student_id,component=component,component_score=score,weight_percent=weights[component],weighted_score=weighted))
            adjustments=sum((Decimal(e.normalized_score) for e in student_events if e.source_type=="MANUAL_ADJUSTMENT"),Decimal("0"))
            total=max(Decimal("0"),min(Decimal("100"),total+adjustments)).quantize(Decimal("0.01")); completeness="READY" if enabled_components.issubset(present) else "PARTIAL"
            self.repo.add(StudentCourseScore(student_course_score_id=str(uuid4()),gradebook_id=book.gradebook_id,course_id=course_id,class_id=class_id,student_id=student_id,total_score=total,completeness=completeness));rankings.append({"student_id":student_id,"total_score":float(total),"completeness":completeness})
            self.build_risks(book, student_id, grouped, total, present, student_events, events)
        rankings.sort(key=lambda x:x["total_score"],reverse=True)
        component_avg={c: self.average([next((i.component_score for i in self.repo.items(book.gradebook_id,s) if i.component==c),Decimal("0")) for s in by_student]) for c in COMPONENTS}
        summary={"status":"READY" if all(x["completeness"]=="READY" for x in rankings) else "PARTIAL","student_count":len(rankings),"avg_assignment":component_avg["ASSIGNMENT"],"avg_quiz":component_avg["QUIZ"],"attendance_rate":component_avg["ATTENDANCE"],"course_average":self.average([Decimal(str(x["total_score"])) for x in rankings]),"ranking":rankings,"source_event_ids":sorted(x.event_id for x in events)}
        self.repo.add(AnalyticsCourseSummary(analytics_course_summary_id=str(uuid4()),gradebook_id=book.gradebook_id,course_id=course_id,class_id=class_id,summary_json=summary,calculated_at=utcnow()))
        self.build_section_analytics(book,events); self.build_lab_analytics(book,events,len(by_student)); book.status="READY"; book.calculated_at=utcnow()
        self.audit("GRADEBOOK_RECALCULATED","gradebook",book.gradebook_id,course_id=course_id,class_id=class_id,details={"policy_version":policy.version_no,"source_event_count":len(events),"status":summary["status"]})
        self.session.commit(); return self.gradebook_dict(book)

    def freeze_fact(self, course_id: str, action: str, class_id: str | None = None) -> dict | None:
        event_type = "course.roster.frozen" if action == "UPSTREAM.COURSE_ROSTER_FROZEN" else "resource.delivery.frozen"
        for row in self.repo.audit_events(course_id=course_id, action=action, result="SUCCESS"):
            if not row.source_event_id or (class_id is not None and row.class_id != class_id):
                continue
            payload = (row.details_json or {}).get("payload") or {}
            try:
                validate_freeze_contract(event_type, row.resource_id, payload)
            except ApiError:
                continue
            return {
                "source_event_id": row.source_event_id,
                "aggregate_id": row.resource_id,
                "occurred_at": row.occurred_at.isoformat() + "Z",
                **payload,
            }
        return None

    def gradebook_integrity(self, book: Gradebook | None) -> dict:
        if not book:
            return {"student_count": 0, "enabled_components": [], "incomplete_students": [], "source_snapshot_present": False, "source_event_ids": [], "unapplied_source_event_ids": [], "missing_source_event_ids": []}
        policy = self.repo.policy_version(book.course_id, book.policy_version)
        weights = {item.component: Decimal(item.weight_percent) for item in self.repo.policy_items(policy.grading_policy_id)} if policy else {}
        enabled = sorted(component for component, weight in weights.items() if weight > 0)
        current_events = self.repo.grade_events(book.course_id, book.class_id)
        summary = self.repo.course_summary(book.gradebook_id)
        raw_snapshot = (summary.summary_json or {}).get("source_event_ids") if summary else None
        snapshot_present = isinstance(raw_snapshot, list) and all(isinstance(item, str) for item in raw_snapshot)
        snapshot_ids = set(raw_snapshot or []) if snapshot_present else set()
        current_ids = {event.event_id for event in current_events}
        snapshot_events = [event for event in current_events if event.event_id in snapshot_ids]
        by_student: dict[str, list] = defaultdict(list)
        for event in snapshot_events:
            by_student[event.student_id].append(event)
        scores = self.repo.scores(book.gradebook_id)
        incomplete = []
        for score in scores:
            present = set(self.authoritative_component_scores(by_student.get(score.student_id, [])).keys())
            missing = sorted(set(enabled) - present)
            if missing or score.completeness != "READY":
                incomplete.append({"student_id": score.student_id, "missing_components": missing})
        return {
            "student_count": len(scores),
            "enabled_components": enabled,
            "incomplete_students": incomplete,
            "source_snapshot_present": snapshot_present,
            "source_event_ids": sorted(snapshot_ids),
            "unapplied_source_event_ids": sorted(current_ids - snapshot_ids),
            "missing_source_event_ids": sorted(snapshot_ids - current_ids),
        }

    def post(self, course_id: str, class_id: str) -> dict:
        self.authorize(course_id,"grading:post",class_id)
        book=self.repo.gradebook(course_id,class_id,lock=True)
        if not book or book.status not in {"READY", "POSTED"}: raise ApiError("GRADING.GRADEBOOK_NOT_READY","只有已重算成绩册可以入账",409)
        integrity = self.gradebook_integrity(book)
        if not integrity["source_snapshot_present"] or integrity["unapplied_source_event_ids"] or integrity["missing_source_event_ids"]:
            raise ApiError("GRADING.GRADEBOOK_STALE", "成绩事实已变化或缺少计算快照，请重新计算后再入账", 409, integrity)
        if not integrity["student_count"] or integrity["incomplete_students"]:
            raise ApiError("GRADING.GRADEBOOK_INCOMPLETE", "存在缺少启用权重分量的学生，不能入账", 409, integrity)
        roster = self.freeze_fact(course_id, "UPSTREAM.COURSE_ROSTER_FROZEN", class_id)
        if not roster:
            raise ApiError("GRADING.ROSTER_FREEZE_REQUIRED", "班级名单尚未形成有效冻结快照，不能入账", 409)
        if roster["member_count"] != integrity["student_count"]:
            raise ApiError("GRADING.ROSTER_COUNT_MISMATCH", "成绩册学生数与冻结名单人数不一致，不能入账", 409, {"roster_member_count": roster["member_count"], "gradebook_student_count": integrity["student_count"], "snapshot_hash": roster["snapshot_hash"]})
        if book.status=="POSTED": return self.gradebook_dict(book)
        book.status="POSTED";book.posted_at=utcnow();self.audit("GRADEBOOK_POSTED","gradebook",book.gradebook_id,course_id=course_id,class_id=book.class_id)
        enqueue_event(self.session,event_type="gradebook.posted",aggregate_type="gradebook",aggregate_id=book.gradebook_id,actor_user_id=self.user.user_id,idempotency_key=f"gradebook-posted:{book.gradebook_id}",payload={"course_id":course_id,"class_id":book.class_id,"policy_version":book.policy_version})
        self.session.commit();return self.gradebook_dict(book)

    def gradebook(self, course_id: str, class_id: str, student_id: str | None = None) -> dict:
        if self.user.role == "student":
            if student_id and student_id != self.user.student_id: raise ApiError("GRADING.STUDENT_SCOPE_DENIED", "学生只能查看本人数据", 403)
            student_id = self.user.student_id
        self.authorize(course_id,"grading:read",class_id,student_id)
        book=self.repo.gradebook(course_id,class_id)
        if not book:return {"course_id":course_id,"status":"PENDING","items":[]}
        if self.user.role == "student" and book.status not in {"POSTED", "LOCKED"}:
            return {"course_id":course_id,"class_id":class_id,"status":"PENDING","reason":"成绩尚未正式入账","items":[]}
        scores=self.repo.scores(book.gradebook_id)
        if student_id:scores=[x for x in scores if x.student_id==student_id]
        return {**self.gradebook_dict(book),"items":[{"student_id":x.student_id,"total_score":number(x.total_score),"completeness":x.completeness} for x in scores]}

    def trace(self, course_id: str, class_id: str, student_id: str) -> dict:
        self.authorize(course_id,"grading:read",class_id,student_id=student_id);book=self.repo.gradebook(course_id,class_id)
        if not book:return {"status":"PENDING","student_id":student_id,"sources":[]}
        if self.user.role == "student" and book.status not in {"POSTED", "LOCKED"}:return {"status":"PENDING","student_id":student_id,"sources":[]}
        score=next((x for x in self.repo.scores(book.gradebook_id) if x.student_id==student_id),None)
        summary=self.repo.course_summary(book.gradebook_id);snapshot_ids=set((summary.summary_json or {}).get("source_event_ids",[])) if summary else set()
        events=[event for event in self.repo.grade_events(course_id,class_id,student_id) if event.event_id in snapshot_ids];items=self.repo.items(book.gradebook_id,student_id)
        components=[{"component":x.component,"score":number(x.component_score),"weight_percent":number(x.weight_percent),"weighted_score":number(x.weighted_score),"sources":[self.event_dict(e) for e in events if self.component(e.source_type)==x.component]} for x in items]
        adjustments=[event for event in events if event.source_type=="MANUAL_ADJUSTMENT"]
        if adjustments:
            adjustment=sum((Decimal(event.normalized_score) for event in adjustments),Decimal("0"))
            components.append({"component":"MANUAL_ADJUSTMENT","score":number(adjustment),"weight_percent":None,"weighted_score":number(adjustment),"sources":[self.event_dict(event) for event in adjustments]})
        return {"status":book.status,"student_id":student_id,"total_score":number(score.total_score) if score else None,"policy_version":book.policy_version,"source_snapshot_event_ids":sorted(snapshot_ids),"components":components}

    def overview(self,course_id,class_id):
        self.authorize(course_id,"analytics:class",class_id);book=self.repo.gradebook(course_id,class_id)
        if not book:return {"course_id":course_id,"status":"PENDING","reason":"等待上游成绩事实与重算"}
        row=self.repo.course_summary(book.gradebook_id);return row.summary_json if row else {"course_id":course_id,"status":"PENDING"}
    def section(self,course_id,class_id,lesson_id):
        self.authorize(course_id,"analytics:class",class_id);book=self.repo.gradebook(course_id,class_id);row=self.repo.section_summary(book.gradebook_id,lesson_id) if book else None
        return row.summary_json if row else {"course_id":course_id,"lesson_id":lesson_id,"status":"PENDING"}
    def labs_by_student(self,course_id,class_id):
        self.authorize(course_id,"analytics:class",class_id);book=self.repo.gradebook(course_id,class_id)
        return {"status":"PENDING","items":[]} if not book else {"status":"READY","items":[{"student_id":x.student_id,"sum_lab_score":number(x.sum_lab_score),"submitted_count":x.submitted_count,"unsubmitted_count":x.unsubmitted_count} for x in self.repo.student_lab_summaries(book.gradebook_id)]}
    def labs_by_lab(self,course_id,class_id):
        self.authorize(course_id,"analytics:class",class_id);book=self.repo.gradebook(course_id,class_id)
        return {"status":"PENDING","items":[]} if not book else {"status":"READY","items":[{"lab_release_id":x.lab_release_id,"max_score":number(x.max_score),"submitted_students":x.submitted_students,"unsubmitted_students":x.unsubmitted_students,"avg_score":number(x.avg_score)} for x in self.repo.lab_summaries(book.gradebook_id)]}
    def risks(self,course_id,class_id,student_id=None):
        self.authorize(course_id,"analytics:read",class_id,student_id=student_id);rows=self.repo.risks(course_id,class_id,student_id)
        return {"status":"READY" if self.repo.gradebook(course_id,class_id) else "PENDING","items":[{"student_id":x.student_id,"risk_type":x.risk_type,"evidence":x.evidence_json,"status":x.status} for x in rows]}

    def learning_summary(self,course_id,class_id,student_id=None):
        if self.user.role == "student":student_id=self.user.student_id
        self.authorize(course_id,"analytics:read",class_id,student_id=student_id);gradebook_row=self.repo.gradebook(course_id,class_id)
        summary_row=self.repo.course_summary(gradebook_row.gradebook_id) if gradebook_row else None;overview=dict(summary_row.summary_json) if summary_row else {"status":"PENDING"}
        if student_id and overview.get("ranking"):
            rank=next((i+1 for i,x in enumerate(overview["ranking"]) if x["student_id"]==student_id),None);overview.pop("ranking",None);overview["student_rank"]=rank
        book=self.gradebook(course_id,class_id,student_id);risks=self.risks(course_id,class_id,student_id)
        states=[overview.get("status","PENDING"),book.get("status","PENDING"),risks.get("status","PENDING")]
        status="PENDING" if all(x=="PENDING" for x in states) else ("READY" if overview.get("status")=="READY" and book.get("status") in {"POSTED","LOCKED"} else "PARTIAL")
        return {"course_id":course_id,"class_id":class_id,"student_id":student_id,"status":status,"grade":book,"risk":risks,"analytics":overview,"missing_upstream":[] if status=="READY" else ["上游班级成员/完整教学事实或成绩入账尚未就绪"]}

    def archive_precheck_data(self, course_id: str, class_id: str, book: Gradebook | None) -> dict:
        events=self.repo.grade_events(course_id,class_id)
        roster=self.freeze_fact(course_id,"UPSTREAM.COURSE_ROSTER_FROZEN",class_id)
        resource_manifest=self.freeze_fact(course_id,"UPSTREAM.RESOURCE_DELIVERY_FROZEN")
        integrity=self.gradebook_integrity(book)
        snapshot_ids=set(integrity["source_event_ids"])
        snapshot_events=[event for event in events if event.event_id in snapshot_ids]
        checks={
            "roster_frozen":bool(roster),
            "roster_student_count_matches":bool(roster and roster["member_count"]==integrity["student_count"]),
            "gradebook_posted":bool(book and book.status=="POSTED"),
            "gradebook_complete":bool(integrity["student_count"] and not integrity["incomplete_students"]),
            "grade_event_traceable":bool(snapshot_events) and not integrity["unapplied_source_event_ids"] and not integrity["missing_source_event_ids"] and all(x.event_id and x.source_id for x in snapshot_events),
            "lab_evidence_referenced":any(x.source_type in {"LAB_CHECKPOINT","LAB_SUBMISSION"} for x in snapshot_events),
            "resources_frozen":bool(resource_manifest),
        }
        blockers=[key for key,value in checks.items() if not value];data={"checks":checks,"blocking":len(blockers),"blocking_items":blockers,"status":"READY" if not blockers else "PRECHECK","roster_snapshot":roster,"resource_manifest":resource_manifest,"gradebook_integrity":integrity}
        return data

    def precheck(self,course_id,class_id):
        self.authorize(course_id,"archives:write",class_id);book=self.repo.gradebook(course_id,class_id,lock=True)
        data=self.archive_precheck_data(course_id,class_id,book)
        blockers=data["blocking_items"]
        archive=self.repo.archive(course_id,class_id,lock=True)
        if archive and archive.status=="ARCHIVED": return {**archive.precheck_json,"status":"ARCHIVED"}
        if archive: archive.precheck_json,archive.status,archive.gradebook_id=data,"READY" if not blockers else "PRECHECK",book.gradebook_id if book else archive.gradebook_id
        else: archive=CourseArchive(course_archive_id=str(uuid4()),course_id=course_id,class_id=class_id,gradebook_id=book.gradebook_id if book else None,status="READY" if not blockers else "PRECHECK",precheck_json=data,manifest_json=None,created_by=self.user.user_id,created_at=utcnow(),archived_at=None);self.repo.add(archive)
        self.audit("COURSE_ARCHIVE_PRECHECK","course_archive",archive.course_archive_id,result="SUCCESS" if not blockers else "BLOCKED",reason=",".join(blockers) or None,course_id=course_id,class_id=class_id,details=data)
        self.session.commit();return data

    def freeze_archive(self,course_id,class_id):
        self.authorize(course_id,"archives:freeze",class_id);existing=self.repo.archive(course_id,class_id)
        if existing and existing.status=="ARCHIVED":return existing.manifest_json
        check=self.precheck(course_id,class_id)
        if check["blocking"]:raise ApiError("ARCHIVE.PRECHECK_BLOCKED","归档门禁尚有阻断项",409,check)
        # Lock in the same order as event ingestion so no score fact can cross
        # the final precheck/archival boundary unnoticed.
        book=self.repo.gradebook(course_id,class_id,lock=True)
        archive=self.repo.archive(course_id,class_id,lock=True)
        if archive.status=="ARCHIVED":return archive.manifest_json
        check=self.archive_precheck_data(course_id,class_id,book)
        if check["blocking"]:raise ApiError("ARCHIVE.PRECHECK_BLOCKED","归档门禁在冻结前发生变化",409,check)
        book.status="LOCKED";book.locked_at=utcnow()
        resource_snapshot = {
            "schema_version": 1,
            "course_id": course_id,
            "class_id": class_id,
            "delivery_manifests": [check["resource_manifest"]],
        }
        contents = [
            ("GRADEBOOK_XLSX", self.gradebook_xlsx(course_id,class_id), f"{course_id}-成绩册.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
            ("ANALYTICS_XLSX", self.analytics_xlsx(course_id,class_id), f"{course_id}-学情.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
            ("RESOURCE_VERSION_MANIFEST_JSON", json.dumps(resource_snapshot, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8"), f"{course_id}-资源版本清单.json", "application/json"),
        ]
        artifacts=[];created_paths=[]
        try:
            for kind,content,name,mime_type in contents:
                artifact, created_path = self.persist_archive_bytes(archive, kind, content, name, mime_type)
                artifacts.append(artifact)
                if created_path: created_paths.append(created_path)
            snapshot_ids=set(check["gradebook_integrity"]["source_event_ids"])
            source_events=[event for event in self.repo.grade_events(course_id,class_id) if event.event_id in snapshot_ids]
            evidence_ids=sorted(x.source_id for x in source_events if x.source_type in {"LAB_CHECKPOINT","LAB_SUBMISSION"})
            evidence_digest=sha256("\n".join(evidence_ids).encode()).hexdigest();evidence_ref="contract://lab-evidence/"+class_id
            self.repo.add(CourseArchiveArtifact(course_archive_artifact_id=str(uuid4()),course_archive_id=archive.course_archive_id,artifact_type="LAB_EVIDENCE_REFERENCES",file_id=None,evidence_ref=evidence_ref,sha256=evidence_digest));artifacts.append({"type":"LAB_EVIDENCE_REFERENCES","evidence_ref":evidence_ref,"sha256":evidence_digest})
            archive.status="ARCHIVED";archive.archived_at=utcnow();archive.manifest_json={"course_id":course_id,"class_id":class_id,"gradebook_id":book.gradebook_id,"policy_version":book.policy_version,"roster_snapshot":check["roster_snapshot"],"resource_manifest":check["resource_manifest"],"source_event_ids":sorted(snapshot_ids),"artifacts":artifacts,"archived_at":archive.archived_at.isoformat()+"Z"}
            self.audit("COURSE_ARCHIVED","course_archive",archive.course_archive_id,course_id=course_id,class_id=class_id,details=archive.manifest_json)
            enqueue_event(self.session,event_type="course.archived",aggregate_type="course_archive",aggregate_id=archive.course_archive_id,actor_user_id=self.user.user_id,idempotency_key=f"course-archived:{archive.course_archive_id}",payload={"course_id":course_id,"class_id":class_id,"gradebook_id":book.gradebook_id})
            self.session.commit();return archive.manifest_json
        except Exception:
            self.session.rollback()
            for path in created_paths:
                try:
                    path.chmod(S_IRUSR | S_IRGRP | S_IROTH | 0o200)
                    path.unlink(missing_ok=True)
                except OSError:
                    pass
            raise

    def persist_archive_bytes(self, archive: CourseArchive, artifact_type: str, content: bytes, original_name: str, mime_type: str) -> tuple[dict, Path | None]:
        if artifact_type not in ARCHIVE_FILE_TYPES or not content:
            raise ApiError("ARCHIVE.ARTIFACT_INVALID", "归档产物类型或内容无效", 422)
        digest=sha256(content).hexdigest();root=archive_storage_root();path=(root/"objects"/digest[:2]/digest).resolve()
        if root not in path.parents:
            raise ApiError("ARCHIVE.STORAGE_PATH_INVALID", "归档存储路径越界", 500)
        path.parent.mkdir(parents=True,exist_ok=True);created=False
        try:
            with path.open("xb") as target:
                target.write(content);target.flush();os.fsync(target.fileno())
            created=True;path.chmod(S_IRUSR|S_IRGRP|S_IROTH)
        except FileExistsError:
            stored=path.read_bytes()
            if len(stored)!=len(content) or sha256(stored).hexdigest()!=digest:
                raise ApiError("ARCHIVE.ARTIFACT_COLLISION", "归档内容寻址对象校验失败", 500)
        file_object=self.session.scalar(select(FileObject).where(FileObject.sha256==digest,FileObject.size_bytes==len(content)))
        if file_object:
            registered=Path(file_object.object_key).resolve()
            if file_object.storage_provider!="local" or file_object.bucket!="course-archives" or registered!=path:
                if created:
                    path.chmod(S_IRUSR|0o200);path.unlink(missing_ok=True)
                raise ApiError("ARCHIVE.FILE_CONTENT_CONFLICT", "相同内容已由其他受控存储登记", 409)
        else:
            file_object=FileObject(file_id=str(uuid4()),storage_provider="local",bucket="course-archives",object_key=str(path),original_name=original_name,mime_type=mime_type,size_bytes=len(content),sha256=digest,created_by=self.user.user_id,created_at=utcnow());self.repo.add(file_object)
        ref=f"/api/v1/archives/courses/{archive.course_id}/artifacts/{artifact_type}?class_id={archive.class_id}"
        artifact=CourseArchiveArtifact(course_archive_artifact_id=str(uuid4()),course_archive_id=archive.course_archive_id,artifact_type=artifact_type,file_id=file_object.file_id,evidence_ref=ref,sha256=digest);self.repo.add(artifact)
        return {"type":artifact_type,"file_id":file_object.file_id,"evidence_ref":ref,"sha256":digest,"size_bytes":len(content)}, path if created else None

    def artifact_download(self,course_id,class_id,artifact_type):
        self.authorize(course_id,"archives:read",class_id);archive=self.repo.archive(course_id,class_id)
        if not archive or archive.status!="ARCHIVED":raise ApiError("ARCHIVE.NOT_READY","课程尚未归档",409)
        if artifact_type not in ARCHIVE_FILE_TYPES:raise ApiError("ARCHIVE.ARTIFACT_NOT_FOUND","归档产物不存在",404)
        artifact=self.repo.archive_artifact(archive.course_archive_id,artifact_type)
        file_object=self.repo.file_object(artifact.file_id) if artifact and artifact.file_id else None
        if not artifact or not file_object or file_object.storage_provider!="local" or file_object.bucket!="course-archives":raise ApiError("ARCHIVE.ARTIFACT_NOT_AVAILABLE","归档产物不可用",404)
        root=archive_storage_root();path=Path(file_object.object_key).resolve()
        if root not in path.parents or not path.is_file():raise ApiError("ARCHIVE.ARTIFACT_NOT_AVAILABLE","归档产物已移出受控存储",404)
        content=path.read_bytes();digest=sha256(content).hexdigest()
        if len(content)!=file_object.size_bytes or digest!=file_object.sha256 or digest!=artifact.sha256:raise ApiError("ARCHIVE.ARTIFACT_INTEGRITY_FAILED","归档产物完整性校验失败",409)
        return content,file_object.original_name,file_object.mime_type,digest

    def archive(self,course_id,class_id):
        self.authorize(course_id,"archives:read",class_id);row=self.repo.archive(course_id,class_id)
        return {"status":"PENDING"} if not row else {"status":row.status,"precheck":row.precheck_json,"manifest":row.manifest_json}
    def audit_list(self,course_id=None,action=None,result=None):
        if "audit:read" not in self.user.permissions:raise ApiError("AUTH.PERMISSION_DENIED","缺少审计读取权限",403)
        unrestricted="grading:all-courses" in self.user.permissions
        if course_id and not unrestricted and course_id not in self.user.course_ids:raise ApiError("AUDIT.COURSE_SCOPE_DENIED","无权读取该课程审计",403)
        course_ids=None if unrestricted else sorted(self.user.course_ids)
        rows=self.repo.audit_events(course_id=course_id,course_ids=course_ids,action=action,result=result);return {"items":[self.audit_dict(x) for x in rows],"page":1,"page_size":len(rows),"total":len(rows)}

    def gradebook_xlsx(self,course_id,class_id):
        data=self.gradebook(course_id,class_id);book=self.repo.gradebook(course_id,class_id)
        component_labels={"ATTENDANCE":"签到","ASSIGNMENT":"作业","QUIZ":"测验","LAB":"实验","INTERACTION":"互动"};rows=[]
        for score in data.get("items",[]):
            values={item.component:number(item.component_score) for item in self.repo.items(book.gradebook_id,score["student_id"])} if book else {}
            trace=self.trace(course_id,class_id,score["student_id"]) if book else {"components":[]}
            adjustment=next((item["score"] for item in trace.get("components",[]) if item["component"]=="MANUAL_ADJUSTMENT"),0)
            rows.append([score["student_id"],*[values.get(component,0) for component in COMPONENTS],adjustment,score["total_score"],score["completeness"],book.policy_version if book else None,book.status if book else "PENDING"])
        return self.make_xlsx("成绩册",["学生标识",*[component_labels[x] for x in COMPONENTS],"人工调整","总评","完整性","规则版本","成绩状态"],rows,f"{course_id}/{class_id}")
    def analytics_xlsx(self,course_id,class_id):
        overview=self.overview(course_id,class_id);gradebook=self.repo.gradebook(course_id,class_id);book=Workbook();sheet=book.active;sheet.title="课程总览";sheet.append(["课程标识",course_id]);sheet.append(["班级标识",class_id]);sheet.append(["生成时间",utcnow().isoformat()+"Z"]);sheet.append(["作业平均",overview.get("avg_assignment")]);sheet.append(["测验平均",overview.get("avg_quiz")]);sheet.append(["平均到课率",overview.get("attendance_rate")]);sheet.append(["课程平均",overview.get("course_average")])
        section_items=[]
        if gradebook:
            for row in self.repo.section_summaries(gradebook.gradebook_id):
                summary=row.summary_json
                for rank in summary.get("ranking",[]):section_items.append([row.lesson_id,summary.get("average"),rank.get("student_id"),rank.get("score")])
        worksheets=[("课程排行",["名次","学生标识","总评","完整性"],[[index,item["student_id"],item["total_score"],item["completeness"]] for index,item in enumerate(overview.get("ranking",[]),1)]),("小节成绩",["课时","平均成绩","学生标识","学生成绩"],section_items),("学生实验完成",["学生标识","实验总分","已提交","未提交"],[[x["student_id"],x["sum_lab_score"],x["submitted_count"],x["unsubmitted_count"]] for x in self.labs_by_student(course_id,class_id)["items"]]),("实验完成统计",["实验发布标识","满分","已提交人数","未提交人数","平均分"],[[x["lab_release_id"],x["max_score"],x["submitted_students"],x["unsubmitted_students"],x["avg_score"]] for x in self.labs_by_lab(course_id,class_id)["items"]]),("风险依据",["学生标识","风险类型","事实依据"],[[x["student_id"],x["risk_type"],json.dumps(x["evidence"],ensure_ascii=False,sort_keys=True)] for x in self.risks(course_id,class_id)["items"]])]
        for title,headers,items in worksheets:
            ws=book.create_sheet(title);ws.append(["课程标识",course_id]);ws.append(["班级标识",class_id]);ws.append(["生成时间",utcnow().isoformat()+"Z"]);ws.append([]);ws.append(headers)
            for x in items:ws.append(x)
        output=BytesIO();book.save(output);return output.getvalue()
    def audit_xlsx(self,course_id=None):
        rows=self.audit_list(course_id)["items"];headers=["时间","操作者","角色","动作","对象类型","对象标识","课程标识","班级标识","学生标识","请求标识","来源地址","结果","原因","详情"]
        return self.make_xlsx("审计",headers,[[x["occurred_at"],x["actor_user_id"],x["actor_role"],x["action"],x["resource_type"],x["resource_id"],x["course_id"],x["class_id"],x["student_id"],x["request_id"],x["ip"],x["result"],x["reason"],json.dumps(x["details"],ensure_ascii=False,sort_keys=True)] for x in rows],course_id or "授权范围")
    def audit_csv(self,course_id=None):
        import csv
        text=__import__('io').StringIO();writer=csv.writer(text);writer.writerow(["时间","操作者","角色","动作","对象类型","对象标识","课程标识","班级标识","学生标识","请求标识","来源地址","结果","原因","详情"])
        for x in self.audit_list(course_id)["items"]:writer.writerow([x[k] for k in ["occurred_at","actor_user_id","actor_role","action","resource_type","resource_id","course_id","class_id","student_id","request_id","ip","result","reason"]]+[json.dumps(x["details"],ensure_ascii=False,sort_keys=True)])
        return ('\ufeff'+text.getvalue()).encode('utf-8')

    def authoritative_component_scores(self, events):
        grouped=defaultdict(list);labs=defaultdict(list)
        for event in events:
            component=self.component(event.source_type)
            if component=="LAB":
                payload=event.payload_json or {}
                lab_key=payload.get("lab_release_id") or payload.get("runtime_instance_id") or event.lesson_id or "legacy-lab"
                labs[lab_key].append(event)
            elif component in COMPONENTS:
                grouped[component].append(Decimal(event.normalized_score))
        for rows in labs.values():
            submissions=[event for event in rows if event.source_type=="LAB_SUBMISSION"]
            if submissions:
                latest=max(submissions,key=self.event_order)
                grouped["LAB"].append(Decimal(latest.normalized_score))
        return grouped

    def build_risks(self,book,student_id,grouped,total,present,student_events,all_events):
        risks=[]
        if grouped.get("ATTENDANCE") and sum(grouped["ATTENDANCE"])/len(grouped["ATTENDANCE"])<80:risks.append(("LOW_ATTENDANCE",{"attendance_rate":float(sum(grouped["ATTENDANCE"])/len(grouped["ATTENDANCE"]))}))
        if total<60:risks.append(("LOW_TOTAL_SCORE",{"total_score":float(total),"threshold":60}))
        missing=[x for x in COMPONENTS if x not in present]
        if missing:risks.append(("MISSING_FACTS",{"missing_components":missing}))
        checkpoint_streak=[];longest_streak=[]
        for event in sorted((x for x in student_events if x.source_type=="LAB_CHECKPOINT"),key=self.event_order):
            if event.payload_json.get("source_event_type")=="lab.checkpoint.failed":
                checkpoint_streak.append(event)
                if len(checkpoint_streak)>len(longest_streak):longest_streak=list(checkpoint_streak)
            else:checkpoint_streak=[]
        if len(longest_streak)>=2:risks.append(("CONSECUTIVE_CHECKPOINT_FAILURES",{"failure_count":len(longest_streak),"source_ids":[x.source_id for x in longest_streak]}))
        available_labs={x.payload_json.get("lab_release_id") for x in all_events if x.source_type.startswith("LAB_") and x.payload_json.get("lab_release_id")};submitted_labs={x.payload_json["lab_release_id"] for x in student_events if x.source_type=="LAB_SUBMISSION"};missing_labs=sorted(available_labs-submitted_labs)
        if len(missing_labs)>=2:risks.append(("MULTIPLE_MISSING_SUBMISSIONS",{"missing_count":len(missing_labs),"lab_release_ids":missing_labs}))
        for kind,evidence in risks:self.repo.add(StudentRiskFlag(student_risk_flag_id=str(uuid4()),gradebook_id=book.gradebook_id,course_id=book.course_id,class_id=book.class_id,student_id=student_id,risk_type=kind,evidence_json=evidence,status="OPEN",created_at=utcnow()))
    def build_section_analytics(self,book,events):
        grouped=defaultdict(list)
        for e in events:
            if e.lesson_id:grouped[e.lesson_id].append(e)
        for lesson_id,rows in grouped.items():
            per=defaultdict(list)
            for e in rows:per[e.student_id].append(Decimal(e.normalized_score))
            scores=[sum(x)/len(x) for x in per.values()];dist={"0-59":0,"60-69":0,"70-79":0,"80-89":0,"90-100":0}
            for score in scores:dist["0-59" if score<60 else "60-69" if score<70 else "70-79" if score<80 else "80-89" if score<90 else "90-100"]+=1
            ranking=sorted([{"student_id":s,"score":float(sum(v)/len(v))} for s,v in per.items()],key=lambda x:x["score"],reverse=True)
            self.repo.add(AnalyticsSectionSummary(analytics_section_summary_id=str(uuid4()),gradebook_id=book.gradebook_id,lesson_id=lesson_id,summary_json={"status":"READY","lesson_id":lesson_id,"average":self.average(scores),"distribution":dist,"ranking":ranking}))
    def build_lab_analytics(self,book,events,student_count):
        students=sorted({event.student_id for event in events});labs=sorted({event.payload_json.get("lab_release_id") for event in events if event.source_type.startswith("LAB_") and event.payload_json.get("lab_release_id")})
        latest={}
        for event in (row for row in events if row.source_type=="LAB_SUBMISSION"):
            key=(event.student_id,event.payload_json["lab_release_id"]);previous=latest.get(key)
            if not previous or self.event_order(event)>self.event_order(previous):latest[key]=event
        for student_id in students:
            rows=[event for (owner,_),event in latest.items() if owner==student_id]
            self.repo.add(AnalyticsStudentLabSummary(analytics_student_lab_summary_id=str(uuid4()),gradebook_id=book.gradebook_id,student_id=student_id,sum_lab_score=sum((Decimal(x.normalized_score) for x in rows),Decimal("0")),submitted_count=len(rows),unsubmitted_count=max(0,len(labs)-len(rows))))
        for lab_id in labs:
            rows=[event for (_,release),event in latest.items() if release==lab_id]
            source_rows=[event for event in events if event.payload_json.get("lab_release_id")==lab_id]
            self.repo.add(AnalyticsLabSummary(analytics_lab_summary_id=str(uuid4()),gradebook_id=book.gradebook_id,lab_release_id=lab_id,max_score=max((Decimal(x.max_score) for x in source_rows),default=Decimal("100")),submitted_students=len(rows),unsubmitted_students=max(0,student_count-len(rows)),avg_score=Decimal(str(self.average([Decimal(x.normalized_score) for x in rows]) or 0))))
    @staticmethod
    def component(source):return "LAB" if source.startswith("LAB_") else source
    @staticmethod
    def event_order(event):return ((event.payload_json or {}).get("source_occurred_at") or event.occurred_at.isoformat(),event.event_id)
    @staticmethod
    def average(values):return round(float(sum(values)/len(values)),2) if values else None
    def policy_dict(self,p):return {"course_id":p.course_id,"version_no":p.version_no,"status":p.status,"effective_at":p.effective_at.isoformat()+"Z","items":[{"component":x.component,"weight_percent":number(x.weight_percent)} for x in self.repo.policy_items(p.grading_policy_id)]}
    @staticmethod
    def gradebook_dict(b):return {"gradebook_id":b.gradebook_id,"course_id":b.course_id,"class_id":b.class_id,"policy_version":b.policy_version,"status":b.status,"calculated_at":b.calculated_at.isoformat()+"Z","posted_at":b.posted_at.isoformat()+"Z" if b.posted_at else None}
    @staticmethod
    def event_dict(e):return {"grade_event_id":e.grade_event_id,"source_type":e.source_type,"source_id":e.source_id,"lab_release_id":e.payload_json.get("lab_release_id"),"lesson_id":e.lesson_id,"raw_score":number(e.raw_score),"max_score":number(e.max_score),"normalized_score":number(e.normalized_score),"source_verification_status":e.source_verification_status,"source_proof_issuer":e.source_proof_issuer,"source_proof_digest":e.source_proof_digest,"occurred_at":e.occurred_at.isoformat()+"Z","event_id":e.event_id}
    @staticmethod
    def audit_dict(x):return {"audit_event_id":x.audit_event_id,"source_event_id":x.source_event_id,"actor_user_id":x.actor_user_id,"actor_role":x.actor_role,"action":x.action,"resource_type":x.resource_type,"resource_id":x.resource_id,"course_id":x.course_id,"class_id":x.class_id,"student_id":x.student_id,"request_id":x.request_id,"ip":x.ip,"result":x.result,"reason":x.reason,"occurred_at":x.occurred_at.isoformat()+"Z","details":x.details_json}
    @staticmethod
    def make_xlsx(title,headers,rows,course_id):
        book=Workbook();sheet=book.active;sheet.title=title;sheet.append(["课程/范围",course_id]);sheet.append(["生成时间",utcnow().isoformat()+"Z"]);sheet.append([]);sheet.append(headers)
        for row in rows:sheet.append(row)
        output=BytesIO();book.save(output);return output.getvalue()
