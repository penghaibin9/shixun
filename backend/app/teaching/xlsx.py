from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill


HEADERS = ["学号*", "姓名*", "班级", "手机号（可选）", "邮箱（可选）"]


def template_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "学生导入"
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4267D5")
    sheet.freeze_panes = "A2"
    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 14
    sheet.column_dimensions["C"].width = 22
    sheet.column_dimensions["D"].width = 18
    sheet.column_dimensions["E"].width = 28
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def parse_members(data: bytes) -> tuple[list[dict], list[dict]]:
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=False)
    except Exception as exc:
        raise ValueError("文件不是有效的 XLSX") from exc
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows or list(rows[0][:5]) != HEADERS:
        raise ValueError("模板表头不匹配")
    valid, errors, seen = [], [], set()
    for row_number, row in enumerate(rows[1:], start=2):
        values = list(row[:5]) + [None] * max(0, 5 - len(row))
        if all(value in (None, "") for value in values):
            errors.append({"row_number": row_number, "student_number": "", "student_name": "", "class_name": "", "phone": "", "email": "", "reason": "空白行"})
            continue
        number, name, class_name, phone, email = values[:5]
        reasons = []
        for value in values:
            if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
                reasons.append("包含公式或危险前缀")
                break
        number, name = str(number or "").strip(), str(name or "").strip()
        if not number:
            reasons.append("学号不能为空")
        if not name:
            reasons.append("姓名不能为空")
        if number in seen:
            reasons.append("批次内学号重复")
        seen.add(number)
        if len(number) > 64 or len(name) > 80 or len(str(phone or "")) > 32 or len(str(email or "")) > 160:
            reasons.append("字段长度超过限制")
        item = {"row_number": row_number, "student_number": number, "student_name": name, "class_name": str(class_name or ""), "phone": str(phone or ""), "email": str(email or "")}
        if reasons:
            errors.append({**item, "reason": "；".join(reasons)})
        else:
            valid.append(item)
    return valid, errors


def error_rows_bytes(errors: list[dict]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "错误行"
    sheet.append(["原行号", *HEADERS, "错误原因"])
    for item in errors:
        sheet.append([item["row_number"], item["student_number"], item["student_name"], item["class_name"], item["phone"], item["email"], item["reason"]])
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()
